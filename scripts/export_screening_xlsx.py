"""Export the whole screening record to a single Excel workbook.

WHY. Without a screening platform, the workbook IS the artefact a human reviewer
reads and a referee would ask for. It has to carry every stage, every reason and
every author, not just the surviving records -- a screening log that shows only
what passed cannot be audited.

SHEETS
  Resumen              counts per stage, and what each stage did
  Candidatos           records awaiting full text -- the working list
  Etapa2_titulo        every title decision with reason and author
  Etapa3_resumen       every abstract decision with reason and author
  Etapa1_reglas        the rule-based exclusions, counted by rule

Regenerate after every screening batch; it is a view, never a source.

USAGE
    python scripts/export_screening_xlsx.py
"""
import collections
import csv
import pathlib

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("openpyxl required: python -m pip install openpyxl")

ROOT = pathlib.Path(__file__).resolve().parent.parent
STAGE1 = ROOT / "revision_sistematica" / "cribado" / "screening_stage1.csv"
STAGE2 = ROOT / "revision_sistematica" / "cribado" / "screening_stage2_decisions.csv"
STAGE3 = ROOT / "revision_sistematica" / "cribado" / "screening_stage3_decisions.csv"
INDEX = ROOT / "quality_reports" / "corpus_identifier_index.txt"
OUT = ROOT / "quality_reports" / "cribado.xlsx"

csv.field_size_limit(10_000_000)

HEAD = PatternFill("solid", fgColor="1F4E79")
HEADF = Font(color="FFFFFF", bold=True)
YES = PatternFill("solid", fgColor="E2EFDA")


def last_per_pmid(path):
    if not path.exists():
        return {}
    with open(path, encoding="utf-8", newline="") as fh:
        out = {}
        for r in csv.DictReader(fh):
            out[r["pmid"]] = r
        return out


def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).fill = HEAD
        ws.cell(1, c).font = HEADF
        ws.cell(1, c).alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    with open(STAGE1, encoding="utf-8", newline="") as fh:
        meta = {r["pmid"]: r for r in csv.DictReader(fh)}
    s2 = last_per_pmid(STAGE2)
    s3 = last_per_pmid(STAGE3)

    known = {}
    for line in INDEX.read_text(encoding="utf-8").splitlines():
        if line.startswith("pmid:"):
            ident, _, sid = line.partition("\t")
            known[ident.split(":", 1)[1]] = sid.strip()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- Resumen ----
    s1_excl = collections.Counter(
        r["stage1_rule"] for r in meta.values() if r["stage1"] == "EXCLUDED")
    n_pool = sum(1 for r in meta.values() if r["stage1"] == "ADVANCE")
    n2_adv = sum(1 for r in s2.values() if r["verdict"] == "ADVANCE")
    n3_ft = sum(1 for r in s3.values() if r["verdict"] == "FULLTEXT")

    res = [
        ["Registros recuperados (PubMed, dos brazos, union)", len(meta), ""],
        ["Etapa 1 - excluidos por regla automatica", sum(s1_excl.values()),
         "PRISMA: 'records removed before screening'"],
        ["Etapa 2 - pozo de cribado por titulo", n_pool, ""],
        ["Etapa 2 - decididos", len(s2), ""],
        ["Etapa 2 - pasan a resumen", n2_adv, ""],
        ["Etapa 3 - decididos", len(s3), ""],
        ["Etapa 3 - pasan a texto completo", n3_ft, "lista de trabajo"],
        ["", "", ""],
        ["Estudios ya incluidos en el corpus (control)", len(known),
         "ninguno debe quedar excluido en ninguna etapa"],
    ]
    sheet(wb, "Resumen", ["Concepto", "n", "Nota"], res, [52, 10, 46])

    # ---- Candidatos ----
    cand = []
    for p, d in s3.items():
        if d["verdict"] != "FULLTEXT":
            continue
        m = meta.get(p, {})
        cand.append([
            "SI" if p in known else "",
            p, m.get("year", ""), " ".join((m.get("title") or "").split()),
            m.get("journal", ""), m.get("doi", ""), d["reason"],
            "https://pubmed.ncbi.nlm.nih.gov/%s/" % p,
        ])
    cand.sort(key=lambda r: (r[2], r[3]))
    ws = sheet(wb, "Candidatos",
               ["Ya incluido", "PMID", "Anio", "Titulo", "Revista", "DOI",
                "Motivo de avance", "PubMed"],
               cand, [11, 11, 7, 70, 30, 26, 58, 40])
    for i in range(2, ws.max_row + 1):
        if ws.cell(i, 1).value == "SI":
            for c in range(1, 9):
                ws.cell(i, c).fill = YES

    # ---- Etapa 2 ----
    rows2 = []
    for p, d in s2.items():
        m = meta.get(p, {})
        rows2.append([p, m.get("year", ""), " ".join((m.get("title") or "").split()),
                      d["verdict"], d["reason"], d["decided_by"], d["decided_at"],
                      "SI" if p in known else ""])
    rows2.sort(key=lambda r: (r[3], r[1], r[0]))
    sheet(wb, "Etapa2_titulo",
          ["PMID", "Anio", "Titulo", "Decision", "Motivo", "Decidido por",
           "Fecha", "Ya incluido"],
          rows2, [11, 7, 70, 11, 60, 40, 20, 11])

    # ---- Etapa 3 ----
    rows3 = []
    for p, d in s3.items():
        m = meta.get(p, {})
        rows3.append([p, m.get("year", ""), " ".join((m.get("title") or "").split()),
                      d["verdict"], d["reason"], d["decided_by"], d["decided_at"],
                      "SI" if p in known else ""])
    rows3.sort(key=lambda r: (r[3], r[1], r[0]))
    sheet(wb, "Etapa3_resumen",
          ["PMID", "Anio", "Titulo", "Decision", "Motivo", "Decidido por",
           "Fecha", "Ya incluido"],
          rows3, [11, 7, 70, 11, 60, 40, 20, 11])

    # ---- Etapa 1 ----
    sheet(wb, "Etapa1_reglas", ["Regla aplicada", "Registros excluidos"],
          sorted(([k or "(sin regla)", v] for k, v in s1_excl.items()),
                 key=lambda r: -r[1]), [95, 20])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("etapa 1 excluidos : %d" % sum(s1_excl.values()))
    print("etapa 2 decididos : %d (%d avanzan)" % (len(s2), n2_adv))
    print("etapa 3 decididos : %d (%d a texto completo)" % (len(s3), n3_ft))
    print("candidatos        : %d" % len(cand))
    print("\nwrote", OUT)


if __name__ == "__main__":
    main()
