"""Genera el formulario CIEGO de riesgo de sesgo de un revisor, en español.

CIEGO SIGNIFICA CIEGO, igual que en `make_extraction_forms.py`: ni un valor
previo, ni el del otro revisor. La concordancia entre D. Valdiviezo y N. Trelles
solo significa algo si el segundo no vio lo que puso el primero.

EL JUICIO ES DE LOS AUTORES. Este guion prepara las preguntas y pone al lado las
frases del artículo que las tocan (`evidencia_riesgo_sesgo.py`), pero no
responde ninguna. El manuscrito reportará la evaluación como hecha por los dos
autores, y lo será.

UN INSTRUMENTO POR DISEÑO, PORQUE NO SON INTERCAMBIABLES. RoB 2 y ROBINS-I no
sirven para un reporte de caso --no hay grupo de comparación, ni asignación, ni
seguimiento--, y aplicarlos daría "alto riesgo" en todos los dominios por
razones que no son un defecto del estudio sino del instrumento. De ahí las dos
listas JBI, que son las que el campo usa para casos y series.

LOS QUE NO SE EVALÚAN TAMBIÉN SE LISTAN. Los 24 estudios sin texto completo van
en su propia hoja marcados como no evaluables. No es un hueco: es el dato que
§3.2 ya mide como sesgo de recuperación, y la fracción que falta concentra 6 de
los 11 estudios comparativos. Si esa hoja no estuviera, la evaluación parecería
más completa de lo que es.

USO
    python scripts/make_rob_forms.py --revisor "Danny Valdiviezo"
    python scripts/make_rob_forms.py --revisor "Nataly Trelles"
"""
import argparse
import collections
import csv
import datetime
import pathlib
import sys

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    raise SystemExit("hace falta openpyxl: python -m pip install openpyxl")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from rob_instruments import (INSTRUMENTOS, SIMPLIFICADOS, COMPARATIVOS,
                             POR_DISENO, NO_EVALUABLE, PENDIENTE)

ROOT = pathlib.Path(__file__).resolve().parent.parent
RS = ROOT / "revision_sistematica"
DEST = RS / "riesgo_sesgo"
EVIDENCIA = DEST / "evidencia_por_dominio.csv"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CAB = PatternFill("solid", fgColor="1F4E79")
CABF = Font(color="FFFFFF", bold=True)
FIJO = PatternFill("solid", fgColor="EDEDED")     # contexto, no se toca
RELLENA = PatternFill("solid", fgColor="FFF2CC")  # lo que rellena el revisor
AVISO = PatternFill("solid", fgColor="FCE4D6")    # no evaluable


def lee(p, clave=None):
    with open(p, encoding="utf-8-sig", newline="") as fh:
        filas = list(csv.DictReader(fh))
    return {f[clave]: f for f in filas} if clave else filas


def corpus():
    """Los 95 con publicación recuperable, con su diseño y su instrumento."""
    fuera = {f["study_id"] for f in
             lee(RS / "cribado" / "exclusiones_tras_texto_completo.csv")}
    legible = {f["study_id"] for f in
               lee(RS / "textos_completos" / "texto_cache" / "_inventario.csv")
               if f.get("estado") == "legible"}
    # `informe_para_extraer` marca el informe representativo de CADA estudio, los
    # 155, incluidos los 60 que solo existen como ficha de registro. Lo que
    # decide si hay algo que evaluar es `situacion`: solo los extraibles y los
    # que existen como resumen de congreso tienen publicacion recuperable, y son
    # los 95. Filtrar por el campo equivocado metia 60 fichas de registro en la
    # hoja de "no evaluables" como si fueran textos que no se consiguieron.
    EXTRAIBLE = {"extraible", "solo-resumen"}
    grupos = {}
    for g in lee(RS / "cribado" / "study_groups.csv"):
        s = "EST-%03d" % int(g["estudio"])
        if (s not in grupos and g.get("informe_para_extraer") == "SI"
                and g.get("situacion") in EXTRAIBLE):
            grupos[s] = g

    # El diseño sale de la extracción adjudicada, que es la que se leyó sobre el
    # articulo; el del resumen no vale para elegir instrumento.
    diseno = {}
    for f in lee(RS / "extraccion" / "extraccion_adjudicada.csv"):
        s, v = f["study_id"], (f.get("study_design") or "").strip()
        if s not in fuera and v and v != "NA" and s not in diseno:
            diseno[s] = v

    # La capa de correcciones va encima, como en el resto del canal: es donde se
    # anotan las adjudicaciones posteriores con su cita y su firma. Nunca se
    # edita `extraccion_adjudicada.csv`, que es el registro de lo que dieron las
    # dos lecturas.
    corr = RS / "extraccion" / "correcciones_tras_texto_completo.csv"
    if corr.exists():
        for f in lee(corr):
            if f.get("campo") == "study_design" and f["study_id"] not in fuera:
                diseno[f["study_id"]] = f["valor_corregido"].strip()

    salida = []
    for s in sorted(grupos):
        if s in fuera:
            continue
        g = grupos[s]
        if s not in legible:
            inst = NO_EVALUABLE
        else:
            inst = POR_DISENO.get(diseno.get(s, ""), PENDIENTE)
        salida.append({"study_id": s, "instrumento": inst,
                       "diseno": diseno.get(s, "(no clasificable)"),
                       "titulo": g.get("titulo", ""), "revista": g.get("revista", ""),
                       "anio": g.get("anio", "")})
    return salida


def evidencia():
    if not EVIDENCIA.exists():
        return {}
    d = collections.defaultdict(lambda: collections.defaultdict(list))
    for f in lee(EVIDENCIA):
        if f["frase"]:
            d[f["study_id"]][f["dominio"]].append(f["frase"])
    return d


CONTEXTO = [("study_id", "Nº", 11, "Clave con la que se cruzan las dos evaluaciones. NO la modifique."),
            ("diseno", "Diseño", 20, "Diseño adjudicado en la extracción. Si no cuadra, dígalo en la última columna."),
            ("titulo", "Título del artículo", 60, ""),
            ("revista", "Revista", 26, ""),
            ("anio", "Año", 6, "")]


def hoja(wb, inst, filas, ev):
    ws = wb.create_sheet(inst["hoja"][:31])
    ws.freeze_panes = "C3"

    ws["A1"] = inst["nombre"]
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = inst["cuando"] + "   ·   Fuente: " + inst["fuente"]
    ws["A2"].font = Font(italic=True, size=9)

    cabeceras = [c[1] for c in CONTEXTO]
    for it in inst["items"]:
        cabeceras.append("%s. %s" % (it["codigo"], it["texto_es"]))
    if inst.get("juicios"):
        cabeceras.append("JUICIO GLOBAL del estudio")
    cabeceras += ["¿En qué frase te apoyaste?", "Dudas o desacuerdo con el diseño"]

    for j, c in enumerate(cabeceras, 1):
        cel = ws.cell(row=3, column=j, value=c)
        cel.fill, cel.font = CAB, CABF
        cel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 78

    n_ctx = len(CONTEXTO)
    for j, (_, _, ancho, _) in enumerate(CONTEXTO, 1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    for j in range(n_ctx + 1, len(cabeceras) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 30

    # Desplegables: uno para responder los items, otro para el juicio global.
    # Una validacion por juego de opciones. RoB 2 no admite las mismas
    # respuestas en todas sus preguntas --3.2 no lleva "Sin informacion"-- y una
    # unica lista para toda la hoja habria ofrecido una opcion que el
    # instrumento prohibe.
    validaciones = {}

    def valida(opciones):
        clave = "|".join(opciones)
        if clave not in validaciones:
            d = DataValidation(type="list", allow_blank=True,
                               formula1='"%s"' % ",".join(opciones))
            ws.add_data_validation(d)
            validaciones[clave] = d
        return validaciones[clave]

    dvj = valida(inst["juicios"]) if inst.get("juicios") else None

    for i, f in enumerate(filas, start=4):
        for j, (campo, _, _, nota) in enumerate(CONTEXTO, 1):
            cel = ws.cell(row=i, column=j, value=f.get(campo, ""))
            cel.fill = FIJO
            cel.alignment = Alignment(wrap_text=(campo == "titulo"), vertical="top")
        for k, it in enumerate(inst["items"]):
            cel = ws.cell(row=i, column=n_ctx + 1 + k)
            cel.fill = RELLENA
            valida(it.get("respuestas") or inst["respuestas"]).add(cel)
            # Las frases del dominio, como comentario de la celda: se leen al
            # pasar por encima y no ensucian la hoja.
            frases = ev.get(f["study_id"], {}).get(it.get("dominio_evidencia", ""), [])
            if frases:
                txt = "\n\n".join(frases)[:2000]
                cel.comment = Comment("Del artículo:\n\n" + txt, "evidencia")
        col = n_ctx + 1 + len(inst["items"])
        if inst.get("juicios"):
            cel = ws.cell(row=i, column=col)
            cel.fill = RELLENA
            dvj.add(cel)
            col += 1
        for _ in range(2):
            ws.cell(row=i, column=col).fill = RELLENA
            col += 1
    return ws


def hoja_no_evaluables(wb, filas, titulo, cabecera, porque):
    ws = wb.create_sheet(titulo[:31])
    ws["A1"] = cabecera
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = porque
    ws["A2"].font = Font(italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True)
    for j, (campo, cab, ancho, _) in enumerate(CONTEXTO, 1):
        cel = ws.cell(row=4, column=j, value=cab)
        cel.fill, cel.font = CAB, CABF
        ws.column_dimensions[get_column_letter(j)].width = ancho
    for i, f in enumerate(filas, start=5):
        for j, (campo, _, _, _) in enumerate(CONTEXTO, 1):
            ws.cell(row=i, column=j, value=f.get(campo, "")).fill = AVISO
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--revisor", required=True)
    ap.add_argument("--simplificado", action="store_true",
                    help="solo los estudios con grupo de comparación, y por "
                         "dominio en vez de pregunta a pregunta")
    a = ap.parse_args()

    filas, ev = corpus(), evidencia()
    # En modo simplificado se evalua SOLO lo comparativo. Los reportes y las
    # series salen del formulario entero, no a una hoja de descartados: no es
    # que no se pudieran evaluar, es que se decidio no evaluarlos, y son cosas
    # distintas que Metodos tiene que distinguir.
    instrumentos = INSTRUMENTOS
    if a.simplificado:
        instrumentos = SIMPLIFICADOS
        filas = [f for f in filas if f["diseno"] in COMPARATIVOS]
    if not ev:
        print("AVISO: no hay evidencia_por_dominio.csv. Ejecuta antes:")
        print("       python scripts/evidencia_riesgo_sesgo.py --csv")

    por = collections.defaultdict(list)
    for f in filas:
        por[f["instrumento"]].append(f)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Cómo se rellena")
    ws.column_dimensions["A"].width = 108
    guia = [
        ("Evaluación del riesgo de sesgo — %s" % a.revisor, True),
        ("Generado el %s" % datetime.date.today().isoformat(), False),
        ("", False),
        ("Una hoja por instrumento. A cada estudio le toca uno según su diseño: no son "
         "intercambiables, y aplicar RoB 2 a un reporte de caso da alto riesgo en todo "
         "por un defecto del instrumento, no del estudio.", False),
        ("", False),
        ("Las celdas amarillas se rellenan; las grises no se tocan.", False),
        ("Muchas celdas llevan un comentario (esquina roja): son las frases del propio "
         "artículo que hablan de ese dominio. Léelas, pero no te fíes solo de ellas — "
         "que un dominio no tenga frases NO significa que el estudio no informe: "
         "significa que el buscador no vio la señal.", False),
        ("", False),
        ("Trabaja a ciegas: no consultes el cuaderno del otro revisor. Los desacuerdos "
         "se resuelven después, por consenso, como se hizo con los 541 de la extracción.", False),
        ("", False),
        ("Si crees que el diseño adjudicado está mal, no fuerces el instrumento: "
         "escríbelo en la última columna y se decide aparte.", False),
    ]
    for i, (t, neg) in enumerate(guia, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=neg, size=13 if neg and i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    resumen = []
    for inst in instrumentos:
        f = por.get(inst["clave"], [])
        if f:
            hoja(wb, inst, f, ev)
        resumen.append((inst["hoja"], len(f)))
    if por.get(NO_EVALUABLE):
        hoja_no_evaluables(
            wb, por[NO_EVALUABLE], "No evaluables",
            "Estudios sin texto completo: no se puede evaluar el riesgo de sesgo",
            "No se rellenan. Van aquí para que consten: la fracción sin texto "
            "concentra 6 de los 11 estudios comparativos, y la evaluación de riesgo "
            "de sesgo hereda ese sesgo de recuperación.")
    if por.get(PENDIENTE):
        hoja_no_evaluables(
            wb, por[PENDIENTE], "Diseño pendiente",
            "Diseño pendiente de decidir: sin diseño no se puede elegir instrumento",
            "Estos estudios se leyeron pero su diseño no quedó clasificado. Decidid "
            "primero qué son y volved a generar el formulario; forzar un instrumento "
            "sobre un diseño que no se sabe da un juicio que no significa nada.")

    DEST.mkdir(parents=True, exist_ok=True)
    slug = a.revisor.lower().replace(" ", "_")
    # Nombre distinto: los dos formularios no son intercambiables y confundirlos
    # mezclaria una evaluacion por dominio con otra pregunta a pregunta.
    sufijo = "_comparativos" if a.simplificado else ""
    salida = DEST / ("riesgo_sesgo%s_%s.xlsx" % (sufijo, slug))
    try:
        wb.save(salida)
    except PermissionError:
        salida = salida.with_name(salida.stem + "_NUEVO.xlsx")
        wb.save(salida)
        print("AVISO: el fichero estaba abierto; se escribe al lado.")
    print("escrito %s" % salida.relative_to(ROOT))
    for h, n in resumen:
        print("   %-26s %3d estudios" % (h, n))
    print("   %-26s %3d estudios" % ("No evaluables (sin texto)", len(por.get(NO_EVALUABLE, []))))
    print("   %-26s %3d estudios" % ("Diseño pendiente", len(por.get(PENDIENTE, []))))
    print("   %-26s %3d estudios" % ("TOTAL", len(filas)))


if __name__ == "__main__":
    main()
