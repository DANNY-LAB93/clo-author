# -*- coding: utf-8 -*-
"""Aplica la ampliacion de NOREC (codigo NOPDF) cuando lleva las dos firmas.

Lee `~/Desktop/FIRMAR_texto_completo_2026-10-06.xlsx` (lo escribe
`make_firma_nopdf.py`). Sin las dos firmas, o sin «APROBAR», no hace nada y lo
dice. Con ellas comprueba que la lista firmada es la misma que la de hoy --si el
corpus cambio por otro camino, una firma sobre otra lista no vale-- y anade una
fila por estudio a `exclusiones_tras_texto_completo.csv` con el codigo NOPDF.
Es idempotente.

Uso:
    python scripts/ingest_firma_nopdf.py              # informe, no escribe
    python scripts/ingest_firma_nopdf.py --escribir
"""
import csv
import pathlib
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from ingest_firma_lectura import son_los_dos, fecha  # noqa: E402
from make_firma_nopdf import candidatos  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
RS = ROOT / "revision_sistematica"
CUADERNO = pathlib.Path.home() / "Desktop" / "FIRMAR_texto_completo_2026-10-06.xlsx"
COPIA = RS / "lectura_pendiente" / "FIRMAR_texto_completo_2026-10-06_firmado.xlsx"
EXCL = RS / "cribado" / "exclusiones_tras_texto_completo.csv"
MOTIVO = {
    "extraible": "Artículo publicado cuyo texto completo no se obtuvo",
    "solo-resumen": "Solo existe como resumen de congreso",
    "solo-registro": "Solo existe como ficha de registro de ensayo, sin artículo publicado",
}


def main():
    escribir = "--escribir" in sys.argv
    if not CUADERNO.exists():
        print("no existe %s: nada que aplicar" % CUADERNO.name)
        return
    wb = load_workbook(CUADERNO, data_only=True)
    fi = wb["Firma"]
    r1, r2 = (str(fi.cell(row=i, column=2).value or "").strip() for i in (4, 5))
    f_firma = fecha(fi.cell(row=6, column=2).value)
    if not son_los_dos(r1, r2):
        print("cuaderno de texto completo sin las dos firmas («%s», «%s»): no se aplica" % (r1, r2))
        return
    h = wb["La enmienda"]
    dec = ""
    for row in h.iter_rows(min_col=1, max_col=2):
        if row[0].value == "Vuestra decisión":
            dec = str(row[1].value or "").strip()
    if not dec.startswith("APROBAR"):
        print("decisión «%s»: no se aplica" % dec)
        return
    hoja = next(ws for ws in wb.worksheets if ws.title.startswith("Los "))
    firmados = [str(hoja.cell(row=i, column=1).value) for i in range(2, hoja.max_row + 1)
                if hoja.cell(row=i, column=1).value]
    hoy = candidatos()
    ya = {r["study_id"] for r in csv.DictReader(open(EXCL, encoding="utf-8"))}
    pendientes = [x for x in hoy if x["estudio"] not in ya]
    if sorted(firmados) != sorted(x["estudio"] for x in hoy) and pendientes:
        print("NO SE ESCRIBE: la lista firmada (%d) no es la de hoy (%d)." % (len(firmados), len(hoy)))
        sys.exit(1)
    nuevas = [{"study_id": x["estudio"], "codigo": "NOPDF",
               "motivo": "%s. Criterio: solo estudios con el PDF del artículo completo accesible "
                         "(ampliación de NOREC firmada el %s)." % (MOTIVO[x["situacion"]], f_firma),
               "cita_del_texto": "(sin PDF del artículo completo)",
               "titulo": x["titulo"], "decidido_por": "D. Valdiviezo y N. Trelles", "fecha": f_firma}
              for x in pendientes]
    print("FIRMADO EL %s por %s y %s: %s" % (f_firma, r1, r2, dec))
    print("exclusiones nuevas: %d" % len(nuevas))
    if not escribir:
        print("(informe; nada escrito. Con --escribir se aplica.)")
        return
    shutil.copy2(CUADERNO, COPIA)
    if nuevas:
        cab = list(csv.DictReader(open(EXCL, encoding="utf-8")).fieldnames)
        with open(EXCL, "a", encoding="utf-8", newline="") as fh:
            csv.DictWriter(fh, fieldnames=cab).writerows(nuevas)
    print("añadidas %d exclusiones NOPDF" % len(nuevas))


if __name__ == "__main__":
    main()
