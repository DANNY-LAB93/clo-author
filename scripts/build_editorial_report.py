"""Autoevaluacion del manuscrito (S0), indice del paquete y manuscritos en Word.

QUIEN ESCRIBE S0. El modelo de lenguaje (Claude Opus 5), a peticion de los
autores. Adopta el punto de vista de un editor que decide si el manuscrito
puede entrar en revision por pares, pero NO es un informe editorial ni una
revision externa, y el documento lo dice en su titulo, en su primera frase y en
la cabecera de cada pagina. Hasta el 2026-10-08 este docstring decia «EL PAPEL
QUE SE ADOPTA AQUI. El de un editor de revista» y S0 se presentaba como
«Informe editorial», en primera persona de editor, sin nombrar al modelo: un
lector del comite podia tomarlo por la valoracion de una persona.

No es un resumen amable de lo hecho: es la lista de lo que falta, con la
consecuencia de cada carencia y quien puede resolverla. Una autoevaluacion que
solo elogia no sirve para publicar.

EL INDICE. Sus estados y descripciones salen del canal y de los registros, no
de frases fijas: hasta el 2026-10-08 daba el riesgo de sesgo por pendiente, el
texto completo por «PARCIAL — 100 %» y solo listaba S0 a S15. Se comprueba en
los dos sentidos: que lo citado existe y que no queda fichero sin citar.

SALIDA
    verificables revisión sistemática/00_INDICE.docx
    verificables revisión sistemática/S0_informe_editorial.docx
    verificables revisión sistemática/manuscrito_*.docx
"""
import csv
import json
import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
# La fecha larga se define una sola vez, en el constructor del paquete. Dos
# anexos la imprimen y no pueden discrepar.
from build_verifiables_package import fecha_larga, MESES

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "verificables revisión sistemática"
RS = ROOT / "revision_sistematica"
QR = ROOT / "quality_reports"
# El limite de palabras del resumen no es un dato del canal: es la norma
# habitual de las revistas, y S0 lo compara con el recuento real.
LIMITE_RESUMEN = 250
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def doc_nuevo(titulo, subtitulo=""):
    d = Document()
    est = d.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(10.5)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)
    d.add_heading(titulo, level=0).alignment = WD_ALIGN_PARAGRAPH.LEFT
    if subtitulo:
        p = d.add_paragraph(subtitulo)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9.5)
    return d


def tabla(d, cabecera, filas, anchos=None):
    t = d.add_table(rows=1, cols=len(cabecera))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(cabecera):
        cel = t.rows[0].cells[i]
        cel.text = ""
        r = cel.paragraphs[0].add_run(str(c))
        r.bold = True
        r.font.size = Pt(9)
    for f in filas:
        cs = t.add_row().cells
        for i, v in enumerate(f):
            cs[i].text = ""
            r = cs[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(9)
    if anchos:
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
    d.add_paragraph()


def nota(d, texto):
    p = d.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(8.5)
    r.italic = True
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


# ------------------------------------------------- cifras que no son escalares
# S0 y el indice citaban cifras tecleadas que envejecieron («dos» informes en
# ruso donde el registro anota tres, «47 estudios con PDF» cuando son 65). Lo
# que no esta en synthesis_scalars.json se cuenta aqui, desde el fichero que lo
# registra, cada vez que se regenera el paquete.

def leer_csv(p):
    """Filas de un CSV del canal; lista vacia si el fichero no esta."""
    p = pathlib.Path(p)
    if not p.exists():
        return []
    with open(p, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def miles(n):
    """7418 -> «7 418», como los manuscritos."""
    return "{:,}".format(n).replace(",", " ")


def fecha_de(iso):
    """«2026-09-14» -> «14 de septiembre de 2026»."""
    a, m, d = (int(x) for x in iso[:10].split("-"))
    return "%d de %s de %d" % (d, MESES[m - 1], a)


def entre_fechas(a, b):
    """Dos fechas ISO -> «entre el 9 y el 22 de septiembre de 2026»."""
    if not a or not b:
        return "sin fecha en el registro"
    if a[:10] == b[:10]:
        return "el " + fecha_de(a)
    if a[:7] == b[:7]:
        return "entre el %d y el %s" % (int(a[8:10]), fecha_de(b))
    return "entre el %s y el %s" % (fecha_de(a), fecha_de(b))


def informes_ruso_por_metadatos():
    """Informes cuyos metadatos declaraban ingles y cuyo cuerpo esta en ruso.

    Salen del registro de la etapa 3, que es solo-anexar: cuenta la ultima
    fila de cada registro. Devuelve (total, vistos en el PDF, vistos solo en
    la web del editor).
    """
    ult = {}
    for r in leer_csv(RS / "cribado" / "screening_stage3_pool_decisions.csv"):
        ult[r["record_id"]] = r
    hit = [r["reason"].lower() for r in ult.values()
           if r["verdict"] == "EXCLUDE" and "metadatos declaraban ingl" in r["reason"].lower()]
    return (len(hit), sum(1 for t in hit if "del pdf" in t),
            sum(1 for t in hit if "web del editor" in t))


def cierre_adjudicacion():
    """Fecha de la ultima fila cerrada del registro de desacuerdos."""
    fs = [r["fecha"] for r in leer_csv(RS / "extraccion" / "extraction_conflicts.csv")
          if r.get("resuelto_por", "").strip() and r.get("fecha", "").strip()]
    return fecha_de(max(fs)) if fs else "(sin fecha en el registro)"


def estado_rob():
    """El estado del riesgo de sesgo que escribe build_rob_table.py."""
    p = QR / "rob_tabla_estado.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def firmas_rob():
    """Juicios de S18: cuantos hay, cuantos llevan dos firmas, primera y ultima fecha."""
    filas = leer_csv(OUT / "S18_cada_juicio_con_su_frase.csv")
    fs = sorted(r["fecha"] for r in filas if r.get("fecha", "").strip())
    dos = sum(1 for r in filas if " y " in r.get("firmado_por", ""))
    return len(filas), dos, (fs[0] if fs else ""), (fs[-1] if fs else "")


# Una cita que empieza por un campo de la ficha entre corchetes
# («[eligibilityCriteria]», «[briefSummary, donde se nombra]») sale de la ficha
# del registro de ensayos, no de un articulo.
CITA_FICHA = re.compile(r"^\[[a-z][A-Za-z]*(?:,[^\]]*)?\]")
# Una cita que empieza por el nombre de una base bibliografica y dos puntos
# («SciELO: «Tipo de literatura: Editorial»») sale de como esa base indexa el
# informe, no de su texto.
CITA_INDEXACION = re.compile(r"^(SciELO|PubMed|Scopus|LILACS|BVS|Europe PMC|"
                             r"OpenAlex|Semantic Scholar)\s*:", re.I)


def resumen_exclusiones():
    """Las exclusiones tras el cribado: cuantas, quien las firmo y que citan.

    Una fila sin cita es la que trae la celda vacia o una nota entre parentesis
    («(sin PDF del articulo completo)»): no hay frase que citar. De las que si
    la llevan se cuentan las que citan la ficha del registro (CITA_FICHA) y las
    que citan la indexacion en una base (CITA_INDEXACION).
    """
    filas = leer_csv(RS / "cribado" / "exclusiones_tras_texto_completo.csv")
    una = [r for r in filas if " y " not in r["decidido_por"]]

    def cita(r):
        return r["cita_del_texto"].strip()
    sin = [r for r in filas if not cita(r) or cita(r).startswith("(")]
    con = [r for r in filas if cita(r) and not cita(r).startswith("(")]
    sin_por_codigo, total_por_codigo = {}, {}
    for r in sin:
        sin_por_codigo[r["codigo"]] = sin_por_codigo.get(r["codigo"], 0) + 1
    for r in filas:
        total_por_codigo[r["codigo"]] = total_por_codigo.get(r["codigo"], 0) + 1
    indexa = [r for r in con if CITA_INDEXACION.match(cita(r))]
    por_quien = {}
    for r in una:
        por_quien[r["decidido_por"]] = por_quien.get(r["decidido_por"], 0) + 1
    return {"n": len(filas), "dos": len(filas) - len(una), "una": len(una),
            "una_quien": sorted(por_quien), "una_por_quien": por_quien,
            "sin_cita": len(sin), "con_cita": len(con),
            "sin_por_codigo": sin_por_codigo, "total_por_codigo": total_por_codigo,
            "cita_ficha": sum(1 for r in con if CITA_FICHA.match(cita(r))),
            "cita_indexacion": sorted(r["study_id"] for r in indexa),
            "bases_indexacion": sorted({CITA_INDEXACION.match(cita(r)).group(1)
                                        for r in indexa})}


def enumera(partes):
    """[a, b, c] -> «a, b y c»."""
    partes = [p for p in partes if p]
    if len(partes) < 2:
        return "".join(partes)
    return ", ".join(partes[:-1]) + " y " + partes[-1]


def frase_citas_s16(ex, S):
    """Cuantas exclusiones de S16 no llevan cita, y que citan las que si.

    El indice decia «71 no llevan cita del articulo, entre ellos los 63 ... y
    los 4 ...». Sin cita del articulo hay mas: de las que llevan cita, las de
    las fichas del registro (14-sep) citan la ficha y EST-205 cita la
    indexacion de SciELO. Todo se cuenta del registro de exclusiones.
    """
    NOMBRE = {"NOPDF": "sin el artículo completo (NOPDF)",
              "NOREC": "sin texto que se pudiera leer (NOREC)"}
    orden = sorted(ex["sin_por_codigo"].items(),
                   key=lambda kv: (kv[0] not in NOMBRE, kv[0] != "NOPDF", -kv[1], kv[0]))
    partes = []
    for c, n in orden:
        todos = n == ex["total_por_codigo"].get(c)
        que = NOMBRE.get(c, "con el código %s" % c)
        partes.append(("los %d %s" % (n, que)) if todos else
                      ("%d de los %d %s" % (n, ex["total_por_codigo"][c], que)))
    texto = "%d no llevan cita: %s." % (ex["sin_cita"], enumera(partes))
    citan = []
    if ex["cita_ficha"]:
        citan.append("%d citan la ficha del registro" % ex["cita_ficha"])
    if ex["cita_indexacion"]:
        citan.append("%d (%s) la indexación en %s" % (
            len(ex["cita_indexacion"]), ", ".join(ex["cita_indexacion"]),
            " y ".join(ex["bases_indexacion"])))
    if citan:
        texto += " De las %d que la llevan, %s." % (ex["con_cita"], enumera(citan))
    # Los mismos cortes que build_auditoria_scalars (s16_*): si discrepan, uno
    # de los dos guiones cuenta mal y la frase no debe salir sin aviso.
    otros = sum(n for c, n in ex["sin_por_codigo"].items() if c not in NOMBRE)
    for clave, valor in (("s16_con_cita", ex["con_cita"]),
                         ("s16_sin_cita_otros", otros)):
        if clave in S and S[clave] != valor:
            print("  AVISO: %s = %s en los escalares y %s contado aquí"
                  % (clave, S[clave], valor))
    # Las demas descripciones del indice no cierran con punto.
    return texto.rstrip(".")


def iniciales(nombre):
    """«D. Valdiviezo» o «DANNY JAVIER VALDIVIEZO VERDUGO» -> «D.V.»."""
    n = nombre.lower()
    if "valdiviezo" in n:
        return "D.V."
    if "trelles" in n:
        return "N.T."
    return nombre.strip()


def quien_firmo(nombre):
    """La celda «decidido_por» en palabras: los dos autores o uno solo."""
    return "los dos autores" if " y " in nombre else "%s solo" % iniciales(nombre)


def firmas_criterio_texto():
    """Quien firmo las exclusiones por falta de texto (NOREC, NOPDF) y cuando.

    S0 decia «Los dos autores firmaron el criterio el 6 de octubre de 2026»,
    y el primer NOREC (2 de septiembre de 2026) lo firmo D.V. solo. Se lee del
    registro de exclusiones: (texto de las firmas, fecha de NOPDF).
    """
    grupos = {}
    for r in leer_csv(RS / "cribado" / "exclusiones_tras_texto_completo.csv"):
        if r["codigo"] in ("NOREC", "NOPDF"):
            clave = (r["fecha"][:10], r["decidido_por"].strip())
            grupos.setdefault(clave, {}).setdefault(r["codigo"], 0)
            grupos[clave][r["codigo"]] += 1
    partes = []
    for (fecha, quien), cuenta in sorted(grupos.items(), reverse=True):
        partes.append("%s, el %s (%s)" % (
            quien_firmo(quien), fecha_de(fecha),
            " y ".join("%d %s" % (n, c) for c, n in sorted(cuenta.items()))))
    f_nopdf = sorted(f for (f, _), c in grupos.items() if "NOPDF" in c)
    return "; ".join(partes), (fecha_de(f_nopdf[-1]) if f_nopdf else "")


def validacion_rayyan():
    """El recribado ciego en Rayyan: quien lo hizo y que paso con lo que incluyo.

    El indice decia «0 falsos negativos en 350» sin decir que lo hizo un solo
    revisor ni que en la primera pasada marco 6 como incluibles y los
    reclasifico despues. Lo dice resultado.json.
    """
    p = RS / "validacion_rayyan" / "resultado.json"
    j = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}
    adj = j.get("adjudicacion", {})
    return {"revisor": j.get("revisor", ""),
            "quien_reclasifica": adj.get("quien", ""),
            "reclasificados": adj.get("reclasificados_a_excluido", 0)}


def primera_pasada():
    """(fecha, «, salvo ...», franja) de la primera pasada; («», «», «») si no consta.

    El 4-ago por la manana el modelo cribo el brazo A de PubMed por titulo
    (screening_stage2_decisions.csv) y por resumen (screening_stage3_decisions.csv).
    El cribado del pozo la sustituyo, y S3 lleva el vigente, con las decisiones
    de resumen que el pozo heredo por PMID de esa pasada. El indice decia que
    S3 trae «cada decision de cribado por titulo», y esa pasada no esta en S3.
    La franja («por la manana») y las heredadas salen de los ficheros: solo se
    cuentan como heredadas de esa pasada las filas cuyo veredicto y motivo estan
    en screening_stage3_decisions.csv.
    """
    t2 = [r["decided_at"] for r in leer_csv(
        RS / "cribado" / "screening_stage2_decisions.csv")
        if r.get("decided_at", "").strip()]
    if not t2:
        return "", "", ""
    s3 = leer_csv(RS / "cribado" / "screening_stage3_decisions.csv")
    horas = [int(t[11:13]) for t in t2 + [r["decided_at"] for r in s3
                                           if r.get("decided_at", "").strip()]]
    franja = ("por la mañana" if max(horas) < 12 else
              "por la tarde" if min(horas) >= 12 and max(horas) < 21 else "")
    de_la_pasada = {(r["verdict"], r["reason"].strip()) for r in s3}
    heredadas = [r for r in leer_csv(
        RS / "cribado" / "screening_stage3_pool_decisions.csv")
        if r.get("decided_by", "").startswith("heredada del brazo")]
    n_her = sum(1 for r in heredadas
                if (r["verdict"], r["reason"].strip()) in de_la_pasada)
    if n_her != len(heredadas):
        print("  AVISO: %d filas «heredadas del brazo» y solo %d coinciden con la "
              "primera pasada de resumen" % (len(heredadas), n_her))
    return (fecha_de(min(t2)),
            (", salvo %d decisiones de resumen que el cribado del pozo heredó "
             "de ella por PMID" % n_her) if n_her else "",
            franja)


_S10 = []


def s10_en_disco(S):
    """Que trae S10 en el paquete: [(nombre, descripcion, es la del cuerpo)].

    S10 era una copia sin rehacer de idioma_texto_completo.csv (13-ago), anterior
    a la comprobacion sobre el cuerpo de los 65 estudios del corpus
    (quality_reports/idioma_cuerpo_pdf.csv, 6-oct), y el indice la describia
    como si fuera esa comprobacion. build_verifiables_package.py la rehace desde
    idioma_cuerpo_pdf.csv. La descripcion no se fija aqui: se lee del fichero
    que haya en disco (si trae la columna del idioma del cuerpo, es la
    comprobacion nueva), y si es el viejo se dice y se avisa. Se calcula una
    vez: lo usan S0 y el indice, y los avisos no deben salir dos veces.
    """
    if _S10:
        return _S10
    cuerpo_ids = {r["study_id"] for r in leer_csv(QR / "idioma_cuerpo_pdf.csv")}
    n_corpus = S["idioma_cuerpo_comprobados"]
    # «S10*» y no «S10_*»: si el fichero viejo se conserva rotulado como
    # historico (S10b_...), el indice tiene que listarlo y decir que es.
    # Es la comprobacion sobre el cuerpo si trae una columna del idioma del
    # cuerpo o si sus estudios son, uno a uno, los de idioma_cuerpo_pdf.csv.
    leidos = []
    for p in sorted(OUT.glob("S10*.csv")):
        datos = leer_csv(p)
        cols = list(datos[0].keys()) if datos else []
        ests = {v.strip() for r in datos for v in r.values()
                if re.fullmatch(r"EST-\d+", (v or "").strip())}
        es = (any("cuerpo" in c.lower() for c in cols)
              or (bool(cuerpo_ids) and ests == cuerpo_ids
                  and len(datos) == len(cuerpo_ids)))
        leidos.append((p, datos, cols, es))
    nuevo = [p.stem for p, _, _, es in leidos if es]
    filas = []
    for p, datos, cols, es in leidos:
        nombre = p.stem + ".xlsx + .csv"
        if es:
            if len(datos) == n_corpus:
                otros = [("%d en español" % S["idioma_cuerpo_espanol"]
                          if S["idioma_cuerpo_espanol"] else ""),
                         ("%d en otro idioma" % S["idioma_cuerpo_otro"]
                          if S["idioma_cuerpo_otro"] else "")]
                reparto = (("los %d en inglés" % S["idioma_cuerpo_ingles"])
                           if S["idioma_cuerpo_ingles"] == n_corpus else
                           enumera(["%d en inglés" % S["idioma_cuerpo_ingles"]] + otros))
                # Como S1: el guion lo escribio el modelo, y se dice.
                desc = ("Idioma del cuerpo del PDF de cada uno de los %d estudios "
                        "del corpus, comprobado por un guion del canal escrito "
                        "por el modelo sobre el texto del artículo y no sobre "
                        "los metadatos: %s"
                        % (n_corpus, reparto))
            else:
                print("  AVISO: %s trae %d filas y el corpus tiene %d estudios"
                      % (p.name, len(datos), n_corpus))
                desc = ("Idioma del cuerpo del PDF, comprobado sobre el texto "
                        "del artículo: %d filas, frente a %d estudios en el corpus"
                        % (len(datos), n_corpus))
        else:
            ids = {(r.get(cols[0]) or "").strip() for r in datos} if cols else set()
            if not nuevo:
                print("  AVISO: %s no es la comprobación sobre el cuerpo de los "
                      "%d estudios del corpus (quality_reports/idioma_cuerpo_pdf.csv)"
                      % (p.name, n_corpus))
            desc = ("Comprobación del idioma sobre el texto completo anterior a "
                    "la comprobación sobre el cuerpo de los %d estudios del "
                    "corpus: %d filas, que cubren %d de esos %d estudios. La "
                    "comprobación sobre el cuerpo está %s"
                    % (n_corpus, len(datos), len(ids & cuerpo_ids), n_corpus,
                       ("en " + ", ".join(nuevo)) if nuevo else
                       "en el repositorio y no en este anexo"))
        filas.append((nombre, desc, es))
    if not filas:
        print("  AVISO: no hay S10 en el paquete")
        filas.append(("S10_idioma_verificado_sobre_texto_completo.xlsx + .csv",
                       "No está en el paquete", False))
    _S10.extend(filas)
    return _S10


def rango(nums):
    """[1, 2] -> «1 y 2»; [1..7] -> «1 a 7»; si hay huecos, la lista."""
    if not nums:
        return ""
    if len(nums) == 1:
        return str(nums[0])
    if nums == list(range(nums[0], nums[-1] + 1)):
        return ("%d y %d" if len(nums) == 2 else "%d a %d") % (nums[0], nums[-1])
    return ", ".join(str(n) for n in nums[:-1]) + " y %d" % nums[-1]


def numerados(carpeta, patron):
    """Los numeros de los ficheros «tabla_3_x.csv» o «figura_1_x.pdf»."""
    return sorted(int(m.group(1)) for p in carpeta.glob("*")
                  for m in [re.match(patron, p.name)] if m)


def largo_titulo(p):
    """Caracteres del titulo de un manuscrito, sin la marca de cursiva."""
    t = pathlib.Path(p).read_text(encoding="utf-8").split("\n", 1)[0]
    return len(re.sub(r"\*", "", t.lstrip("# ").strip()))


def seccion_del_maestro(desde, hasta):
    """El texto del maestro entre dos encabezados, para comprobar que dice algo."""
    t = (ROOT / "paper" / "manuscrito_revision_sistematica.md").read_text(encoding="utf-8")
    i = t.find(desde)
    j = t.find(hasta, i + 1) if i >= 0 else -1
    return t[i:j] if i >= 0 and j > i else ""


def informe_editorial(S):
    # QUIEN LO ESCRIBE, A LA VISTA. Se titulaba «Informe editorial», con el
    # subtitulo «Evaluacion previa a revision por pares. Revista de
    # microbiologia clinica / revisiones sistematicas», que es la cabecera de
    # un informe de revista, y hablaba en primera persona de editor. Lo redacto
    # el modelo y viaja en el sobre: ahora lo dicen el titulo, la primera frase
    # y la cabecera de cada pagina. El nombre del fichero no se toca aqui: lo
    # citan la guia, la lista de anexos y el sobre, que son de otros guiones.
    AVISO = ("Autoevaluación redactada por el modelo de lenguaje (Claude Opus 5) "
             "a petición de los autores. No es un informe de un editor ni una "
             "revisión externa.")
    d = doc_nuevo("Autoevaluación del manuscrito, redactada por el modelo de "
                  "lenguaje")
    for sec in d.sections:
        r = sec.header.paragraphs[0].add_run(AVISO)
        r.italic = True
        r.font.size = Pt(8)
    d.add_paragraph().add_run(AVISO).bold = True
    p = d.add_paragraph(
        # La fecha sale de fecha_larga(): se regenera con el paquete.
        "Estado a %s. Se escribió adoptando el punto de vista de un editor que "
        "decide si el manuscrito puede entrar en revisión por pares, pero las "
        "valoraciones son del modelo, no de una revista. El modelo hizo buena "
        "parte del trabajo que aquí se evalúa —la declaración de uso de "
        "inteligencia artificial del manuscrito dice cuál—, de modo que esta "
        "autoevaluación no es independiente." % fecha_larga())
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9.5)
    # S0 remite a «§2.9 y Tabla 3» y viaja tambien en el sobre de la revista,
    # cuyo manuscrito no numera secciones y cuya Tabla 3 es otra (la de
    # completitud). Sin este aviso las remisiones llevaban a otra parte.
    p = d.add_paragraph(
        "Las remisiones a secciones (§) y a tablas son las del manuscrito "
        "maestro (el informe extendido). El manuscrito preparado para la "
        "revista no numera las secciones y numera las tablas de otro modo; S1 "
        "da la equivalencia de las secciones.")
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9.5)
    f_pp = primera_pasada()
    s10_cuerpo = any(es for _, _, es in s10_en_disco(S))

    d.add_heading("Valoración del modelo", level=1)
    p = d.add_paragraph()
    r = p.add_run("REVISIÓN MAYOR ANTES DE ENVIAR A PARES.")
    r.bold = True
    d.add_paragraph(
        "El trabajo metodológico que sostiene este manuscrito está por encima "
        "de lo habitual en el campo, y su hallazgo central es relevante y "
        "publicable. " +
        # Dos redacciones, segun el estado. Con desacuerdos abiertos el reparo
        # es la adjudicacion. Sin ellos decia «No es un consenso de los dos: el
        # registro anota ... a uno solo en 525» y que lo unico que seguia
        # bloqueando era el registro del protocolo. Lo primero va mas alla de lo
        # comprobado (el registro anota un unico responsable; no consta que no
        # hubiera consenso) y lo segundo callaba por que el veredicto sigue en
        # pie: el cribado por el modelo, la falta de independencia y las
        # enmiendas, que no se arreglan cerrando tareas.
        ("No puede, sin embargo, entrar en revisión por pares como revisión "
         "sistemática completa mientras los desacuerdos de la extracción no estén "
         "adjudicados: la extracción por duplicado sí está hecha (%d de %d "
         "estudios). Quedan %d de los %d desacuerdos sin firmar, y las cifras de "
         "desenlace se apoyan en la extracción adjudicada: ese conjunto tiene "
         "todavía casillas abiertas, que salen vacías y se declaran como tales."
         % (S["extraccion_estudios_ambos"], S["extraccion_estudios_r1"],
            S["extraccion_conflictos_sin_firmar"], S["extraccion_desacuerdos"])
         if S["extraccion_conflictos_sin_firmar"] else
         "La extracción por duplicado está hecha (%d de %d estudios) y su "
         "adjudicación se cerró el %s: %d de los %d desacuerdos están "
         "adjudicados; el registro anota la firma de los dos en %d y un único "
         "responsable en %d, y solo esas %d constan como consenso. Los %d "
         "restantes se cerraron por una regla mecánica, sin firma conjunta, "
         "declarada como tal. El veredicto se mantiene por lo que no se arregla "
         "cerrando tareas: las decisiones del cribado de títulos y resúmenes las "
         "emitió el modelo como revisor único, sin duplicación, con un "
         "vocabulario de motivos que no se fijó antes de empezar; ni la "
         "extracción ni el riesgo de sesgo acreditan la independencia de sus dos "
         "lecturas; las enmiendas de elegibilidad se adoptaron con el cribado "
         "terminado, y el protocolo no se registró. El manuscrito lo declara "
         "todo."
         % (S["extraccion_estudios_ambos"], S["extraccion_estudios_r1"],
            cierre_adjudicacion(),
            S["extraccion_conflictos_firmados"], S["extraccion_desacuerdos"],
            S["conflictos_resueltos_por_los_dos"],
            S["conflictos_resueltos_por_uno"],
            S["conflictos_resueltos_por_los_dos"],
            S["extraccion_cerrados_por_regla"])) +
        " Lo que sigue distingue lo que impide publicar de lo que solo mejora "
        "el manuscrito.")

    n_ruso, n_ruso_pdf, n_ruso_web = informes_ruso_por_metadatos()
    d.add_heading("Lo que el manuscrito hace bien, a juicio del modelo", level=1)
    for t in [
        # Decia «mas amplia que la de las revisiones publicadas» y que BVS y
        # SciELO aportan «la literatura donde se concentran los disenos
        # comparativos». El manuscrito no sostiene ni lo uno ni lo otro: los
        # comparativos se concentran en la literatura de Europa del Este (§4.2),
        # que la enmienda de idioma dejo fuera.
        "La búsqueda es amplia: %d fuentes en dos corrientes —Scopus en dos "
        "brazos, %d consultas en total—, sin restricción de idioma, con "
        "literatura regional (BVS y SciELO) y registros de ensayos. Sin "
        "restricción de idioma recuperó %d estudios rusos y %d ucranianos, la "
        "literatura donde se concentran los diseños comparativos; la enmienda de "
        "idioma la dejó fuera, y el manuscrito lo declara (§4.2)."
        % (S["fuentes_distintas_n"], S["fuentes_brazos_n"],
           S["rusos_antes_de_la_enmienda"], S["ucranianos_antes_de_la_enmienda"]),
        # Decia que la auditoria «permite afirmar, y no solo esperar, que el
        # cribado no perdio estudios elegibles conocidos». Solo se comprobaron
        # los 40 con PMID de los 42 fijados, el conjunto salio de un metaanalisis
        # hecho con modelos y sirvio para disenar las ecuaciones, y Ronit 2024
        # se perdio con la enmienda de idioma.
        "La auditoría de controles positivos es infrecuente. Antes del cribado "
        "se fijaron %d estudios conocidos como elegibles, procedentes del "
        "metaanálisis exploratorio de julio de 2026, que se hizo con asistencia "
        "de modelos de lenguaje; la comprobación automática alcanzaba a los %d "
        "con PMID, y el cribado no perdió ninguno de esos %d hasta la enmienda "
        "de idioma, que excluyó en la etapa de resumen a Ronit 2024. Las "
        "ecuaciones de búsqueda se diseñaron para recuperar ese conjunto, así "
        "que la auditoría acota lo que pierde la selección, no la sensibilidad "
        "de la búsqueda. Hoy siguen en el corpus %d de los %d (§2.5 dice cuáles "
        "salieron y por qué). Que se declare la regla de exclusión mal calibrada "
        "que la auditoría detectó en una ejecución intermedia la hace más "
        "creíble, no menos."
        % (S["controles_fijados"], S["controles_comprobables_pmid"],
           S["controles_comprobables_pmid"], S["controles_en_corpus"],
           S["controles_fijados"]),
        # Decia «El vocabulario cerrado de exclusion, fijado antes de empezar».
        # El registro lo desmiente: el primer dia de cribado de titulos los
        # motivos van en texto libre y los codigos se usaron desde el dia
        # siguiente. Lo que merece elogio es que el manuscrito ya lo declare.
        # Decia «las exclusiones del primer dia se asignaron despues a un
        # codigo» y «Ese mismo 4 de agosto se anadieron, a mitad de lote, las
        # reglas». El primer dia hubo, por la manana, una primera pasada del
        # modelo sobre el brazo A de PubMed, que el cribado del pozo sustituyo
        # y que no esta en S3: el codigo se asigno a las exclusiones del primer
        # dia del cribado vigente. Las dos reglas entraron en esa pasada, en
        # los guiones de lote del modelo, que las propuso; las decidieron los
        # autores (hoja de hechos, 13-oct). La pasiva callaba quien las metio.
        "El cribado acabó usando un vocabulario cerrado de %d motivos, y el "
        "manuscrito declara que no se fijó antes de empezar (§2.4): las "
        "decisiones del primer día de cribado de títulos, el 4 de agosto de "
        "2026, llevan el motivo en texto libre; los códigos, cuya redacción "
        "procede de esos textos, se usaron desde el 5 de agosto, y las "
        "exclusiones del primer día que siguen en el cribado vigente se "
        "asignaron después a un código. "
        % len(S["exclusiones_titulo"]) +
        (("El %s%s, el modelo había hecho una primera pasada sobre "
          "el brazo A de PubMed, que el cribado del pozo sustituyó y que consta "
          "en el repositorio y no en S3%s. En esa pasada entraron, a mitad de "
          "lote, "
          % (f_pp[0], (", " + f_pp[2]) if f_pp[2] else "", f_pp[1]))
         if f_pp[0] else
         "Ese mismo 4 de agosto entraron, a mitad de lote, ") +
        "las reglas que dejan fuera las endolisinas aisladas, por la "
        "intervención, y los informes veterinarios, por la población, que "
        "propuso el modelo y metió en sus guiones de lote, y que decidieron los "
        "autores. Declararlo así es lo correcto; presentarlo como un vocabulario "
        "fijado de antemano no lo sería.",
        # Era «La seccion 3.5»: al entrar la completitud por desenlace como
        # 3.5, la viabilidad de la sintesis paso a la 3.6.
        "La sección 3.6 argumenta explícitamente por qué NO se agrupa. Es lo "
        "contrario de lo habitual, y es lo correcto cuando los requisitos del "
        "diseño no se cumplen.",
        # Decia que la declaracion «separa las tareas en que el modelo juzgo de
        # aquellas en que no» y que «ningun revisor humano adjudico decisiones
        # de cribado individuales». La declaracion canonica ya no dice ni lo uno
        # ni lo otro, y no se resume con palabras propias: se remite a ella.
        "La declaración de uso de inteligencia artificial está en la sección "
        "Declaraciones del manuscrito y se copia literal en S7. Nombra el modelo "
        "y su versión y recorre, fase por fase, lo que hizo. Esta autoevaluación "
        "no la resume: remite a ella. Entre lo que declara está la redacción de "
        "los borradores de los anexos, y este documento es uno de ellos.",
        # Decia «en los 47 estudios con PDF» y que la comprobacion «detecto dos
        # estudios» en ruso. Hoy los 65 del corpus tienen PDF y se comprobaron
        # sobre el cuerpo, y el registro de la etapa 3 anota tres informes, no
        # dos: el tercero solo se vio en la web del editor. Se cuentan desde el
        # registro (informes_ruso_por_metadatos). Esas exclusiones las firma en
        # el registro el modelo (decided_by), y el texto lo dice.
        # Remitia a «S9 y S10» y S10 era una copia del 13-ago, anterior a la
        # comprobacion sobre el cuerpo: ahora se remite a S10 solo si S10 es
        # esa comprobacion (s10_en_disco), y la clase de evidencia, a S9.
        "El criterio de idioma no se dio por supuesto. Se tomó del campo de "
        "idioma de la fuente y, donde había texto completo, se comprobó sobre "
        "él; esa comprobación, que hizo el modelo, destapó %d informes cuyos "
        "metadatos declaraban inglés y "
        "cuyo cuerpo está en ruso, que se excluyeron: %d se vieron en el PDF y "
        "%d solo en la web del editor. El 6 de octubre de 2026 se repitió sobre "
        "el cuerpo de los %d estudios del corpus vigente, y los %d están en "
        "inglés%s. Los %d informes del corpus quedan con su clase de evidencia y "
        "ninguno sin prueba (S9)."
        % (n_ruso, n_ruso_pdf, n_ruso_web, S["idioma_cuerpo_comprobados"],
           S["idioma_cuerpo_ingles"],
           (" (S10)" if s10_cuerpo else
            "; esa comprobación consta en el repositorio y no en S10, que es "
            "anterior"),
           S["informes_agrupados"]),
    ]:
        d.add_paragraph(t, style="List Bullet")

    d.add_heading("Sobre la restricción de idioma", level=1)
    d.add_paragraph(
        # Decia «Como editor lo acepto declarado, pero advierto»: primera
        # persona de un editor que no existe. Redaccion impersonal.
        "La revisión restringe la elegibilidad a informes en inglés o español. "
        "El criterio se adoptó con el cribado ya cerrado y está declarado como "
        "enmienda en §2.9, con su impacto medido. Declarado así es aceptable, "
        "pero cabe esperar la objeción que se le va a plantear, y el manuscrito "
        "debe seguir respondiéndola con claridad: el propio artículo documenta "
        "que los diseños comparativos se concentran en la literatura de Europa "
        "del Este, y el criterio deja fuera precisamente esa literatura. La "
        "respuesta actual —que la búsqueda sí la recuperó, que la revisión "
        "documenta su existencia y su tamaño, y que declara no poder "
        "pronunciarse sobre su contenido— es defendible. Debe mantenerse en "
        "Discusión y no diluirse.")
    tabla(d, ["Efecto del criterio", "Antes de la enmienda",
              "Después de la enmienda", "Corpus vigente"],
          # Todas del canal. Estaban escritas a mano y envejecieron: el informe
          # llegó a decir 185 estudios donde el manuscrito decía 184, y así en
          # casi todas. Un anexo que contradice al manuscrito lo hunde.
          # La columna «Despues» era el corpus vigente, y asi atribuia al idioma
          # lo que salio despues al releer los textos y con NOPDF (los
          # comparativos pasaban de 41 a 2, cuando el idioma se llevo 18). Cada
          # perdida contra su propio antes, como en build_synthesis_scalars: la
          # columna del medio es *_antes_de_releer. La fila de procedencia rusa
          # no tiene escalar despues de la enmienda y pasa a la nota.
          [("Estudios", S["estudios_antes_de_la_enmienda"],
            S["estudios_antes_de_releer"], S["estudios"]),
           ("Con publicación recuperable", S["extraibles_antes_de_la_enmienda"],
            S["extraibles_antes_de_releer"], S["estudios_extraibles"]),
           ("Diseños comparativos", S["comparativos_antes_de_la_enmienda"],
            S["comparativos_antes_de_releer"], S["estudios_comparativos"]),
           ("Ensayos aleatorizados", S["ecas_antes_de_la_enmienda"],
            S["ecas_antes_de_releer"], S["ecas"])],
          [5.4, 3.7, 3.7, 3.7])
    nota(d, "«Después de la enmienda» es el corpus el día en que se adoptó. El "
            "corpus vigente recoge además lo que salió al releer los textos "
            "completos y con la segunda enmienda (solo estudios con el artículo "
            "completo), que no es efecto del idioma. La procedencia rusa pasó de "
            "%d estudios antes de la enmienda a %d en el corpus vigente, sobre "
            "todo por este criterio. De los %d informes que excluyó, %d se "
            "detectaron solo al comprobar el texto completo: sus metadatos "
            "declaraban inglés y el cuerpo está en ruso. Se pierde además Ronit "
            "2024, del conjunto de control positivo."
            % (S["rusos_antes_de_la_enmienda"], S["procedencia"].get("Rusia", 0),
               S["enmienda_idioma_informes_excluidos"], n_ruso))

    R = estado_rob()
    rob_hecho = bool(R.get("completa")) and not R.get("celdas_pendientes")
    n18, n18_dos, f18_ini, f18_fin = firmas_rob()
    sin_texto = S["texto_completo_no_obtenido"]
    firmas_texto, f_nopdf = firmas_criterio_texto()
    # «(bloqueantes)» a secas, con B1, B3 y B4 ya resueltos en la tabla.
    d.add_heading("Condiciones para publicar: las que bloquean y las ya "
                  "resueltas", level=1)
    tabla(d,
          ["#", "Qué falta", "Por qué bloquea", "Quién lo resuelve"],
          # B1 daba por pendiente la extraccion cuando esta hecha al 98 %. Lo
          # que bloquea hoy es la adjudicacion de los desacuerdos, que es otra
          # cosa y la resuelven otras personas. Las cifras salen de los
          # escalares para que esta fila no se vuelva a quedar vieja.
          # Se titulaba «Adjudicacion por consenso» y daba como responsables
          # «Los dos revisores»: el registro anota a los dos solo en 14 y un
          # unico responsable en 525. Solo esas 14 constan como consenso.
          [("B1", "Adjudicación de los desacuerdos de la extracción",
            ("La extracción por duplicado está completa (%d de %d estudios, "
             "%s %%), pero %d de los %d desacuerdos siguen sin firmar, y las "
             "cifras de desenlace se apoyan en la extracción adjudicada."
             % (S["extraccion_estudios_ambos"], S["extraccion_estudios_r1"],
                S["extraccion_doble_pct"], S["extraccion_conflictos_sin_firmar"],
                S["extraccion_desacuerdos"])
             if S["extraccion_conflictos_sin_firmar"] else
             "RESUELTO el %s. Extracción por duplicado completa (%d de %d "
             "estudios, %s %%; los %d del corpus vigente la tienen) y %d de %d "
             "desacuerdos adjudicados: el registro anota la firma de los dos en "
             "%d y un único responsable en %d (D.V. en %d, N.T. en %d); solo "
             "esas %d constan como consenso. Los %d restantes se cerraron por "
             "una regla mecánica, sin firma conjunta, y se declaran como tal. "
             "Sin independencia acreditada: la columna de texto libre «¿Dónde "
             "viste estos datos?» coincide carácter por carácter entre los dos "
             "cuadernos en %d de %d filas. La fila se conserva porque era el "
             "reparo principal y el lector tiene que poder ver qué cayó y qué no."
             % (cierre_adjudicacion(),
                S["extraccion_estudios_ambos"], S["extraccion_estudios_r1"],
                S["extraccion_doble_pct"], S["extraccion_doble_corpus"],
                S["extraccion_conflictos_firmados"],
                S["extraccion_desacuerdos"], S["conflictos_resueltos_por_los_dos"],
                S["conflictos_resueltos_por_uno"], S["conflictos_resueltos_solo_dv"],
                S["conflictos_resueltos_solo_nt"],
                S["conflictos_resueltos_por_los_dos"],
                S["extraccion_cerrados_por_regla"],
                S["extraccion_localizacion_identicas"],
                S["extraccion_localizacion_comparables"])),
            ("Los dos revisores, sobre ADJUDICACION_conflictos.xlsx"
             if S["extraccion_conflictos_sin_firmar"] else
             "Adjudicado según el registro: D.V. (%d), N.T. (%d), los dos (%d); "
             "%d por regla" % (S["conflictos_resueltos_solo_dv"],
                               S["conflictos_resueltos_solo_nt"],
                               S["conflictos_resueltos_por_los_dos"],
                               S["extraccion_cerrados_por_regla"]))),
           # Decia «PROSPERO admite registro aunque la revision este en marcha»
           # y recomendaba registrar en 1-2 dias. PROSPERO no admite una
           # revision con la extraccion hecha (CLAUDE.md, estado a 2026-10-07),
           # y la de esta lo esta: la ausencia ya no tiene arreglo, solo se
           # declara, como hace §2.1.
           ("B2", "Registro del protocolo",
            "La revisión no se registró ni depositó un protocolo previo, y el "
            "manuscrito lo declara (§2.1), que es la vía honesta. PROSPERO ya "
            "no es una opción: no admite revisiones con la extracción hecha. "
            "Muchas revistas devuelven sin revisar un manuscrito sin registro; "
            "aquí la ausencia no tiene arreglo y solo cabe declararla.",
            "Sin arreglo retroactivo; declarado"),
           # El 46,5 % y el 63,4 % eran de una version anterior del corpus.
           # Desde NOPDF (2026-10-06) la plantilla pendiente salia con «El
           # 100,0 % obtenido sigue siendo insuficiente» y «concentra el 0,0 %».
           # Sin textos que falten, la fila se da por resuelta y dice como: por
           # construccion, excluyendo lo que no se obtuvo.
           ("B3", "Recuperación del texto completo",
            ("Faltan %d de los %d textos completos (%s %% obtenido), y la "
             "fracción que falta no es aleatoria: concentra el %s %% de los "
             "diseños comparativos. Un revisor objetará que las características "
             "descritas no representan al campo."
             % (sin_texto, S["estudios_extraibles"], coma(S["texto_completo_pct"]),
                coma(S["comparativos_sin_texto_pct"]))
             if sin_texto else
             # «Codigo NOPDF» a secas no bastaba: las cifras del criterio
             # cuentan tambien los NOREC, y la fecha sale del registro.
             "RESUELTO POR CONSTRUCCIÓN, NO POR RECUPERACIÓN. Desde el %s "
             "solo entran estudios con el PDF del artículo completo (segunda "
             "enmienda, códigos NOREC y NOPDF), así que la recuperación es del "
             "%s %% (%d de %d) por definición: lo que faltaba no se recuperó, "
             "se excluyó. Justo antes del criterio el texto completo cubría %d "
             "de %d estudios con artículo o resumen (%s %%). El sesgo no "
             "desaparece, cambia de sitio: contando las dos fechas, el criterio "
             "se lleva %d de los %d diseños comparativos y %d de los %d ensayos "
             "aleatorizados, y el manuscrito lo declara como efecto del "
             "criterio (§2.9 y Tabla 3)."
             % (f_nopdf, coma(S["texto_completo_pct"]),
                S["texto_completo_obtenido"],
                S["estudios_extraibles"], S["texto_completo_obtenido"],
                S["texto_completo_antes_de_nopdf_base"],
                coma(S["texto_completo_antes_de_nopdf_pct"]),
                S["criterio_texto_comparativos"],
                S["criterio_texto_comparativos_antes"],
                S["criterio_texto_ecas"], S["criterio_texto_ecas_antes"])),
            # Decia «Los dos autores firmaron el criterio el 6 de octubre de
            # 2026», y el primer NOREC, del 2 de septiembre, lo firmo D.V. solo.
            ("Biblioteca de la universidad; petición interbibliotecaria"
             if sin_texto else
             "Firmado según el registro por " + firmas_texto)),
           # Seguia como bloqueante («Depende de B1 y B3») con el riesgo de
           # sesgo ya emitido y firmado. Se da por resuelto si el estado que
           # escribe build_rob_table.py dice que esta completo, y con lo que
           # ese estado y S18 registran, incluida la falta de independencia.
           ("B4", "Evaluación del riesgo de sesgo",
            ("RESUELTO. Evaluado en los %d estudios comparativos evaluables (%d "
             "con RoB 2 y %d con ROBINS-I): %d juicios, %s, firmados %s. Sin "
             "independencia acreditada: hubo "
             "dos cuadernos, uno a nombre de cada autor, pero compartían %s "
             "caracteres de notas y el registro del 5 de septiembre de 2026 "
             "concluye que uno salió del otro. Se declaró evaluación por "
             "consenso, sobre un cuaderno que generó el modelo con los juicios "
             "coincidentes precargados y los discrepantes en blanco. Concordancia "
             "de %d en %d juicios (%s %%; kappa %s), que no mide dos "
             "evaluaciones independientes; %s hacia el cuaderno de D.V. GRADE no "
             "se aplica: no hay estimador agrupado cuya certeza calificar."
             % (R["evaluables"], R["por_instrumento"].get("RoB 2", 0),
                R["por_instrumento"].get("ROBINS-I", 0), R["celdas_totales"],
                ("todos con las dos firmas en S18" if n18 == n18_dos == R["celdas_totales"]
                 else "%d de los %d de S18 con las dos firmas" % (n18_dos, n18)),
                entre_fechas(f18_ini, f18_fin),
                miles(S["rob_notas_identicas_caracteres"]),
                S["rob_acuerdo_bruto"], S["rob_juicios_comparables"],
                coma(S["rob_acuerdo_pct"]), coma(S["rob_kappa"]),
                ("los %d desacuerdos se resolvieron todos" % S["rob_desacuerdos"]
                 if S["rob_resueltos_hacia_r1"] == S["rob_desacuerdos"] else
                 "%d de los %d desacuerdos se resolvieron"
                 % (S["rob_resueltos_hacia_r1"], S["rob_desacuerdos"])))
             if rob_hecho else
             "PRISMA 2020 ítems 11 y 18. Faltan %d de los %d juicios."
             % (R.get("celdas_pendientes", 0), R.get("celdas_totales", 0))),
            ("Los dos autores (firmado)" if rob_hecho else
             "Los dos revisores, con RoB 2 y ROBINS-I")),
           ("B5", "Formularios ICMJE de conflicto de interés",
            "Uno firmado por autor. Requisito administrativo sin excepción.",
            "Cada autor"),
           # Decia «Contribucion de autores en formato CRediT. Exigido por la
           # mayoria de revistas», como si faltara entera. El CRediT esta; lo que
           # falta es la parte de N.T. y su aprobacion escrita de la version
           # final, que ningun documento puede dar por hecha.
           # Decia solo que N.T. no habia aprobado ni revisado la version final
           # y «Sin esa aprobacion», como si fuera lo unico que falta. N.T.
           # tampoco ha aceptado responder por el trabajo (hoja de hechos,
           # 13-oct), que es otro criterio del ICMJE.
           ("B6", "Contribución de N.T. y aprobación de la versión final",
            "El CRediT está en el manuscrito, pero la parte de N.T. lleva un "
            "marcador hasta que conste por escrito. N.T. no ha aprobado ni "
            "revisado por escrito la versión final, ni ha aceptado responder "
            "por el trabajo: D.V. asume la responsabilidad por la exactitud e "
            "integridad del trabajo; la de N.T. queda pendiente de su "
            "aprobación escrita. Sin esos actos de N.T. no se cumplen los "
            "criterios de autoría del ICMJE.",
            "N.T., por escrito")],
          [1.2, 4.6, 6.4, 4.4])

    # La tabla de cuatro ensayos «prioritarios» se retira. Era del corpus de
    # agosto, recomendaba recuperar EST-118 (Leitner), excluido por NOREC el
    # 2-sep, y desde NOPDF no queda texto que recuperar: salia «No hacen falta
    # los 0 textos completos que faltan». Si vuelve a faltar texto, la seccion
    # reaparece con las cifras del canal y sin una lista tecleada.
    if sin_texto:
        d.add_heading("Prioridad dentro de B3, si el tiempo es limitado", level=1)
        d.add_paragraph(
            "No hacen falta los %d textos completos que faltan para que el "
            "manuscrito resista la revisión. Los que deciden son los de diseño "
            "comparativo y, dentro de ellos, los ensayos aleatorizados: "
            "recuperarlos cambia la objeción de «no ha leído los ensayos "
            "principales» a «no ha leído parte de la literatura regional», que "
            "es una limitación declarable." % sin_texto)

    d.add_heading("Observaciones que mejoran el manuscrito (no bloquean)", level=1)
    # Cada observacion que depende del estado se comprueba contra el canal y se
    # retira cuando ya no aplica. Salian tres caducadas: «241 palabras ...
    # cumple el limite» con un resumen de 302; «repetir la salvedad en la
    # leyenda de la Tabla 2», que ya la lleva; y «la Figura 1 incluye una caja
    # de recuperacion de texto completo», que no se dibuja desde NOPDF.
    maestro = ROOT / "paper" / "manuscrito_revision_sistematica.md"
    ingles = ROOT / "paper" / "manuscript_systematic_review_en.md"
    t2 = ROOT / "paper" / "tablas" / "tabla_2_completitud.md"
    t2_txt = t2.read_text(encoding="utf-8") if t2.exists() else ""
    salvedad_t2 = ("*Nota.*" in t2_txt
                   and "resumen" in t2_txt.split("*Nota.*", 1)[1].lower())
    caja_fig1 = not (S.get("estudios_solo_registro", 0) == 0 and sin_texto == 0)
    endolisinas_en_discusion = "endolisin" in seccion_del_maestro(
        "## 4. Discusión", "## 5.").lower()
    resumen_max = max(S["palabras_resumen_es"], S["palabras_resumen_en"])
    obs = [
        # Las longitudes intermedias (162/157 y 141/141) iban tecleadas; se
        # quitan, y la final se mide sobre el titulo de cada manuscrito.
        "El título se revisó dos veces el 14 de septiembre de 2026. Partía de "
        "«revisión sistemática de la estructura y la verificabilidad del "
        "cuerpo de evidencia clínica». Primero se acortó adoptando la fórmula "
        "«revisión sistemática de la literatura» que usa la revista de destino. "
        "Después se sustituyó «verificabilidad» por «completitud del reporte» "
        "—«reporting completeness»—, y quedó en %d caracteres en castellano y "
        "%d en inglés. Dos pérdidas "
        % (largo_titulo(maestro), largo_titulo(ingles)) +
        "conscientes, y las dos se declaran aquí. La primera: el título ya no "
        "nombra la estructura, que es la otra mitad de la pregunta de "
        "investigación; esa pregunta no se tocó y sigue formulada sobre las "
        "tres propiedades. La segunda: «completitud del reporte» nombra dos "
        "de los tres componentes medidos y deja fuera la recuperabilidad del "
        "informe, que se sigue nombrando en el cuerpo. Se prefirió un título "
        "que dice menos a uno que dice otra cosa. Sobre el término adoptado "
        "conviene adelantar una respuesta: buena parte de la literatura que "
        "se titula «reporting completeness» puntúa adherencia a una guía "
        "—CONSORT, CARE, STROBE—, y esta revisión no aplicó ninguna: contó "
        "presencia y ausencia de variables y de numerador y denominador, no "
        "adherencia normativa.",
        ("El resumen tiene %d palabras en la versión española y %d en la "
         "inglesa, y cumple el límite habitual de %d. Comprobar el límite exacto "
         "de la revista de destino."
         if resumen_max <= LIMITE_RESUMEN else
         "El resumen tiene %d palabras en la versión española y %d en la "
         "inglesa, y pasa del límite habitual de %d: hay que recortarlo, o "
         "comprobar que la revista de destino admite más.")
        % (S["palabras_resumen_es"], S["palabras_resumen_en"], LIMITE_RESUMEN),
    ]
    if not salvedad_t2:
        obs.append(
            "Las cifras de la sección 3.4 miden el reporte EN EL RESUMEN. El "
            "manuscrito lo aclara, y hace bien, pero conviene repetir esa "
            "salvedad en la leyenda de la Tabla 2, porque las tablas se leen "
            "sueltas.")
    if caja_fig1:
        obs.append(
            "La Figura 1 incluye una caja de recuperación de texto completo que "
            "no forma parte del flujo PRISMA. Está explicada en el pie, pero "
            "algún revisor la leerá como parte del diagrama. Considerar moverla "
            "a una figura suplementaria.")
    if not endolisinas_en_discusion:
        obs.append(
            "Falta un párrafo sobre por qué se excluyeron las endolisinas. La "
            "decisión es defendible y está en Métodos (§2.2; §2.4 declara que la "
            "regla se añadió a mitad de lote el 4 de agosto de 2026), pero "
            "merece una frase en Discusión, porque hay literatura que las cuenta "
            "como fagoterapia.")
    obs.append("Declarar en Métodos qué se hará si un estudio aparece publicado "
               "entre la búsqueda y el envío.")
    for t in obs:
        d.add_paragraph(t, style="List Bullet")

    d.add_heading("Sobre la vía de publicación", level=1)
    # Con la extraccion sin adjudicar habia tres salidas. Esa disyuntiva cayo:
    # la tabla seguia diciendo «No es alcanzable en 24 horas con 65 estudios»,
    # ofrecia GRADE (que no se aplica) y hablaba de «las cuatro limitaciones»
    # (el manuscrito declara ocho). Sin desacuerdos pendientes se dice lo que
    # se envia; la tabla de opciones vuelve solo si vuelven a quedar.
    if S["extraccion_conflictos_sin_firmar"]:
        d.add_paragraph(
            "Si el plazo no permite completar B1, hay dos salidas legítimas y "
            "una que no lo es.")
        tabla(d, ["Opción", "Qué se envía", "Viabilidad"],
              [("A. Completar la adjudicación",
                "Revisión sistemática con la completitud por desenlace y el "
                "riesgo de sesgo de los comparativos",
                "La correcta, si el plazo lo permite: quedan %d desacuerdos."
                % S["extraccion_conflictos_sin_firmar"]),
               ("B. Enviar como está, reencuadrado",
                "Revisión sistemática de la estructura del cuerpo de evidencia, "
                "de su recuperabilidad y de la completitud de su reporte, sin "
                "estimaciones de eficacia, con sus limitaciones declaradas",
                "Admisible en revistas metodológicas y en revistas del área que "
                "publican evidence mapping."),
               ("C. Publicar proporciones agrupadas",
                "Metaanálisis de proporciones de éxito",
                "NO. Los datos no lo sostienen y la sección 3.6 explica por qué. "
                "Hacerlo sería reproducir el defecto que este trabajo denuncia.")],
              [4.2, 6.4, 6.0])
    else:
        # Decia «El manuscrito se envia como ...», como cosa hecha. Si se envia
        # o no lo deciden los autores: el borrador del protocolo de la fase
        # definitiva dice que el informe piloto no se enviara como tal, y el
        # sobre esta preparado para enviarlo (hoja de hechos, 13-oct).
        d.add_paragraph(
            "La disyuntiva que planteaba esta sección cuando la extracción "
            "estaba sin adjudicar ya no existe. Enviar o no el manuscrito lo "
            "deciden los autores; si se envía, va como "
            "revisión sistemática de la estructura del cuerpo de evidencia, de "
            "su recuperabilidad y de la completitud de su reporte, con la "
            "completitud por desenlace (§3.5) y el riesgo de sesgo de los "
            "comparativos (§3.7), y sin estimaciones de eficacia. GRADE no se "
            "aplica: no hay un estimador agrupado cuya certeza calificar. "
            "Publicar proporciones agrupadas sigue descartado: los datos no lo "
            "sostienen, la sección 3.6 explica por qué, y hacerlo sería "
            "reproducir el defecto que este trabajo denuncia.")
    nota(d, "Que un cuerpo de evidencia no admita la síntesis que se le viene "
            "aplicando es un resultado por derecho propio, y el manuscrito lo "
            "sostiene con datos y no con opinión.")
    d.save(OUT / "S0_informe_editorial.docx")


def md_a_docx(origen, destino, titulo):
    """Vuelca el manuscrito a Word conservando la jerarquia y las negritas."""
    t = pathlib.Path(origen).read_text(encoding="utf-8")
    d = Document()
    est = d.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(11)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)
    for linea in t.split("\n"):
        l = linea.rstrip()
        if not l or l.strip() == "---":
            continue
        if l.startswith("#"):
            n = len(l) - len(l.lstrip("#"))
            d.add_heading(re.sub(r"\*", "", l.lstrip("# ").strip()),
                          level=min(n - 1, 4) if n > 1 else 0)
            continue
        estilo = None
        if l.startswith("- "):
            estilo, l = "List Bullet", l[2:]
        p = d.add_paragraph(style=estilo)
        # negrita y cursiva de Markdown, respetando las citas [@clave]
        for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", l):
            if not trozo:
                continue
            if trozo.startswith("**") and trozo.endswith("**"):
                p.add_run(trozo[2:-2]).bold = True
            elif trozo.startswith("*") and trozo.endswith("*"):
                p.add_run(trozo[1:-1]).italic = True
            else:
                p.add_run(trozo)
    d.save(destino)



def coma(x):
    """73.4 -> «73,4». El documento esta en castellano y el manuscrito usa coma."""
    return ("%s" % x).replace(".", ",")

def nombre_s5():
    """El listado lleva el recuento en el nombre, asi que se busca en disco."""
    hit = sorted(OUT.glob("S5_listado_*_estudios.csv"))
    # Como el resto de anexos de datos, se entrega en dos formatos: el .xlsx
    # para leer y el .csv para reejecutar el canal.
    return (hit[0].stem + ".xlsx + .csv") if hit else "S5_listado_estudios.xlsx + .csv"


def contenido_figuras_y_tablas():
    """Lo que trae de verdad la carpeta «figuras y tablas», leido del disco."""
    ft = OUT / "figuras y tablas"
    if not ft.is_dir():
        return "Figuras y tablas del manuscrito"
    figs = numerados(ft, r"figura_(\d+)_.*\.pdf$")
    tabs = numerados(ft, r"tabla_(\d+)_.*\.csv$")
    partes = []
    if figs:
        partes.append(("Figura %s" if len(figs) == 1 else "Figuras %s")
                      % rango(figs) + " en PDF vectorial")
    if tabs:
        partes.append(("tabla %s" if len(tabs) == 1 else "tablas %s")
                      % rango(tabs) + " del manuscrito en CSV")
    texto = "; ".join(partes) or "Figuras y tablas del manuscrito"
    if any(ft.glob("tabla_criterios*.csv")):
        texto += ", más la tabla de criterios de elegibilidad"
    return texto


def indice(S):
    R = estado_rob()
    rob_hecho = bool(R.get("completa")) and not R.get("celdas_pendientes")
    n18, n18_dos, _, _ = firmas_rob()
    ex = resumen_exclusiones()
    ray = validacion_rayyan()
    _, f_nopdf = firmas_criterio_texto()
    f_pp = primera_pasada()
    sin_texto = S["texto_completo_no_obtenido"]
    # Las remisiones del indice («§2.1», «Sección 3.2», «tablas 1 a 7») son del
    # maestro, y el indice viaja tambien en el sobre de la revista, cuyo
    # manuscrito no numera secciones y numera las tablas de otro modo.
    d = doc_nuevo("Verificables — Revisión sistemática de fagoterapia en "
                  "Pseudomonas aeruginosa multirresistente",
                  "Índice del paquete. Cada archivo responde a un ítem de "
                  "PRISMA 2020, a un requisito del ICMJE o a una comprobación "
                  "metodológica. Las remisiones a secciones (§) y a tablas son "
                  "las del manuscrito maestro (el informe extendido), no las "
                  "del manuscrito preparado para la revista.")
    tabla(d, ["Archivo", "Qué contiene", "Responde a"],
          # Decia «Uso interno», y el paquete lo manda a la revista. El
          # manuscrito declara S0 como anexo del material suplementario, asi
          # que lo que estaba mal era la etiqueta, no que viaje: se adjunta a
          # proposito, para que el comite vea que carencias se reconocen.
          # Despues decia «Evaluacion editorial previa» y «PRISMA item 27»: lo
          # redacto el modelo en el papel de un editor, y no es un informe de
          # una revista ni responde a ese item (disponibilidad de datos).
          [("S0_informe_editorial.docx",
            "Autoevaluación redactada por el modelo de lenguaje (Claude Opus 5) "
            "a petición de los autores; no es un informe de un editor ni una "
            "revisión externa. Qué bloquea la publicación y qué solo la mejora",
            "Se adjunta a propósito; no es un requisito de PRISMA ni del ICMJE"),
           # Los dos decian «Envio», y el ingles «para envio a revista». A la
           # revista va manuscrito_JSR_final, que deriva del maestro; el ingles
           # no viaja en el sobre.
           ("manuscrito_es.docx",
            "Manuscrito maestro completo, en español, en Word; el de la "
            "revista deriva de él", "Referencia"),
           ("manuscrito_revision_sistematica.pdf",
            "El mismo manuscrito maestro, maquetado en PDF", "Referencia"),
           ("manuscript_en.docx",
            "Versión en inglés del manuscrito maestro, en Word", "Referencia"),
           ("manuscript_systematic_review_en.pdf",
            "La versión en inglés, maquetada en PDF", "Referencia"),
           ("S1_lista_PRISMA_2020.docx",
            "Lista de comprobación ítem por ítem, con estado y ubicación",
            "PRISMA 2020, ítem 27"),
           # Prometia «fecha» por fuente, y no la hay: la ejecucion de Scopus
           # cuya exportacion entro al corpus no tiene fecha registrada.
           ("S2_estrategias_de_busqueda.docx",
            "Sintaxis literal por fuente, con su número de resultados",
            "PRISMA 2020, ítem 7; el 6, en parte: la ejecución de Scopus cuya "
            "exportación entró al corpus no tiene fecha registrada"),
           # No decia quien emitio las decisiones. Despues decia «cada decision
           # de cribado por titulo», y la primera pasada del 4-ago sobre el
           # brazo A de PubMed no esta en S3: S3 lleva el cribado vigente.
           ("S3_decisiones_etapa2_titulo.xlsx + .csv",
            "Registro solo-anexar de cada decisión del cribado vigente por "
            "título, con su motivo y su marca de tiempo; las emitió el modelo de "
            "lenguaje (Claude Opus 5) como revisor único, sin duplicación" +
            ((". La primera pasada del %s sobre el brazo A de PubMed, que el "
              "cribado del pozo sustituyó, consta en el repositorio y no en "
              "S3%s" % f_pp[:2]) if f_pp[0] else ""),
            "PRISMA 2020, ítem 8"),
           ("S3_decisiones_etapa3_resumen.xlsx + .csv",
            "Ídem para el cribado por resumen, con la corriente PRISMA; una "
            "corrección es una fila nueva y la anterior se conserva",
            "PRISMA 2020, ítems 8 y 16"),
           ("S4_pre_extraccion_desde_resumen.xlsx + .csv",
            "Pre-extracción desde el resumen de los %d estudios, emitida por "
            "el modelo de lenguaje y marcada como parcial"
            % S["estudios_extraibles"],
            "PRISMA 2020, ítems 9 y 10"),
           # «Si se obtuvo el texto completo»: desde NOPDF lo tienen todos.
           (nombre_s5(),
            "Los %d estudios del corpus con su situación, diseño, procedencia "
            "e informes agrupados; la agrupación la propuso un guion escrito "
            "por el modelo" % S["estudios"], "PRISMA 2020, ítem 17"),
           # Decia «Prueba de que el cribado no perdio estudios elegibles
           # conocidos»: se fijaron 42, la comprobacion alcanzaba a los 40 con
           # PMID, y la enmienda de idioma excluyo a uno de ellos.
           ("S6_auditoria_controles_positivos.docx",
            "Auditoría del cribado contra los %d controles positivos fijados "
            "antes de empezarlo: no perdió ninguno de los %d comprobables por "
            "PMID hasta la enmienda de idioma, que excluyó a Ronit 2024; hoy "
            "siguen en el corpus %d de los %d"
            % (S["controles_fijados"], S["controles_comprobables_pmid"],
               S["controles_en_corpus"], S["controles_fijados"]),
            "Metodológico; refuerza el ítem 8"),
           ("S7_declaraciones_ICMJE.docx",
            "Financiación, conflictos, CRediT, registro, datos y la declaración "
            "de uso de IA, copiada literal del manuscrito",
            "ICMJE; PRISMA ítems 24 a 27"),
           # «Justifica el 46,5 % de §3.2»: esa cifra ya no esta en el
           # manuscrito, que da el 100 % por construccion y la recuperacion
           # medida justo antes del criterio.
           ("S8_recuperacion_texto_completo.xlsx + .csv",
            "Vía de recuperación resuelta para cada informe sin DOI editorial",
            "Sección 3.2 del manuscrito"),
           ("S9_idioma_por_informe_y_clase_de_evidencia.xlsx + .csv",
            # «Incluido»: S9 trae todos los que superaron el cribado, tambien
            # los que salieron despues.
            "Idioma de cada informe que superó el cribado, con la clase de "
            "evidencia y la "
            "prueba concreta", "PRISMA 2020, ítem 5"),
           # Decia «Idioma determinado sobre el texto completo de cada PDF
           # recuperado» de una copia del 13-ago que no cubria el corpus vigente.
           # Se describe lo que trae el fichero en disco (s10_en_disco).
           *[(nombre, desc, "PRISMA 2020, ítem 5")
             for nombre, desc, _ in s10_en_disco(S)],
           ("S11_concordancia_entre_extractores.pdf",
            "Concordancia entre las dos extracciones, variable a variable, antes "
            "de resolver nada; su independencia no está acreditada",
            "PRISMA 2020, ítem 10a"),
           # «Donde esta hoy cada desacuerdo y con que se resuelve»: no queda
           # ninguno abierto, y no decia a quien anota el registro.
           ("S12_resolucion_de_conflictos.pdf",
            "Cómo se cerró cada uno de los %d desacuerdos de la extracción y a "
            "quién anota el registro como responsable" % S["extraccion_desacuerdos"],
            "PRISMA 2020, ítem 10a"),
           # No decia quien redacto las reglas ni que el registro no recoge su
           # ratificacion por los dos.
           ("S13_reglas_de_extraccion_de_desenlaces.pdf",
            "Definición de erradicación y de éxito clínico, con su corrección "
            "documentada. Las redactó el modelo de lenguaje; según declara D.V., "
            "las decidieron los dos autores, y el registro del 12 de agosto de "
            "2026 las dejó como propuesta sin ratificar",
            "PRISMA 2020, ítems 10a y 13"),
           # «Tras el consenso» y «la firmo el consenso»: el registro anota a los
           # dos solo en una parte de los desacuerdos y un unico responsable en
           # el resto. Las etiquetas son las que escribe hoy S15.
           ("S14_extraccion_adjudicada.xlsx + .csv",
            "Extracción definitiva: un valor por casilla, brazo a brazo, tras "
            "la adjudicación de los desacuerdos",
            "PRISMA 2020, ítems 10a y 19"),
           ("S15_procedencia_de_cada_casilla.xlsx + .csv",
            "De dónde sale el valor de cada casilla de S14: acuerdo, adjudicada "
            "por los dos, adjudicada por uno, corregida contra el texto, una "
            "sola lectura o abierta", "PRISMA 2020, ítem 9"),
           # El indice se quedaba en S15 y el paquete trae hasta S24.
           ("S16_excluidos_tras_leer_el_texto_completo.xlsx + .csv",
            # Decia «71 no llevan cita del articulo, entre ellos los 63 ...»: sin
            # cita del articulo hay mas, porque parte de las que llevan cita
            # citan la ficha del registro o la indexacion (frase_citas_s16).
            "Los %d estudios que el cribado admitió y que salieron después, con "
            "su código de motivo y quién firmó cada exclusión (%s). %s"
            % (ex["n"],
               "; ".join(["los dos autores en %d" % ex["dos"]] +
                         ["%s en %d" % (quien_firmo(q), ex["una_por_quien"][q])
                          for q in ex["una_quien"]]),
               frase_citas_s16(ex, S)),
            "PRISMA 2020, ítem 16b"),
           ("S17_tabla6_brazo_a_brazo.xlsx + .csv",
            "La Tabla 6 desglosada: los requisitos para agrupar, brazo a brazo, "
            "con el primero que falla y por qué", "PRISMA 2020, ítem 13a"),
           ("S18_cada_juicio_con_su_frase.xlsx + .csv",
            "Los %d juicios de riesgo de sesgo, uno por fila, con lo que sostiene "
            "cada uno en el cuaderno firmado: una cita literal en %d, una nota "
            "sin cita en %d y nada en %d"
            % (S["celdas_tabla5"], S["citas_literales"], S["notas_del_revisor"],
               S["juicios_sin_frase"]),
            "PRISMA 2020, ítems 11 y 18"),
           ("S18b_frases_cosechadas_por_tema.xlsx + .csv",
            "Las frases de cada artículo que el modelo de lenguaje localizó de "
            "forma automática para cada dominio de riesgo de sesgo y puso como "
            "comentario en los formularios; no sustituyen a lo que sostiene cada "
            "juicio (S18)", "PRISMA 2020, ítem 11"),
           ("S19_cribado_con_modelo_de_lenguaje.docx",
            "Qué modelo emitió las decisiones de cribado, cuándo y sobre qué "
            "ficheros; qué examinó D.V., según su declaración, y qué no quedó "
            "registrado", "PRISMA 2020, ítem 8"),
           ("S20_de_que_informe_salio_cada_brazo.xlsx + .csv",
            "De qué informe salió cada brazo extraído, y con qué prueba se "
            "estableció el vínculo", "PRISMA 2020, ítem 10a"),
           ("S21_solapamiento_de_pacientes.xlsx + .csv",
            "Los %d pares de estudios que podían describir a los mismos "
            "pacientes, con la señal que los marcó, el veredicto tras leer los "
            "dos artículos y la frase firmada en que se apoya"
            % S["pares_solapamiento_examinados"],
            "Metodológico; pacientes contados dos veces"),
           ("S22_completitud_resumen_frente_a_texto.xlsx + .csv",
            "Para cada variable de estratificación, cuántos estudios la declaran "
            "en el texto completo, cuántos la mencionan sin poder clasificarla y "
            "cuántos callan", "Metodológico; separa el resumen del texto"),
           ("S23_clase_de_resistencia_comprobada.xlsx + .csv",
            "La clase de resistencia comprobada contra el antibiograma de cada "
            "artículo, con la frase o la tabla en que se apoya y qué cambió en "
            "la extracción", "PRISMA 2020, ítem 10a"),
           ("S24_comparativos_reextraidos.xlsx + .csv",
            "Los estudios con diseño comparativo, reextraídos: comparador, "
            "asignación, tiempo cero, cointervenciones, pérdidas, estimando y "
            "si hay un contraste para P. aeruginosa", "PRISMA 2020, ítem 10a"),
           ("00_GUIA_DEL_MATERIAL_SUPLEMENTARIO.pdf",
            "Guía en lengua llana: qué pregunta contesta cada fichero",
            "Orientación para el revisor"),
           # Decia «Figuras 1 y 2 ...; tablas 1 a 4 en CSV», tecleado, y la
           # carpeta trae mas tablas. Se lee del disco.
           ("figuras y tablas/", contenido_figuras_y_tablas(),
            "PRISMA 2020, ítem 16a")],
          [5.6, 7.4, 3.6])

    # Cada estado se escribe desde el canal. Salian caducados: «Riesgo de sesgo
    # y GRADE: PENDIENTE — depende de la adjudicacion» con el riesgo de sesgo
    # firmado y GRADE descartado; «Recuperacion de texto completo: PARCIAL —
    # 65 de 65 (100.0 %)», que se contradice; «Cribado: COMPLETO — 65
    # estudios», que es el corpus tras el texto completo y no decia quien
    # emitio las decisiones; «Rayyan: 0 falsos negativos en 350» sin decir que
    # fue un solo revisor ni lo que paso con los 6 de la primera pasada;
    # «PROSPERO: PENDIENTE», que ya no admite la revision; y «Manuscrito:
    # COMPLETOS», sin la aprobacion escrita de N.T.
    rev = iniciales(ray["revisor"])
    un_revisor = bool(ray["revisor"]) and " y " not in ray["revisor"]
    incl1 = S["validacion_primera_pasada_incluidos"]
    if un_revisor:
        # «recribio» era una errata por «recribo».
        rayyan = ("COMPLETA — un solo revisor (%s) recribó a ciegas %d "
                  "registros excluidos" % (rev, S["validacion_muestra"]))
    else:
        rayyan = ("COMPLETA — %d registros excluidos, recribados a ciegas"
                  % S["validacion_muestra"])
    if incl1:
        rayyan += ("; en la primera pasada, sin los criterios a la vista, marcó "
                   "%d como incluibles y, al contrastarlos con los criterios, %s "
                   "%s como excluidos"
                   % (incl1, iniciales(ray["quien_reclasifica"]) or rev,
                      "los reclasificó" if ray["reclasificados"] == incl1 else
                      "reclasificó %d" % ray["reclasificados"]))
    rayyan += (": %d falsos negativos en %d"
               % (S["validacion_falsos_negativos"], S["validacion_muestra"]))
    tabla(d, ["Elemento", "Estado"],
          [("Búsqueda en %d fuentes" % S["fuentes_distintas_n"], "COMPLETA"),
           ("Cribado por título y resumen",
            "COMPLETO — decisiones emitidas por el modelo de lenguaje como "
            "revisor único, sin duplicación; %d informes pasaron a texto "
            "completo (%d estudios)"
            % (S["informes_a_texto_completo"], S["estudios_antes_de_releer"])),
           ("Elegibilidad a texto completo",
            "COMPLETA — %d de %d estudios excluidos (%d por no disponer del "
            "artículo completo); %d en el corpus"
            % (S["estudios_excluidos_tras_texto_completo"],
               S["estudios_antes_de_releer"],
               S["excluidos_texto_completo_NOPDF"], S["estudios"])),
           ("Verificación del criterio de idioma",
            "COMPLETA — %d informes, ninguno sin prueba; los %d estudios del "
            "corpus, comprobados sobre el cuerpo del PDF"
            % (S["informes_agrupados"], S["idioma_cuerpo_comprobados"])),
           ("Agrupación de informes en estudios",
            "COMPLETA — la propuso un guion escrito por el modelo (S5)"),
           ("Pre-extracción desde resumen",
            "COMPLETA — %d de %d, emitida por el modelo de lenguaje"
            % (S["estudios_extraibles"], S["estudios_extraibles"])),
           ("Recuperación de texto completo",
            ("PARCIAL — %d de %d (%s %%)"
             % (S["texto_completo_obtenido"], S["estudios_extraibles"],
                coma(S["texto_completo_pct"]))
             if sin_texto else
             "%s %% POR CONSTRUCCIÓN — %d de %d: desde el %s solo entran "
             "estudios con el artículo completo (segunda enmienda); justo antes "
             "del criterio, %d de %d (%s %%)"
             % (coma(S["texto_completo_pct"]), S["texto_completo_obtenido"],
                S["estudios_extraibles"], f_nopdf, S["texto_completo_obtenido"],
                S["texto_completo_antes_de_nopdf_base"],
                coma(S["texto_completo_antes_de_nopdf_pct"])))),
           ("Extracción por duplicado",
            "COMPLETA — %d de %d estudios (%s %%), sin independencia acreditada"
            % (S["extraccion_estudios_ambos"], S["extraccion_estudios_r1"],
               S["extraccion_doble_pct"])),
           # «(14 por los dos, 525 por uno solo)» se leia como que lo decidio
           # uno solo; lo comprobado es a quien anota el registro.
           ("Adjudicación de los desacuerdos",
            ("PENDIENTE — %d de %d sin firmar; bloquea reportar desenlaces"
             % (S["extraccion_conflictos_sin_firmar"], S["extraccion_desacuerdos"])
             if S["extraccion_conflictos_sin_firmar"] else
             "COMPLETA — %d de %d adjudicados: el registro anota a los dos en %d "
             "y un único responsable en %d (D.V. en %d, N.T. en %d); %d "
             "cerrados por una regla mecánica, sin firma conjunta"
             % (S["extraccion_conflictos_firmados"], S["extraccion_desacuerdos"],
                S["conflictos_resueltos_por_los_dos"],
                S["conflictos_resueltos_por_uno"],
                S["conflictos_resueltos_solo_dv"],
                S["conflictos_resueltos_solo_nt"],
                S["extraccion_cerrados_por_regla"]))),
           ("Validación del cribado en Rayyan", rayyan),
           ("Riesgo de sesgo",
            ("COMPLETO — %d comparativos evaluables (%d RoB 2, %d ROBINS-I), %d "
             "juicios%s; sin independencia acreditada entre los dos cuadernos"
             % (R["evaluables"], R["por_instrumento"].get("RoB 2", 0),
                R["por_instrumento"].get("ROBINS-I", 0), R["celdas_totales"],
                ", firmados por los dos autores"
                if n18 == n18_dos == R["celdas_totales"] else "")
             if rob_hecho else
             "PENDIENTE — faltan %d de los %d juicios"
             % (R.get("celdas_pendientes", 0), R.get("celdas_totales", 0)))),
           ("GRADE", "NO SE APLICA — no hay estimador agrupado cuya certeza "
                     "calificar"),
           ("Registro del protocolo",
            "NO SE REGISTRÓ — la ausencia se declara (§2.1); PROSPERO ya no "
            "admite una revisión con la extracción hecha"),
           ("Manuscrito y figuras",
            # Faltaba que N.T. tampoco ha aceptado responder por el trabajo.
            "REDACTADOS — los borradores los redactó el modelo de lenguaje; N.T. "
            "no ha aprobado ni revisado por escrito la versión final ni ha "
            "aceptado responder por el trabajo")],
          [8.0, 8.6])
    # Decia «Todas las cifras ... se generan desde el canal» sin decir quien
    # escribio el canal, y que la comprobacion automatica «falla si alguna
    # cifra del texto no procede de el»: solo mira si el valor existe entre
    # los escalares, no si es el que toca.
    nota(d, "Las cifras de este paquete y de los manuscritos las genera el canal "
            "de datos, cuyo código escribió y ejecutó el modelo de lenguaje. Unas "
            "comprobaciones automáticas recorren los manuscritos y fallan si una "
            "cifra no figura entre las del canal; que figure no garantiza que "
            "sea la que corresponde, y para eso una parte de las frases se "
            "comprueba contra la cifra concreta que debe llevar.")
    d.save(OUT / "00_INDICE.docx")

    # Un indice que nombra un fichero inexistente es peor que no tener indice:
    # el editor lo pide, no lo encuentra y deja de fiarse del resto del paquete.
    citados = set()
    for tab in d.tables:
        for fila in tab.rows:
            n = fila.cells[0].text.strip()
            if n and "." in n:
                citados.add(n)
    # Una entrada como «S3_x.xlsx + .csv» nombra dos ficheros. Comprobarla como
    # si fuera un nombre literal hace que el aviso salte sobre anexos que si
    # estan, y un aviso que siempre miente se deja de leer.
    def existe(nombre):
        partes = [x.strip() for x in nombre.split("+")]
        raiz = partes[0].rsplit(".", 1)[0]
        reales = [x if not x.startswith(".") else raiz + x for x in partes]
        return all((OUT / r).exists() for r in reales)

    faltan = sorted(n for n in citados if not existe(n))
    if faltan:
        print("  AVISO: el indice cita ficheros que no existen: %s"
              % ", ".join(faltan))
    else:
        print("  indice comprobado: todos los ficheros citados existen")

    # Y al reves. El indice se quedo en S15 mientras el paquete crecia hasta
    # S24, y la comprobacion de arriba no lo veia: solo mira lo citado.
    def reales(nombre):
        partes = [x.strip() for x in nombre.split("+")]
        raiz = partes[0].rsplit(".", 1)[0]
        return {x if not x.startswith(".") else raiz + x for x in partes}
    cubiertos = set().union(*(reales(n) for n in citados)) | {"00_INDICE.docx"}
    sin_indice = sorted(p.name for p in OUT.iterdir()
                        if p.is_file() and p.name not in cubiertos)
    if sin_indice:
        print("  AVISO: hay ficheros en el paquete que el indice no lista: %s"
              % ", ".join(sin_indice))
    else:
        print("  indice comprobado: lista todos los ficheros del paquete")


def main():
    OUT.mkdir(exist_ok=True)
    S = json.loads((ROOT / "quality_reports" / "synthesis_scalars.json")
                   .read_text(encoding="utf-8"))
    informe_editorial(S)
    print("  S0  informe editorial")
    md_a_docx(ROOT / "paper" / "manuscrito_revision_sistematica.md",
              OUT / "manuscrito_es.docx", "Manuscrito (español)")
    md_a_docx(ROOT / "paper" / "manuscript_systematic_review_en.md",
              OUT / "manuscript_en.docx", "Manuscript (English)")
    print("  manuscritos en Word, español e inglés")
    indice(S)
    print("  00  índice del paquete")
    return 0


if __name__ == "__main__":
    sys.exit(main())
