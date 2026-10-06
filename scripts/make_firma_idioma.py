# -*- coding: utf-8 -*-
"""Tres estudios cuyo idioma solo se probo sobre el resumen, para firmar su exclusion.

DE DONDE SALE

El 2026-10-06 D. Valdiviezo recordo que el corpus solo admite estudios
redactados en ingles o en espanol, y se comprobo el idioma de los 131 estudios
del corpus. Los 65 con PDF estan en ingles sobre el CUERPO; las 42 fichas de
registro son de registros que publican en ingles. De los 24 sin texto completo,
el idioma solo se juzgo sobre el titulo y el resumen, y en tres la revista no
publica de forma nativa en ingles:

  · EST-045 y EST-059, Infektsionnye Bolezni (Rusia). El 2026-09-01 se
    conservaron porque el DOI lleva a la ficha inglesa del editor; esa ficha no
    enlaza el articulo ni dice en que idioma esta.
  · EST-173, Surgical Chronicles (Grecia). Sin DOI; nunca se comprobo mas alla
    del resumen.

D. Valdiviezo respondio «EXCLUYE». Excluir es un acto de autoria y lo firman los
dos, asi que va a este cuaderno; `ingest_firma_idioma.py` lo aplica con las dos
firmas.

EL CODIGO. IDI afirma que el articulo NO esta en ingles ni en espanol, y eso no
esta comprobado. NOREC dice exactamente lo que pasa: «texto completo no
recuperado: no se pudo verificar contra el articulo». Ojo con la nota 4 de
CLAUDE.md: aqui no se extiende NOREC a los 24 sin texto, sino a tres con una
razon concreta para dudar del criterio, y se declara asi.

SALIDA
    ~/Desktop/FIRMAR_idioma_2026-10-06.xlsx
"""
import pathlib
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CUADERNO = pathlib.Path.home() / "Desktop" / "FIRMAR_idioma_2026-10-06.xlsx"
OPCIONES = ["EXCLUIR con NOREC (idioma no verificable)",
            "EXCLUIR con IDI (consta que no es inglés ni español)",
            "MANTENER y declarar que solo se verificó el resumen"]

FILAS = [
    ("EST-045", "Infektsionnye Bolezni (Rusia), 2017",
     "Concept of individualized medicine based on personalized phage therapy for intensive "
     "care unit patients suffering from healthcare-associated infections",
     "Resumen en inglés (1 161 caracteres). DOI 10.20953/1729-9225-2017-4-49-54 → ficha inglesa "
     "del editor (phdynasty.ru/en/), que no enlaza el artículo ni declara su idioma. Sin texto "
     "completo: no es acceso abierto."),
    ("EST-059", "Infektsionnye Bolezni (Rusia), 2017",
     "Anti-phage antibody response in phage therapy against healthcare-associated infections (HAIs)",
     "Resumen en inglés (1 337 caracteres). DOI 10.20953/1729-9225-2017-1-35-40 → ficha inglesa del "
     "editor, sin enlace al artículo. Zenodo 1997795 guarda un .docx de 17 756 bytes, que no se "
     "ha descargado."),
    ("EST-173", "Surgical Chronicles (Grecia), 2016",
     "Use of bacteriophages in the treatment of infected wounds in patients who have allergy to "
     "antibiotics",
     "Resumen en inglés (1 994 caracteres). Sin DOI ni otro identificador consultable; nunca se "
     "comprobó más allá del resumen."),
]
QUE_CAMBIA = ("Ninguno tiene texto completo ni datos de desenlace extraídos (EXTRACTION_INCOMPLETE). "
              "Sale del corpus y de los 24 sin texto completo; no toca ningún brazo con datos.")


def main():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if CUADERNO.exists():
        wb = load_workbook(CUADERNO)
        if any(c.value not in (None, "") for ws in wb.worksheets for row in ws.iter_rows()
               for c in row if c.fill and str(c.fill.fgColor.rgb).endswith("FCE4D6")):
            raise SystemExit("NO SE ESCRIBE: el cuaderno ya tiene respuestas.")

    negrita, grande = Font(bold=True), Font(bold=True, size=13)
    blanca, azul = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    ojo = PatternFill("solid", fgColor="FCE4D6")
    arriba = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()
    s = wb.active
    s.title = "Idioma"
    s.cell(row=1, column=1, value="Tres estudios cuyo idioma solo se probó sobre el resumen").font = grande
    s.cell(row=2, column=1, value="D. Valdiviezo pidió excluirlos el 6 de octubre. NOREC es el código "
                                  "exacto si nadie ha visto el artículo; IDI solo si consta que no está "
                                  "en inglés ni en español.")
    cab = ["Estudio", "Revista y año", "Título", "Qué se comprobó", "Qué cambia si sale",
           "Vuestra decisión", "Por qué"]
    for j, (c, w) in enumerate(zip(cab, [10, 26, 44, 60, 40, 34, 36]), start=1):
        x = s.cell(row=3, column=j, value=c)
        x.font, x.fill, x.alignment = blanca, azul, arriba
        s.column_dimensions[x.column_letter].width = w
    for i, (e, rev, tit, comp) in enumerate(FILAS, start=4):
        for j, v in enumerate([e, rev, tit, comp, QUE_CAMBIA], start=1):
            s.cell(row=i, column=j, value=v).alignment = arriba
        for j in (6, 7):
            s.cell(row=i, column=j).fill = ojo
            s.cell(row=i, column=j).alignment = arriba
        dv = DataValidation(type="list", allow_blank=True, formula1='"%s"' % ",".join(OPCIONES))
        s.add_data_validation(dv)
        dv.add(s.cell(row=i, column=6))
        s.row_dimensions[i].height = 110

    fi = wb.create_sheet("Firma")
    fi.column_dimensions["A"].width = 46
    fi.column_dimensions["B"].width = 54
    fi.cell(row=1, column=1, value="Firma de los dos autores").font = negrita
    fi.cell(row=2, column=1, value="Se firman juntas o no se firma ninguna.")
    for i, k in enumerate(["Revisor 1 (nombre completo)", "Revisor 2 (nombre completo)",
                           "Fecha (AAAA-MM-DD)",
                           "¿Habéis leído los artículos citados en cada hoja?"], start=4):
        fi.cell(row=i, column=1, value=k).font = negrita
        fi.cell(row=i, column=2, value="").fill = ojo
    wb.save(CUADERNO)
    print("cuaderno: %s (%d estudios)" % (CUADERNO, len(FILAS)))


if __name__ == "__main__":
    main()
