"""Pasa el resumen estructurado a .docx, con las cuatro versiones y su recuento.

POR QUE EXISTE. `paper/resumen_estructurado.md` es de donde se copia y se pega
en el portal de la revista, y para eso el markdown vale. Pero el resumen viaja
tambien a un director o a un coautor, y ahi hace falta Word.

QUE NO ES. No es un documento de envio. El resumen que se manda ya va dentro de
`manuscrito_JSR_final.docx`: esto es la hoja de trabajo con las cuatro
versiones juntas, para elegir y pegar. La cabecera lo dice en cada pagina.

EL RECUENTO NO SE TECLEA. Las cabeceras de seccion del .md ya traen las
palabras que conto `build_structured_abstract.py` --como las cuenta Word,
etiquetas y signos incluidos-- y se copian tal cual. Las dos versiones por
etiquetas se pasan de 250 a proposito y no son las que se envian; el .docx
marca cual es cual para que nadie pegue la que no toca.

Salida:
    paper/docx/resumen_estructurado.docx

Uso:
    python scripts/build_resumen_docx.py
"""
import pathlib
import re
import sys

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches, Mm, RGBColor
except ImportError:
    raise SystemExit("hace falta python-docx: python -m pip install python-docx")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_jsr_docx import escribe

ROOT = pathlib.Path(__file__).resolve().parent.parent
FUENTE = ROOT / "paper" / "resumen_estructurado.md"
DEST = ROOT / "paper" / "docx"
TNR = "Times New Roman"
LIMITE = 250

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def documento():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, lado, Inches(1))
    n = d.styles["Normal"]
    n.font.name, n.font.size = TNR, Pt(11)
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.3
    cab = s.header.paragraphs[0]
    cab.text = ("Resumen estructurado — hoja de trabajo. El resumen que se envía "
                "ya va dentro de manuscrito_JSR_final.docx")
    cab.runs[0].font.size, cab.runs[0].font.name = Pt(9), TNR
    cab.runs[0].italic = True
    return d


def veredicto(titulo):
    """De «... — 251 palabras» saca el recuento y dice si esa version se envia.

    La regla la fija `build_structured_abstract.py` y aqui solo se repite: las
    de parrafo corrido son las que van al manuscrito, las de etiquetas no.
    """
    m = re.search(r"—\s*(\d+)\s*palabras", titulo)
    n = int(m.group(1)) if m else None
    corrido = "corrido" in titulo.lower() or "single paragraph" in titulo.lower()
    if not corrido:
        return n, "No se envía. Versión por etiquetas, de referencia."
    if n is not None and n > LIMITE:
        return n, "SE PASA DEL LÍMITE de %d. Revisar antes de enviar." % LIMITE
    return n, "Esta es la que va en el manuscrito. Dentro del límite de %d." % LIMITE


def main():
    if not FUENTE.exists():
        raise SystemExit("no encuentro %s; corre antes build_structured_abstract.py"
                         % FUENTE.name)
    d = documento()
    secciones = 0
    for l in FUENTE.read_text(encoding="utf-8").split("\n"):
        l = l.rstrip()
        if not l:
            continue
        if l.startswith("# "):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            escribe(p, l[2:], size=15, negrita=True)
            continue
        if l.startswith("## "):
            secciones += 1
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(18 if secciones > 1 else 14)
            p.paragraph_format.space_after = Pt(2)
            escribe(p, l[3:], size=12, negrita=True)
            n, nota = veredicto(l[3:])
            aviso = d.add_paragraph()
            aviso.paragraph_format.space_after = Pt(7)
            aviso.paragraph_format.line_spacing = 1.0
            r = aviso.add_run(nota)
            r.font.name, r.font.size, r.italic = TNR, Pt(9), True
            r.font.color.rgb = (RGBColor(0xA0, 0x30, 0x20)
                                if "PASA" in nota else RGBColor(0x55, 0x55, 0x55))
            continue
        # La nota de cabecera del .md explica de donde sale el fichero; en el
        # .docx sobra el «no editar a mano», que se refiere al markdown.
        if l.startswith("Generado por") or l.startswith("No editar a mano"):
            continue
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribe(p, l, size=11)

    DEST.mkdir(parents=True, exist_ok=True)
    salida = DEST / "resumen_estructurado.docx"
    try:
        d.save(salida)
    except PermissionError:
        salida = DEST / "resumen_estructurado_NUEVO.docx"
        d.save(salida)
        print("AVISO: estaba abierto en Word; escrito al lado como %s" % salida.name)
    print("escrito %s  (%d KB)" % (salida.relative_to(ROOT), salida.stat().st_size // 1024))
    print("  %d versiones del resumen" % secciones)


if __name__ == "__main__":
    main()
