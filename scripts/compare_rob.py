"""Compara las dos evaluaciones de riesgo de sesgo: concordancia y conflictos.

QUÉ MIDE. Por cada pregunta de cada instrumento: porcentaje de acuerdo y kappa
de Cohen sobre el valor tal cual lo eligieron del desplegable. No hay que
normalizar nada porque el vocabulario es cerrado: o eligieron la misma opción o
no.

LA KAPPA NO SIEMPRE EXISTE, y aquí va a pasar mucho. Si los dos revisores
respondieron "No" a la misma pregunta en los cuarenta reportes de caso, no hay
variación y la kappa sale 0/0. Se informa "no calculable (sin variación)", no un
cero que parecería desacuerdo total. Con 3 estudios en RoB 2, casi ninguna de
sus 22 preguntas va a tener kappa; eso NO es un defecto de la evaluación, es lo
que ocurre cuando el denominador es tres.

NO PISA LAS FIRMAS. Al volver a ejecutarse recupera las resoluciones ya escritas
por la clave (estudio, instrumento, pregunta) y dice cuántas conservó. El
comparador de la extracción llegó a reescribir el fichero con las columnas de
resolución EN BLANCO, y ejecutar el comando documentado habría borrado 561
firmas. No se quita ese paso.

UNA CELDA VACÍA NO ES UN DESACUERDO. Si uno respondió y el otro no llegó a esa
fila, es cobertura, no discrepancia: se cuenta aparte. Mezclarlo hundiría la
kappa atribuyendo a desacuerdo lo que es trabajo a medias.

USO
    python scripts/compare_rob.py
    python scripts/compare_rob.py --a ruta/uno.xlsx --b ruta/otro.xlsx
"""
import argparse
import collections
import csv
import datetime
import pathlib
import sys

try:
    import openpyxl
except ImportError:
    raise SystemExit("hace falta openpyxl: python -m pip install openpyxl")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from rob_instruments import INSTRUMENTOS, SIMPLIFICADOS

ROOT = pathlib.Path(__file__).resolve().parent.parent
DEST = ROOT / "revision_sistematica" / "riesgo_sesgo"
CONFLICTOS = DEST / "riesgo_sesgo_conflictos.csv"
CONFLICTOS_SIMPLE = DEST / "riesgo_sesgo_comparativos_conflictos.csv"
INFORME = ROOT / "quality_reports" / "rob_agreement.md"
INFORME_SIMPLE = ROOT / "quality_reports" / "rob_agreement_comparativos.md"

COLS = ["study_id", "instrumento", "item", "pregunta", "valor_a", "valor_b",
        "frase_a", "frase_b", "resolucion", "resuelto_por", "fecha"]

# A partir de cuantas respuestas comparables el acuerdo perfecto deja de ser
# creible. Con cinco celdas dos revisores pueden coincidir en todo; con ochenta
# no. El umbral es deliberadamente bajo: mas vale preguntar de mas.
SOSPECHA_ACUERDO_TOTAL = 20

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def kappa(pares):
    """Kappa de Cohen. Devuelve (valor|None, motivo si es None)."""
    n = len(pares)
    if n == 0:
        return None, "sin filas comparables"
    cats = sorted({x for p in pares for x in p})
    if len(cats) < 2:
        return None, "no calculable (sin variación: una sola categoría)"
    po = sum(1 for a, b in pares if a == b) / n
    ca = collections.Counter(a for a, _ in pares)
    cb = collections.Counter(b for _, b in pares)
    pe = sum((ca[c] / n) * (cb[c] / n) for c in cats)
    if abs(1 - pe) < 1e-12:
        return None, "no calculable (acuerdo esperado = 1)"
    return (po - pe) / (1 - pe), None


def etiqueta(k):
    if k is None:
        return "—"
    for lim, txt in ((0.81, "casi perfecta"), (0.61, "sustancial"),
                     (0.41, "moderada"), (0.21, "aceptable"), (0.0, "leve")):
        if k >= lim:
            return txt
    return "pobre"


def lee_libro(p, instrumentos=None):
    """{(estudio, instrumento, item): valor}. El juicio global va como item 'GLOBAL'."""
    wb = openpyxl.load_workbook(p, data_only=True)
    fuera = {}
    for inst in (instrumentos or INSTRUMENTOS):
        hoja = inst["hoja"][:31]
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        n_ctx = 5                       # las columnas de contexto del formulario
        codigos = [it["codigo"] for it in inst["items"]] + ["GLOBAL"]
        for fila in ws.iter_rows(min_row=4):
            s = fila[0].value
            if not s:
                continue
            for k, cod in enumerate(codigos):
                j = n_ctx + k
                if j >= len(fila):
                    break
                v = fila[j].value
                fuera[(str(s).strip(), inst["clave"], cod)] = (
                    str(v).strip() if v is not None else "")
    return fuera


def notas(p, instrumentos=None):
    """Las dos columnas de texto libre del final: {(estudio, inst): (frase, duda)}.

    Van DESPUES del juicio global, asi que `lee_libro` no las ve --y no debe
    verlas: se colarian como si fueran items y el ingestor las trataria como
    juicios--. Se leen aparte porque hacen falta al resolver: sentarse a firmar
    un desacuerdo viendo dos etiquetas y ningun motivo no es resolverlo, es
    elegir. Aqui se recogen para que el fichero de conflictos lleve al lado la
    frase en que se apoyo cada uno.
    """
    wb = openpyxl.load_workbook(p, data_only=True)
    fuera = {}
    for inst in (instrumentos or INSTRUMENTOS):
        hoja = inst["hoja"][:31]
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        # 5 de contexto + los items + el juicio global; lo siguiente es la
        # frase, y despues la duda sobre el diseno.
        j = 5 + len(inst["items"]) + 1
        for fila in ws.iter_rows(min_row=4):
            s = fila[0].value
            if not s or j >= len(fila):
                continue
            def v(k):
                return (str(fila[k].value).strip()
                        if k < len(fila) and fila[k].value is not None else "")
            fuera[(str(s).strip(), inst["clave"])] = (v(j), v(j + 1))
    return fuera


def previas(destino=None):
    """Las resoluciones ya firmadas, para no perderlas al reescribir."""
    destino = destino or CONFLICTOS
    if not destino.exists():
        return {}
    with open(destino, encoding="utf-8-sig", newline="") as fh:
        return {(r["study_id"], r["instrumento"], r["item"]): r
                for r in csv.DictReader(fh)
                if (r.get("resolucion") or "").strip()}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--a", default=str(DEST / "riesgo_sesgo_danny_valdiviezo.xlsx"))
    ap.add_argument("--b", default=None)
    ap.add_argument("--simplificado", action="store_true",
                    help="los formularios por dominio de los comparativos")
    args = ap.parse_args()
    instrumentos = SIMPLIFICADOS if args.simplificado else INSTRUMENTOS
    conflictos_p = CONFLICTOS_SIMPLE if args.simplificado else CONFLICTOS
    informe_p = INFORME_SIMPLE if args.simplificado else INFORME
    if args.simplificado and args.a == str(DEST / 'riesgo_sesgo_danny_valdiviezo.xlsx'):
        args.a = str(DEST / 'riesgo_sesgo_comparativos_danny_valdiviezo.xlsx')
    if args.b is None:
        args.b = str(DEST / ('riesgo_sesgo_comparativos_nataly_trelles.xlsx'
                             if args.simplificado
                             else 'riesgo_sesgo_nataly_trelles.xlsx'))

    pa, pb = pathlib.Path(args.a), pathlib.Path(args.b)
    for p in (pa, pb):
        if not p.exists():
            raise SystemExit("no encuentro %s" % p)
    A, B = lee_libro(pa, instrumentos), lee_libro(pb, instrumentos)
    NA, NB = notas(pa, instrumentos), notas(pb, instrumentos)
    guardadas = previas(conflictos_p)

    pregunta = {}
    for inst in instrumentos:
        for it in inst["items"]:
            pregunta[(inst["clave"], it["codigo"])] = it["texto_es"]
        pregunta[(inst["clave"], "GLOBAL")] = "JUICIO GLOBAL del estudio"

    por_item = collections.defaultdict(list)
    conflictos, sin_pareja, vacias = [], 0, 0
    for clave in sorted(set(A) | set(B)):
        s, inst, cod = clave
        va, vb = A.get(clave, ""), B.get(clave, "")
        if not va and not vb:
            vacias += 1
            continue
        if not va or not vb:
            sin_pareja += 1
            continue
        por_item[(inst, cod)].append((va, vb))
        if va != vb:
            prev = guardadas.get(clave, {})
            conflictos.append({
                "study_id": s, "instrumento": inst, "item": cod,
                "pregunta": pregunta.get((inst, cod), ""),
                "valor_a": va, "valor_b": vb,
                "frase_a": NA.get((s, inst), ("", ""))[0],
                "frase_b": NB.get((s, inst), ("", ""))[0],
                "resolucion": prev.get("resolucion", ""),
                "resuelto_por": prev.get("resuelto_por", ""),
                "fecha": prev.get("fecha", "")})

    DEST.mkdir(parents=True, exist_ok=True)
    with open(conflictos_p, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, COLS)
        w.writeheader()
        w.writerows(conflictos)

    dudas = []
    for clave, (_, duda) in sorted(NA.items()):
        if duda:
            dudas.append(("revisor A", clave[0], duda))
    for clave, (_, duda) in sorted(NB.items()):
        if duda:
            dudas.append(("revisor B", clave[0], duda))

    comparadas = sum(len(v) for v in por_item.values())
    # "Escrita" y "firmada" no son lo mismo, y el ingestor solo se queda con
    # las firmadas. Contarlas juntas aqui haria creer que hay mas consenso del
    # que hay.
    escritas = sum(1 for c in conflictos if c["resolucion"])
    firmadas = sum(1 for c in conflictos if c["resolucion"] and c["resuelto_por"])
    lineas = ["# Concordancia en la evaluación del riesgo de sesgo", "",
              "Generado el %s por `compare_rob.py`." % datetime.date.today().isoformat(),
              "", "| | |", "|---|---:|",
              "| Respuestas comparables (los dos contestaron) | %d |" % comparadas,
              "| Desacuerdos | %d |" % len(conflictos),
              "| Resueltos y firmados | %d |" % firmadas,
              "| Con valor escrito pero SIN firma (no se ingieren) | %d |"
              % (escritas - firmadas),
              "| Sin pareja (solo uno contestó) | %d |" % sin_pareja,
              "| Sin contestar por ninguno | %d |" % vacias, ""]

    if dudas:
        # Esto va ANTES que la concordancia: si el diseno adjudicado esta mal,
        # el instrumento esta mal, y la kappa de sus juicios no significa nada.
        lineas += ["## Dudas sobre el diseño adjudicado", "",
                   "Resolver esto primero: cambia qué instrumento se aplica.", ""]
        lineas += ["- **%s**, %s: %s" % d for d in dudas]
        lineas += [""]

    if comparadas:
        acuerdo = 100.0 * (comparadas - len(conflictos)) / comparadas
        lineas += ["Acuerdo global: **%.1f %%** sobre %d respuestas.\n" % (acuerdo, comparadas)]

    # EL ACUERDO PERFECTO NO ES UN RESULTADO, ES UN SINTOMA. Dos personas que
    # juzgan por separado ochenta dominios discrepan en alguno: es exactamente
    # lo que la doble evaluacion existe para medir. Cero desacuerdos sobre
    # muchas celdas casi siempre significa que los dos cuadernos no son
    # independientes --uno copiado del otro, o los dos rellenados por la misma
    # persona--, y entonces la kappa de 1,00 no mide concordancia entre
    # revisores: mide que el fichero es el mismo. Es lo primero que comprueba
    # un arbitro, y el guion no puede presentarlo como un hallazgo.
    acuerdo_total = comparadas >= SOSPECHA_ACUERDO_TOTAL and not conflictos
    if acuerdo_total:
        lineas[1:1] = [
            "",
            "> **ATENCIÓN: los dos cuadernos coinciden en las %d respuestas, sin "
            "una sola discrepancia.**" % comparadas,
            ">",
            "> Eso no es lo que produce una evaluación independiente por dos "
            "revisores; es lo que produce un cuaderno copiado del otro, o los "
            "dos rellenados por la misma persona. Hay que resolver si las dos "
            "evaluaciones son de verdad independientes antes de usar nada de "
            "este informe.",
            ">",
            "> Mientras no se resuelva, la kappa de más abajo **no mide "
            "concordancia entre revisores**, no debe reportarse, y la frase de "
            "Métodos que declara evaluación independiente con resolución por "
            "consenso sería falsa.",
        ]

    kk = []
    for inst in instrumentos:
        filas = [(cod, por_item[(inst["clave"], cod)])
                 for cod in [it["codigo"] for it in inst["items"]] + ["GLOBAL"]
                 if por_item.get((inst["clave"], cod))]
        if not filas:
            continue
        lineas += ["## %s" % inst["nombre"], "",
                   "| Pregunta | n | Acuerdo | Kappa | |", "|---|---:|---:|---:|---|"]
        for cod, pares in filas:
            k, motivo = kappa(pares)
            ac = 100.0 * sum(1 for a, b in pares if a == b) / len(pares)
            if k is not None:
                kk.append(k)
            lineas.append("| %s | %d | %.0f %% | %s | %s |" % (
                cod, len(pares), ac,
                "%.2f" % k if k is not None else "—",
                etiqueta(k) if k is not None else motivo))
        lineas.append("")

    if kk:
        kk.sort()
        med = kk[len(kk) // 2] if len(kk) % 2 else (kk[len(kk) // 2 - 1] + kk[len(kk) // 2]) / 2
        lineas += ["Kappa mediana sobre las %d preguntas en que es calculable: "
                   "**%.2f** (%s).\n" % (len(kk), med, etiqueta(med))]
    else:
        lineas += ["Ninguna pregunta admite kappa todavía: hace falta variación "
                   "en las respuestas, y con pocos estudios por instrumento no "
                   "siempre la hay.\n"]

    informe_p.write_text("\n".join(lineas), encoding="utf-8")

    print("comparadas %d respuestas · %d desacuerdos (%d firmados, %d escritos "
          "sin firma)" % (comparadas, len(conflictos), firmadas,
                          escritas - firmadas))
    print("sin pareja %d · sin contestar %d" % (sin_pareja, vacias))
    if acuerdo_total:
        print()
        print("ATENCION: %d respuestas y CERO desacuerdos." % comparadas)
        print("  Una evaluacion independiente por dos revisores no sale asi.")
        print("  Comprueba si los dos cuadernos son de verdad independientes")
        print("  antes de usar la kappa o de ingerir nada.")
    if dudas:
        print("%d duda(s) sobre el diseño adjudicado, al principio del informe"
              % len(dudas))
    print("escrito %s" % conflictos_p.relative_to(ROOT))
    print("escrito %s" % informe_p.relative_to(ROOT))


if __name__ == "__main__":
    main()
