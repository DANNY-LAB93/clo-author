"""Convierte los anexos del paquete en cuadernos que un revisor pueda leer.

EL PROBLEMA

Los anexos se exportan como CSV con los nombres internos del canal:
`record_id`, `verdict`, `decided_by`, `pathogen_scope`, `extraction_status`. Un
revisor que abre S3 encuentra 13 917 filas de `EXCLUDE` y `ADVANCE` sin saber
qué es una fila, cuántas debería haber, ni qué significa la columna. El fichero
es correcto y no comunica nada.

QUE HACE

Por cada CSV del paquete escribe un .xlsx con dos hojas:

  «Léeme»       qué es este anexo, cuántas filas tiene, qué significa cada
                columna, y a qué sección del manuscrito corresponde
  los datos     con las cabeceras en castellano, los códigos traducidos, la
                primera fila congelada y filtros puestos

QUE NO HACE

No toca los CSV. Se conservan tal cual, con sus nombres internos, porque son la
copia reproducible: quien quiera reejecutar el canal necesita las claves que el
canal usa, no una traducción. El .xlsx es para leer; el .csv es para verificar.
Un paquete que solo trae la versión bonita pierde la trazabilidad, y uno que
solo trae la cruda no la comunica.

Uso:
    python scripts/build_readable_annexes.py
"""
import csv
import datetime
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
PAQ = ROOT / "verificables revisión sistemática"
csv.field_size_limit(200_000_000)
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# Qué es cada anexo, en una frase que un revisor entienda sin abrir el fichero.
QUE_ES = {
    "S3_decisiones_etapa2_titulo": (
        "Toda decisión de cribado por TÍTULO, una por fila, con su motivo y su marca de tiempo.",
        "Sección 2.4. El registro es solo-anexar: una corrección es una fila nueva y ambas permanecen."),
    "S3_decisiones_etapa3_resumen": (
        "Toda decisión de cribado por RESUMEN, una por fila, con su motivo y su marca de tiempo.",
        "Sección 2.4. Solo llegan aquí los registros que superaron el cribado por título."),
    "S4_pre_extraccion_desde_resumen": (
        "La pre-extracción desde el resumen, que es la fuente de TODA cifra que el artículo publica.",
        "Sección 2.6. Cada fila va marcada PARCIAL: describe lo que el resumen declara, no el texto completo."),
    "S5_listado_184_estudios": (
        "Los 184 estudios tras agrupar los informes, con su situación y si se obtuvo el texto completo.",
        "Sección 3.1. La unidad es el estudio, no el informe: un mismo estudio puede tener varios informes."),
    "S8_recuperacion_texto_completo": (
        "Cómo se localizó cada texto completo, o por qué no se pudo.",
        "Sección 3.2. Explica el sesgo de recuperación que el artículo mide."),
    "S9_idioma_por_informe_y_clase_de_evidencia": (
        "El idioma de cada informe y de dónde se determinó, informe a informe.",
        "Sección 2.9. Sostiene la enmienda de idioma y su impacto medido."),
    "S10_idioma_verificado_sobre_texto_completo": (
        "El idioma comprobado sobre el PDF, no sobre lo que declaraba la base de datos.",
        "Sección 2.9. El nombre de la revista no determina el idioma del artículo."),
    "S21_solapamiento_de_pacientes": (
        "Cada par de estudios que podía describir a los mismos pacientes, con el veredicto "
        "tras leer los dos artículos y los pacientes que comparten, una clave por persona.",
        "Sección 4.4, limitación séptima. Un paciente que sale en tres pares tiene la misma "
        "clave en los tres y se cuenta una vez."),
    "S23_clase_de_resistencia_comprobada": (
        "La clase de resistencia de cada estudio con texto completo, comprobada contra el "
        "antibiograma del artículo, con la conclusión firmada por los dos autores.",
        "Sección 3.4. La columna «queda abierto» dice lo que la lectura no resolvió y no "
        "se cambió en la extracción."),
    "S24_comparativos_reextraidos": (
        "Los estudios con diseño comparativo, reextraídos: qué comparan, desde cuándo, "
        "hasta cuándo y si hay un contraste para P. aeruginosa.",
        "Sección 3.7. «Grupo sin fago» y «contraste para P. aeruginosa» son la "
        "codificación del texto firmado que va al lado."),
}

# Cabeceras. Lo que el canal llama X, un revisor lo entiende como Y.
ETIQUETA = {
    "record_id": "Identificador del informe", "id": "Identificador",
    "id_provisional": "Identificador del estudio", "estudio": "Estudio",
    "arm_id": "Brazo", "orden": "Orden", "verdict": "Decisión",
    "reason": "Motivo", "decided_by": "Quién decidió", "decided_at": "Cuándo",
    "corriente": "Corriente (base o registro)", "titulo": "Título",
    "revista": "Revista", "anio": "Año", "situacion": "Situación",
    "n_informes": "Informes del estudio", "diseno": "Diseño",
    "n_brazo": "Pacientes en el brazo", "procedencia": "País",
    "texto_completo": "¿Texto completo obtenido?", "clave_de_estudio": "Clave del estudio",
    "tipo": "Tipo de informe", "identificador_resuelto": "Identificador resuelto",
    "via": "Cómo se localizó", "nota": "Nota", "duplicado_interno": "Duplicado interno",
    "fuentes": "Fuentes que lo aportaron", "paginas": "Páginas", "palabras": "Palabras",
    "idioma_texto_completo": "Idioma del texto completo", "confianza": "Confianza",
    "evidencia": "Evidencia", "idioma_previo": "Idioma declarado antes",
    "coincide": "¿Coinciden?", "idioma_declarado": "Idioma declarado",
    "fuente_declarado": "De dónde sale el idioma declarado",
    "idioma_del_texto": "Idioma del texto", "clase_de_evidencia": "Clase de evidencia",
    "prueba": "Prueba", "estado": "Estado",
    "caracteres_de_resumen": "Caracteres del resumen",
    "titulo_alternativo": "Título alternativo",
    "pathogen_scope": "¿Solo P. aeruginosa o mixto?",
    "resistance_class": "Clase de resistencia",
    "resistance_class_source": "De dónde sale la clase",
    "dtr_status": "¿Cumple criterio DTR?", "route": "Vía de administración",
    "modality": "¿Fago solo o con antibiótico?", "study_design": "Diseño del estudio",
    "extraction_status": "Estado de la extracción", "n_arm": "Pacientes en el brazo",
    "en_corpus_actual": "¿Sigue en el corpus?", "study_id": "Estudio",
    "estado_final": "¿Quedó incluido al final?",
    "motivo_de_la_exclusion": "Por qué se excluyó",
    "campo": "Variable", "resolucion": "Valor acordado",
    "resuelto_por": "Quién lo resolvió", "fecha": "Fecha",
}

# Códigos internos -> castellano. Solo donde el código no se entiende solo.
VALOR = {
    "EXCLUDE": "excluido", "ADVANCE": "pasa a resumen", "FULLTEXT": "pasa a texto completo",
    "Pseudomonas-only": "solo P. aeruginosa",
    "mixed-pathogen-with-Pseudomonas-subgroup": "mixto, con subgrupo separable",
    "not-classifiable": "no clasificable", "below-MDR-threshold": "por debajo de MDR",
    "author-reported": "declarado por el autor",
    "independently-verified": "verificado de forma independiente",
    "yes": "sí", "no": "no", "not-derivable": "no derivable",
    "PARTIAL": "parcial", "COMPLETE": "completa",
    "extraible": "extraíble", "solo-registro": "solo ficha de registro",
    "solo-resumen": "solo resumen",
    "case report": "reporte de caso único", "case series": "serie de casos",
    "RCT": "ensayo aleatorizado", "non-randomised trial": "ensayo no aleatorizado",
    "prospective cohort": "cohorte prospectiva",
    "retrospective cohort": "cohorte retrospectiva",
    "base": "base bibliográfica", "registro": "registro de ensayos",
    # Vias y modalidades. Se traducen porque son de las columnas que un revisor
    # lee primero; los identificadores de ensayo y los DOI NO se tocan, que son
    # claves y traducirlas seria destruirlas.
    "phage+antibiotic combination": "fago + antibiótico",
    "phage monotherapy": "fago en monoterapia",
    "topical/local": "tópica o local", "inhaled/nebulized": "inhalada o nebulizada",
    "IV": "intravenosa", "oral": "oral", "intravenous": "intravenosa",
    "other": "otra", "unknown": "no consta", "none": "ninguno",
    "no-latino": "no latino",
    "INCLUIDO": "incluido", "EXCLUIDO DESPUES": "excluido después",
    "ADMITIDO": "admitido",
}


def main():
    if not PAQ.exists():
        print("no existe %s" % PAQ)
        return 1
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    AZUL = PatternFill("solid", fgColor="1F4E79")
    hechos = 0

    for csvf in sorted(PAQ.glob("S*.csv")):
        clave = csvf.stem
        with open(csvf, encoding="utf-8-sig", newline="") as fh:
            filas = list(csv.reader(fh))
        if not filas:
            continue
        cab, datos = filas[0], filas[1:]
        que, donde = QUE_ES.get(clave, ("Anexo del material suplementario.", ""))

        wb = openpyxl.Workbook()

        # ---- Léeme ---------------------------------------------------------
        p = wb.active
        p.title = "Léeme"
        p.column_dimensions["A"].width = 42
        p.column_dimensions["B"].width = 74
        lineas = [
            (clave.replace("_", " "), 15, True),
            ("", 11, False),
            ("Qué es", 12, True), (que, 11, False),
            ("", 11, False),
            ("Dónde se usa", 12, True), (donde, 11, False),
            ("", 11, False),
            ("Cuántas filas", 12, True), ("%d filas de datos" % len(datos), 11, False),
            ("", 11, False),
            ("Qué significa cada columna", 12, True),
        ]
        f = 1
        for txt, tam, neg in lineas:
            c = p.cell(row=f, column=1, value=txt)
            c.font = Font(size=tam, bold=neg, color="1F4E79" if neg and tam >= 12 else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            f += 1
        for col in cab:
            p.cell(row=f, column=1, value=ETIQUETA.get(col, col)).font = Font(bold=True, size=10)
            p.cell(row=f, column=2, value="columna interna: %s" % col).font = Font(size=10, color="808080")
            f += 1
        f += 1
        p.cell(row=f, column=1, value="El CSV con los nombres internos acompaña a este fichero: "
               "es la copia reproducible.").font = Font(size=10, italic=True)
        p.cell(row=f + 2, column=1,
               value="Generado el %s" % datetime.date.today().isoformat()).font = Font(size=9, color="808080")

        # ---- datos ---------------------------------------------------------
        h = wb.create_sheet("Datos")
        h.append([ETIQUETA.get(c, c) for c in cab])
        for c in h[1]:
            c.fill, c.font = AZUL, Font(color="FFFFFF", bold=True, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        h.row_dimensions[1].height = 30
        for fila in datos:
            h.append([VALOR.get((v or "").strip(), v) for v in fila])
        h.freeze_panes = "A2"
        h.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cab)), h.max_row)
        for i, col in enumerate(cab, start=1):
            ancho = 46 if col in ("titulo", "reason", "evidencia", "nota",
                                  "titulo_alternativo", "via", "motivo_de_la_exclusion") else 20
            h.column_dimensions[get_column_letter(i)].width = ancho

        sal = PAQ / (clave + ".xlsx")
        wb.save(sal)
        hechos += 1
        print("  %-48s %6d filas" % (sal.name[:48], len(datos)))

    print()
    print("%d anexos convertidos. Los CSV se conservan intactos." % hechos)
    return 0


if __name__ == "__main__":
    sys.exit(main())
