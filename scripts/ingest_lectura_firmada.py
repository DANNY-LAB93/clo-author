# -*- coding: utf-8 -*-
"""Ingiere la lectura firmada del 2026-09-30: resistencia, comparativos y solapamiento.

DE DONDE SALE

`build_trabajo_pendiente.py` escribio el 2026-09-23 los tres listados de
articulos que habia que abrir para los puntos 2, 7 y 15 del encargo de
correccion integral. Los dos autores los devolvieron rellenos y firmados
(«DANNY VALDIVIEZO/ NATALY TRELLES», fecha 2026-09-30) en
`revision_sistematica/lectura_pendiente/lectura_firmada_2026-09-30.xlsx`, que es
una copia intacta del cuaderno: el generador sobrescribe `quality_reports/`.

QUE HACE

1. Comprueba que cada fila de las tres hojas tiene firma, fecha y todas las
   columnas que habia que rellenar. Una fila a medias no se ingiere.
2. Escribe las tres hojas como CSV firmados, con el texto de los autores
   intacto y, al lado, la CODIFICACION que este guion hace de ese texto. La
   codificacion esta escrita aqui abajo, estudio por estudio, y cada una lleva
   una guarda: si la frase firmada no dice lo que la codificacion supone (no
   contiene «no verificable», o no nombra la clase que se le asigna), el
   guion se niega. Traducir texto libre a un vocabulario es una interpretacion,
   y tiene que poder revisarse fila a fila.
3. Anade a `correcciones_tras_texto_completo.csv` las correcciones de clase de
   resistencia, de su procedencia y del criterio DTR que la lectura deja
   INEQUIVOCAS para el brazo entero. Antes comprueba que el valor de hoy es el
   que se corrige. Es idempotente.

QUE NO HACE

· No aplica lo que la lectura deja abierto: un brazo con pacientes de clases
  distintas, un DTR que cambia segun el aislado o segun se lea la «I» de
  EUCAST, una clase declarada que la lectura refuta sin decir cual es la
  buena. Eso va a `PENDIENTE` y al cuaderno de firma, no a la extraccion.
· No excluye a nadie. Nueve estudios no tienen, segun la lectura, ningun
  paciente con P. aeruginosa tratado con fagos; excluir es un acto de autoria
  y se firma aparte (`make_firma_lectura.py`).
· No toca el manuscrito: las cifras salen de los escalares.

Uso:
    python scripts/ingest_lectura_firmada.py              # informe, no escribe
    python scripts/ingest_lectura_firmada.py --escribir
"""
import collections
import csv
import pathlib
import re
import sys

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
RS = ROOT / "revision_sistematica"
LP = RS / "lectura_pendiente"
FIRMADO = LP / "lectura_firmada_2026-09-30.xlsx"
EXTR = RS / "extraccion" / "extraccion_adjudicada.csv"
PROC = RS / "extraccion" / "extraccion_adjudicada_procedencia.csv"
CORR = RS / "extraccion" / "correcciones_tras_texto_completo.csv"
OUT_RES = LP / "lectura_resistencia_firmada.csv"
OUT_COM = LP / "lectura_comparativos_firmada.csv"
OUT_SOL = LP / "lectura_solapamiento_firmada.csv"

FECHA = "2026-09-30; 2026-09-30"
FIRMA_CORR = "DANNY VALDIVIEZO; NATALY TRELLES"
FIRMA_HOJA = "DANNY VALDIVIEZO/ NATALY TRELLES"

VER = "independently-verified"

# --------------------------------------------------------------- RESISTENCIA
# verificacion:
#   antibiograma  -- la clase se comprueba contra un antibiograma impreso
#                    (tabla, figura o tabla en el texto)
#   texto         -- el articulo no imprime el antibiograma, pero describe la
#                    sensibilidad con bastante precision para asignar la clase
#                    («sensible solo a colistina»)
#   no verificable-- no hay con que comprobar la clase declarada
#   no aplica     -- el estudio no documenta ningun paciente con P. aeruginosa
#                    tratado con fagos
# relacion: como queda la clase verificada frente a la que la extraccion trae.
# cambios: {(brazo, campo): valor nuevo}. Solo lo inequivoco para el brazo.
# pendiente: lo que la lectura deja abierto y vuelve a los autores.
# guarda: textos que la frase firmada TIENE que contener (en minuscula).
R = collections.OrderedDict()


def res(est, verif, clase="", relacion="", cambios=None, pendiente="", guarda=()):
    R[est] = dict(verif=verif, clase=clase, relacion=relacion,
                  cambios=cambios or {}, pendiente=pendiente, guarda=guarda)


NV = ("no verificable",)
res("EST-001", "no verificable", guarda=NV)
res("EST-002", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER}, guarda=("mdr (verificada",))
res("EST-003", "no verificable", guarda=NV)
res("EST-004", "antibiograma", "MDR",
    "coincide en el brazo A; la PDR del brazo B no es verificable",
    {("A", "resistance_class_source"): VER},
    pendiente="DTR: «sí en pacientes 1 y 3», pero el artículo no dice en qué "
              "brazo están; el brazo A sigue en «yes» y el B en «no».",
    guarda=("mdr confirmada", "pdr (2 pacientes) no verificable"))
res("EST-006", "antibiograma", "XDR", "coincide",
    {("A", "resistance_class_source"): VER}, guarda=("xdr confirmada",))
res("EST-007", "no verificable", guarda=NV)
res("EST-008", "no verificable", guarda=NV)
res("EST-009", "no aplica", guarda=("no aplica", "no se aisló p. aeruginosa"))
res("EST-010", "texto", "PDR", "coincide",
    {("A", "dtr_status"): "yes"}, guarda=("pdr (según el texto", "dtr = sí"))
res("EST-011", "texto", "MDR", "coincide",
    {("A", "dtr_status"): "yes"}, guarda=("mdr confirmada", "dtr = sí"))
res("EST-012", "no verificable", guarda=NV)
res("EST-013", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER}, guarda=("mdr confirmada",))
res("EST-014", "no verificable", guarda=NV)
res("EST-015", "antibiograma", "XDR", "menos grave que la declarada (PDR)",
    {("A", "resistance_class"): "XDR", ("A", "resistance_class_source"): VER},
    guarda=("xdr (no pdr)",))
res("EST-016", "no verificable", guarda=NV)
res("EST-019", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER},
    pendiente="DTR mixto: sí en el paciente HVR y no en LFV; el brazo (2 "
              "pacientes) sigue en «no».",
    guarda=("mdr confirmada", "dtr = sí en hvr"))
res("EST-021", "no verificable", guarda=NV)
res("EST-024", "no verificable", guarda=("pdr declarada, no verificable",))
res("EST-027", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER, ("A", "dtr_status"): "not-derivable"},
    guarda=("mdr confirmada", "dtr no derivable"))
res("EST-030", "antibiograma", "MDR", "no confirma la declarada (XDR)",
    {("A", "dtr_status"): "no"},
    guarda=("mdr confirmada; xdr declarada pero no verificable", "dtr = no"))
for e in ("EST-031", "EST-034", "EST-038", "EST-039"):
    res(e, "no verificable", guarda=NV)
res("EST-040", "texto", "XDR", "coincide", guarda=("xdr (según el texto",))
res("EST-042", "antibiograma", "", "refuta la declarada (PDR)",
    {("A", "dtr_status"): "no"},
    pendiente="Clase: la lectura refuta la PDR (paciente 1 por debajo del "
              "umbral; paciente 2 por debajo antes del fago y ≥ MDR después), "
              "pero no dice qué clase codificar para el brazo de 2 pacientes.",
    guarda=("pdr no confirmada; el antibiograma la contradice", "dtr = no"))
res("EST-043", "antibiograma", "XDR", "más grave que la declarada (MDR)",
    {("A", "resistance_class"): "XDR", ("A", "resistance_class_source"): VER},
    guarda=("xdr confirmada (mdr declarada se queda corta)",))
res("EST-044", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER}, guarda=("mdr confirmada",))
res("EST-046", "texto", "XDR", "la declarada no era clasificable",
    {("A", "resistance_class"): "XDR", ("A", "dtr_status"): "yes"},
    guarda=("xdr (según el texto", "dtr = sí"))
res("EST-047", "antibiograma", "XDR", "más grave que la declarada (MDR)",
    {("A", "resistance_class"): "XDR", ("A", "resistance_class_source"): VER},
    pendiente="DTR: sí en el aislado 2 y no en el aislado 1 del mismo "
              "paciente; sigue en «no».",
    guarda=("xdr confirmada (mdr declarada se queda corta)",))
res("EST-048", "no verificable",
    pendiente="La extracción dice «MDR, declarada por el autor», y la lectura "
              "dice que el artículo no declara ninguna clase (solo meropenem "
              "sensible). ¿Se corrige a «not-classifiable»?",
    guarda=("el artículo no declara mdr/xdr/pdr",))
for e in ("EST-049", "EST-050", "EST-051"):
    res(e, "no verificable", guarda=NV)
res("EST-053", "texto", "below-MDR-threshold", "la declarada no era clasificable",
    {("A", "resistance_class"): "below-MDR-threshold"},
    guarda=("below-mdr-threshold (verificada con el texto)",))
res("EST-055", "no aplica", guarda=("no aplica", "no hay p. aeruginosa"))
res("EST-057", "texto", "below-MDR-threshold", "la declarada no era clasificable",
    {("A", "resistance_class"): "below-MDR-threshold"},
    guarda=("below-mdr-threshold (verificada con el texto)",))
res("EST-058", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER},
    pendiente="DTR: «sí» si la «I» de EUCAST cuenta como no sensible (Kadri), "
              "«no» si se lee como sensible con exposición aumentada; sigue en «no».",
    guarda=("mdr confirmada",))
res("EST-061", "antibiograma", "MDR", "coincide (el paciente 1 es además XDR)",
    {("A", "resistance_class_source"): VER, ("A", "dtr_status"): "yes"},
    guarda=("xdr confirmada en el paciente 1", "paciente 2: mdr confirmada",
            "dtr = sí en ambos"))
res("EST-062", "texto", "XDR", "coincide", guarda=("xdr (según el texto",))
for e in ("EST-070", "EST-074", "EST-075"):
    res(e, "no verificable", guarda=NV)
res("EST-077", "antibiograma", "MDR",
    "coincide en el caso 2; los casos 1, 3 y 4 no son verificables",
    guarda=("mdr confirmada al límite", "no verificable"))
res("EST-080", "no verificable", guarda=NV)
res("EST-085", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER}, guarda=("mdr confirmada",))
res("EST-088", "antibiograma", "below-MDR-threshold",
    "la declarada no era clasificable",
    {("A", "resistance_class"): "below-MDR-threshold",
     ("A", "resistance_class_source"): VER},
    guarda=("below-mdr-threshold (verificada)",))
res("EST-091", "no verificable", guarda=NV)
res("EST-094", "antibiograma", "below-MDR-threshold", "coincide",
    {("A", "resistance_class_source"): VER},
    guarda=("below-mdr-threshold (verificada)",))
res("EST-095", "antibiograma", "MDR", "coincide (al límite)",
    {("A", "resistance_class_source"): VER, ("A", "dtr_status"): "not-derivable"},
    guarda=("mdr confirmada al límite", "dtr no derivable"))
res("EST-096", "no verificable", guarda=NV)
res("EST-105", "texto", "XDR", "la declarada no era clasificable",
    {("A", "resistance_class"): "XDR", ("A", "dtr_status"): "yes"},
    guarda=("xdr (según el texto", "dtr = sí"))
res("EST-106", "no verificable",
    pendiente="La lectura dice que lo documentado apunta a «below-MDR-threshold» "
              "(ceftazidima sensible durante años, ciprofloxacino intermedio) y "
              "la extracción trae «MDR, declarada por el autor». ¿Se corrige?",
    guarda=("mdr no verificable; lo documentado apunta a below-mdr-threshold",))
res("EST-108", "no verificable", guarda=NV)
for e in ("EST-116", "EST-117", "EST-121"):
    res(e, "no verificable", guarda=NV)
res("EST-124", "antibiograma", "",
    "XDR en un aislado previo; por debajo del umbral al tratar",
    pendiente="Clase y DTR: el aislado PA02 (abril de 2021) es XDR y DTR, pero "
              "los del día de la fagoterapia (junio de 2021) están por debajo "
              "del umbral. ¿Qué aislado manda?",
    guarda=("xdr confirmada en el aislado pa02", "below-mdr-threshold en el momento del tratamiento"))
res("EST-129", "no aplica", guarda=("no aplica", "no fue diana de los fagos"))
res("EST-132", "no aplica",
    guarda=("no aplica", "no se informa cuántos pacientes la tenían"))
res("EST-134", "no verificable", guarda=NV)
res("EST-135", "no aplica", guarda=("no aplica", "no tratada con fagos"))
res("EST-146", "no verificable", guarda=NV)
res("EST-152", "no aplica", guarda=("no aplica", "no hay ningún paciente con p. aeruginosa"))
res("EST-164", "antibiograma", "MDR", "coincide",
    {("A", "resistance_class_source"): VER, ("A", "dtr_status"): "yes"},
    guarda=("mdr confirmada en los 2 pacientes", "dtr = sí en ambos"))
res("EST-169", "antibiograma", "", "mixta: un paciente por debajo del umbral y otro XDR",
    pendiente="Clase y DTR: F12 está por debajo del umbral y M17 es XDR y DTR, "
              "en un solo brazo de 2 pacientes. Codificarlo exige partir el "
              "brazo, con el desenlace de cada paciente por separado.",
    guarda=("mixta, verificable con la fig. 4", "f12: below-mdr-threshold", "m17"))
res("EST-170", "no aplica", guarda=("no aplica", "k. pneumoniae"))
res("EST-181", "no aplica", guarda=("no aplica", "staphylococcus spp."))
res("EST-184", "no verificable", guarda=NV)
res("EST-208", "no aplica", guarda=("no aplica", "staphylococcus aureus sensible a meticilina"))

# ------------------------------------------------------------- COMPARATIVOS
# grupo_sin_fago: si hay un brazo que NO recibe fago.
# contraste_pa: si el articulo da un contraste entre brazos para P. aeruginosa.
C = collections.OrderedDict()
C["EST-004"] = ("no", "no", ("no: no hay grupo control",))
C["EST-008"] = ("si", "parcial", ("sí: placebo", "parcial:"))
C["EST-021"] = ("si", "si", ("sí: atención estándar", "sí: hr"))
C["EST-029"] = ("sin texto", "sin texto", ("sin texto completo",))
C["EST-055"] = ("si", "no: ningún paciente tenía P. aeruginosa",
                ("sí: grupo iv", "no para p. aeruginosa: ningún paciente"))
C["EST-070"] = ("no", "no", ("no. es un registro descriptivo",))
C["EST-096"] = ("no", "no", ("no hay brazo de control tratado",))
C["EST-108"] = ("no", "no", ("no hay grupo control sin fagos", "no para la pregunta de la revisión"))
C["EST-116"] = ("no", "no", ("no. es un análisis descriptivo",))
C["EST-132"] = ("si", "no: el artículo no desglosa por organismo",
                ("sí: grupo c", "no para p. aeruginosa: el artículo no desglosa"))
C["EST-146"] = ("no", "no", ("no: serie retrospectiva",))
C["EST-152"] = ("no", "no", ("no hay control sin fago",))
C["EST-157"] = ("sin texto", "sin texto", ("sin texto completo",))
C["EST-165"] = ("sin texto", "sin texto", ("sin texto completo",))

# ------------------------------------------------------------- SOLAPAMIENTO
# Una clave por PACIENTE, no por par: la nina del Berlin Heart sale en tres
# pares y es una sola persona. Las claves P1-P11 son las de
# `audita_pendientes.py`; P12-P20 las abre esta lectura.
K = {
    "P1": "niña de 10 años, Berlin Heart EXCOR",
    "P2": "niña de 7 años, infección osteoarticular",
    "P3": "varón de 25 años, osteomielitis craneal",
    "P4": "varón de 96 años, infección protésica",
    "P5": "infección espinal panresistente",
    "P6": "varón de 21 años, osteomielitis femoral, Riga",
    "P7": "septicemia sensible solo a colistina",
    "P8": "niño pequeño, trasplante hepático, Bruselas",
    "P9": "infección de la línea motriz de un LVAD (Racenis 2023)",
    "P10": "osteomielitis pélvica, Lovaina (Onsea P1)",
    "P11": "osteomielitis femoral, Lovaina (Onsea P2)",
    "P12": "varón de 46 años, osteomielitis",
    "P13": "varón de 69 años, mastoiditis",
    "P14": "mujer de 41 años, infección por silicona",
    "P15": "varón de 32 años, infección ósea",
    "P16": "varón de 67 años, trasplante pulmonar (UCSD)",
    "P17": "varón de 60 años, LVAD (UCSD)",
    "P18": "varón de 82 años, LVAD (UCSD)",
    "P19": "niño de 12 años, fascitis necrosante",
    "P20": "infección de injerto vascular (Blasco 2023)",
}
# (estudio_1, estudio_2): (claves, certeza). Sin claves = descartado.
S = {
    ("EST-009", "EST-038"): ((), ""), ("EST-009", "EST-044"): ((), ""),
    ("EST-009", "EST-080"): ((), ""), ("EST-009", "EST-134"): ((), ""),
    ("EST-012", "EST-015"): ((), ""), ("EST-016", "EST-164"): ((), ""),
    ("EST-038", "EST-044"): ((), ""), ("EST-038", "EST-080"): ((), ""),
    ("EST-038", "EST-134"): ((), ""), ("EST-044", "EST-080"): ((), ""),
    ("EST-051", "EST-108"): ((), "inferido por fechas"),
    ("EST-062", "EST-164"): ((), ""), ("EST-080", "EST-134"): ((), ""),
    ("EST-003", "EST-012"): (("P4",), ""),
    ("EST-003", "EST-015"): (("P3",), ""),
    ("EST-003", "EST-070"): (("P12", "P2", "P1", "P13", "P14", "P15"), ""),
    ("EST-003", "EST-077"): (("P1",), ""),
    ("EST-003", "EST-095"): (("P2",), ""),
    ("EST-070", "EST-077"): (("P1",), ""),
    ("EST-070", "EST-095"): (("P2",), ""),
    ("EST-034", "EST-049"): (("P16",), ""),
    ("EST-034", "EST-061"): (("P16",), ""),
    ("EST-049", "EST-061"): (("P16",), ""),
    ("EST-034", "EST-077"): (("P17", "P18"), ""),
    ("EST-062", "EST-108"): (("P8",), ""),
    ("EST-010", "EST-108"): (("P5",), ""),
    ("EST-088", "EST-108"): (("P19",), "muy probable"),
    ("EST-046", "EST-108"): (("P7",), ""),
    ("EST-108", "EST-164"): (("P10", "P11"), ""),
    ("EST-016", "EST-108"): (("P6",), ""),
    ("EST-002", "EST-108"): (("P20",), "por cita"),
    ("EST-108", "EST-124"): (("P9",), "por cita"),
}


def leer(p, enc="utf-8"):
    with open(p, encoding=enc, newline="") as fh:
        return list(csv.DictReader(fh))


def hoja(wb, nombre):
    it = wb[nombre].iter_rows(values_only=True)
    cab = [str(c) for c in next(it)]
    out = []
    for row in it:
        d = {k: ("" if v is None else v) for k, v in zip(cab, row)}
        if any(str(v).strip() for v in d.values()):
            out.append(d)
    return cab, out


def fecha(v):
    return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v)


def primera_cita(apoyo):
    """La primera frase entre comillas angulares, o la primera transcripcion."""
    m = re.search(r"«(.+?)»", apoyo or "", re.S)
    if m:
        return m.group(1)
    return (apoyo or "").split(" | ")[0]


def comprueba_firmas(filas, rellenar, donde, fallos):
    for f in filas:
        clave = f.get("study_id") or "%s + %s" % (f.get("estudio_1"), f.get("estudio_2"))
        if str(f.get("firmado_por", "")).strip().upper() != FIRMA_HOJA:
            fallos.append("%s %s: firma «%s»" % (donde, clave, f.get("firmado_por")))
        if fecha(f.get("fecha")) != "2026-09-30":
            fallos.append("%s %s: fecha %s" % (donde, clave, fecha(f.get("fecha"))))
        for c in rellenar:
            if not str(f.get(c, "")).strip():
                fallos.append("%s %s: «%s» vacío" % (donde, clave, c))


def main():
    escribir = "--escribir" in sys.argv
    wb = load_workbook(FIRMADO, data_only=True)
    fallos = []

    _, h2 = hoja(wb, "2 resistencia")
    _, h7 = hoja(wb, "7 comparativos")
    _, h15 = hoja(wb, "15 solapamiento")
    comprueba_firmas(h2, ["clase_verificada", "criterio_o_antibiograma_en_que_te_apoyas"],
                     "hoja 2", fallos)
    comprueba_firmas(h7, ["brazo_control_identificado", "tipo_y_unidad_de_asignacion",
                          "tiempo_cero", "tiempo_de_evaluacion_del_desenlace",
                          "cointervenciones_y_antibiotico_concomitante",
                          "perdidas_y_datos_faltantes", "estimando_reportado",
                          "hay_contraste_extraible"], "hoja 7", fallos)
    comprueba_firmas(h15, ["veredicto_tras_leer", "pacientes_duplicados",
                           "frase_en_que_te_apoyas"], "hoja 15", fallos)

    # --- la codificacion cubre exactamente lo firmado, y la frase la sostiene
    ids2 = [f["study_id"] for f in h2]
    if sorted(ids2) != sorted(R):
        fallos.append("hoja 2: estudios firmados %s; codificados %s"
                      % (sorted(set(ids2) ^ set(R)), "difieren"))
    for f in h2:
        c = R.get(f["study_id"])
        if not c:
            continue
        txt = str(f["clase_verificada"]).lower()
        for g in c["guarda"]:
            if g not in txt:
                fallos.append("hoja 2 %s: la codificación «%s» supone «%s» y la "
                              "frase firmada no lo dice" % (f["study_id"], c["verif"], g))
    ids7 = [f["study_id"] for f in h7]
    if sorted(ids7) != sorted(C):
        fallos.append("hoja 7: estudios firmados y codificados difieren: %s"
                      % sorted(set(ids7) ^ set(C)))
    for f in h7:
        g = C.get(f["study_id"])
        if not g:
            continue
        txt = (str(f["brazo_control_identificado"]) + " " +
               str(f["hay_contraste_extraible"])).lower()
        for x in g[2]:
            if x not in txt:
                fallos.append("hoja 7 %s: la codificación supone «%s»" % (f["study_id"], x))
    pares = {(f["estudio_1"], f["estudio_2"]) for f in h15}
    if pares != set(S):
        fallos.append("hoja 15: pares firmados y codificados difieren: %s"
                      % sorted(pares ^ set(S)))
    for f in h15:
        claves, _ = S.get((f["estudio_1"], f["estudio_2"]), ((), ""))
        ver = str(f["veredicto_tras_leer"]).lower()
        n = re.match(r"\s*(\d+)", str(f["pacientes_duplicados"]))
        if claves and not ver.startswith("solapan"):
            fallos.append("hoja 15 %s+%s: codificado como solapamiento y el veredicto "
                          "dice «%s»" % (f["estudio_1"], f["estudio_2"], ver[:30]))
        if not claves and not ver.startswith("no solapan"):
            fallos.append("hoja 15 %s+%s: codificado como descartado y el veredicto "
                          "dice «%s»" % (f["estudio_1"], f["estudio_2"], ver[:30]))
        # El numero de pacientes que firmaron, frente a las claves. EST-164
        # tiene 4 pacientes en EST-108, pero solo 2 con P. aeruginosa, que es
        # lo que el brazo de EST-164 extrae.
        if n and int(n.group(1)) != len(claves):
            if not ((f["estudio_1"], f["estudio_2"]) == ("EST-108", "EST-164")
                    and "2 con p. aeruginosa" in str(f["pacientes_duplicados"]).lower()):
                fallos.append("hoja 15 %s+%s: firmaron %s pacientes y hay %d claves"
                              % (f["estudio_1"], f["estudio_2"], n.group(1), len(claves)))

    # --- las correcciones, contra el valor de HOY
    adj = {(r["study_id"], r["arm_id"]): r for r in leer(EXTR)}
    proc = {(r["study_id"], r["arm_id"]): r for r in leer(PROC)}
    ya = leer(CORR) if CORR.exists() else []
    hechas = {(r["study_id"], r["arm_id"], r["campo"], r["valor_corregido"]) for r in ya}
    nuevas, iguales = [], 0
    por_est = {f["study_id"]: f for f in h2}
    for est, c in R.items():
        for (brazo, campo), nuevo in c["cambios"].items():
            fila = adj.get((est, brazo))
            if fila is None:
                fallos.append("%s %s no existe en la extracción" % (est, brazo))
                continue
            if (est, brazo, campo, nuevo) in hechas:
                continue
            hoy = fila[campo]
            if hoy == nuevo:
                iguales += 1
                continue
            f = por_est[est]
            nuevas.append({
                "study_id": est, "arm_id": brazo, "campo": campo,
                "valor_anterior": hoy,
                "procedencia_anterior": proc.get((est, brazo), {}).get(campo, ""),
                "valor_corregido": nuevo,
                "cita_literal": primera_cita(f["criterio_o_antibiograma_en_que_te_apoyas"]),
                "motivo": "Lectura firmada del 2026-09-30, hoja «2 resistencia»: "
                          + " ".join(str(f["clase_verificada"]).split()),
                "firmado_por": FIRMA_CORR,
                "fecha": FECHA,
            })

    # ------------------------------------------------------------ el informe
    cuenta = collections.Counter(c["verif"] for c in R.values())
    print("LECTURA FIRMADA DEL 2026-09-30")
    print("  hoja 2  · %d estudios: %s" % (len(R), ", ".join(
        "%s %d" % (k, cuenta[k]) for k in ("antibiograma", "texto",
                                           "no verificable", "no aplica"))))
    rel = collections.Counter(c["relacion"] for c in R.values() if c["relacion"])
    for k, v in rel.most_common():
        print("            %2d  %s" % (v, k))
    print("  hoja 7  · %d comparativos: grupo sin fago %d, contraste para P. aeruginosa %s"
          % (len(C), sum(1 for g in C.values() if g[0] == "si"),
             ", ".join(e for e, g in C.items() if g[1] in ("si", "parcial"))))
    conf = [p for p, (k, _) in S.items() if k]
    pac = {x for k, _ in S.values() for x in k}
    print("  hoja 15 · %d pares: %d con pacientes compartidos, %d descartados; "
          "%d pacientes distintos" % (len(S), len(conf), len(S) - len(conf), len(pac)))
    print()
    print("correcciones a la extracción: %d nuevas, %d ya al día"
          % (len(nuevas), iguales))
    for n in nuevas:
        print("  %s %s %-24s %s -> %s" % (n["study_id"], n["arm_id"], n["campo"],
                                        n["valor_anterior"], n["valor_corregido"]))
    pend = [(e, c["pendiente"]) for e, c in R.items() if c["pendiente"]]
    print()
    print("abierto, vuelve a los autores: %d" % len(pend))
    for e, p in pend:
        print("  %s  %s" % (e, p[:110]))
    if fallos:
        print()
        print("NO SE ESCRIBE NADA (%d problemas):" % len(fallos))
        for f in fallos:
            print("  · %s" % f)
        sys.exit(1)
    if not escribir:
        print()
        print("(informe; nada escrito. Con --escribir se aplica.)")
        return

    # ------------------------------------------------------------ escritura
    def escribe(ruta, filas):
        with open(ruta, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(filas[0].keys()))
            w.writeheader()
            w.writerows(filas)

    out = []
    for f in h2:
        c = R[f["study_id"]]
        out.append({
            "study_id": f["study_id"],
            "clase_en_la_extraccion_antes": f["clase_hoy"],
            "procedencia_antes": f["fuente_de_la_clase_hoy"],
            "verificacion": c["verif"],
            "clase_verificada_codificada": c["clase"],
            "relacion_con_la_declarada": c["relacion"],
            "cambia_en_la_extraccion": "; ".join(
                "%s %s -> %s" % (b, k, v) for (b, k), v in c["cambios"].items()),
            "queda_abierto": c["pendiente"],
            "clase_verificada_firmada": " ".join(str(f["clase_verificada"]).split()),
            "criterio_o_antibiograma_firmado": " ".join(
                str(f["criterio_o_antibiograma_en_que_te_apoyas"]).split()),
            "firmado_por": f["firmado_por"], "fecha": fecha(f["fecha"]),
        })
    escribe(OUT_RES, out)

    out = []
    for f in h7:
        g = C[f["study_id"]]
        d = {"study_id": f["study_id"], "disenos": f["disenos"],
             "grupo_de_comparacion_firmado_2026_09_22": f["grupo_de_comparacion_firmado"],
             "grupo_sin_fago": g[0], "contraste_para_p_aeruginosa": g[1]}
        for k in ("brazo_control_identificado", "tipo_y_unidad_de_asignacion",
                  "tiempo_cero", "tiempo_de_evaluacion_del_desenlace",
                  "cointervenciones_y_antibiotico_concomitante",
                  "perdidas_y_datos_faltantes", "estimando_reportado",
                  "hay_contraste_extraible"):
            d[k] = " ".join(str(f[k]).split())
        d["firmado_por"], d["fecha"] = f["firmado_por"], fecha(f["fecha"])
        out.append(d)
    escribe(OUT_COM, out)

    out = []
    for f in h15:
        claves, certeza = S[(f["estudio_1"], f["estudio_2"])]
        out.append({
            "estudio_1": f["estudio_1"], "estudio_2": f["estudio_2"],
            "par_nuevo": "si" if "añadido" in str(f["senal_que_lo_marco"]) else "no",
            "senal_que_lo_marco": f["senal_que_lo_marco"],
            "veredicto": ("SOLAPAMIENTO CONFIRMADO POR LECTURA" if claves
                          else "DESCARTADO POR LECTURA"),
            "certeza": certeza,
            "pacientes_compartidos": len(claves),
            "claves_de_paciente": "; ".join("%s %s" % (k, K[k]) for k in claves),
            "veredicto_firmado": " ".join(str(f["veredicto_tras_leer"]).split()),
            "pacientes_duplicados_firmado": str(f["pacientes_duplicados"]),
            "frase_firmada": " ".join(str(f["frase_en_que_te_apoyas"]).split()),
            "firmado_por": f["firmado_por"], "fecha": fecha(f["fecha"]),
        })
    escribe(OUT_SOL, out)

    if nuevas:
        cab = list(ya[0].keys()) if ya else list(nuevas[0].keys())
        with open(CORR, "a", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=cab)
            w.writerows(nuevas)
    print()
    print("escrito %s, %s, %s" % (OUT_RES.name, OUT_COM.name, OUT_SOL.name))
    print("añadidas %d correcciones a %s" % (len(nuevas), CORR.name))


if __name__ == "__main__":
    main()
