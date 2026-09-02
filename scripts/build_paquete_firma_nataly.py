"""Arma el paquete que N. Trelles necesita para firmar las cinco correcciones.

QUÉ MANDA, Y POR QUÉ ASÍ

Las cinco casillas corregidas no eran desacuerdos: tres las escribieron igual
los dos revisores y dos estaban cerradas por consenso. Reabrirlas exige que la
segunda revisora vea **la misma prueba** que las reabrió, no un resumen de
ella. Por eso cada fila lleva la cita literal del artículo y dónde está, y el
paquete incluye los PDF de los que se pueden comprobar.

EST-118 va sin PDF a propósito: ese es justamente su problema.

Salida:
    revision_sistematica/extraccion/firma_nataly/
        mensaje.md
        firma_correcciones_modalidad.xlsx
        pdf_para_revisar/EST-nnn.pdf

Uso:
    python scripts/build_paquete_firma_nataly.py
"""
import csv
import pathlib
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
EXTR = ROOT / "revision_sistematica" / "extraccion"
CORR = EXTR / "correcciones_tras_texto_completo.csv"
PDFS = ROOT / "revision_sistematica" / "textos_completos" / "pdf"
OUT = EXTR / "firma_nataly"

ES = {"phage monotherapy": "fago solo",
      "phage+antibiotic combination": "fago + antibiótico",
      "NA": "NA (no lo sabemos)"}

MENSAJE = """Nataly, necesito que revises cinco casillas y me digas si estás de
acuerdo.

Leyendo los textos completos uno por uno encontré cinco brazos donde el campo de
modalidad, el de fago solo o fago con antibiótico, no coincide con lo que dice el
artículo. Tres de esas casillas las pusimos igual las dos, y dos las habíamos
cerrado por consenso. Corregirlas es reabrir algo que ya dimos por cerrado, así
que no lo doy por hecho sin tu firma.

Van en firma_correcciones_modalidad.xlsx, una por fila. Cada una trae la frase
del artículo que la contradice, tal cual, y en qué parte está. Los PDF de los
cuatro que se pueden comprobar están en la carpeta pdf_para_revisar.

EST-118 es distinto y por eso no tiene PDF: es un ensayo aleatorizado del que
nunca conseguimos el texto completo, así que la modalidad la sacamos del
resumen. Propongo dejarlo en NA, que en nuestro esquema quiere decir que no lo
sabemos.

En cada fila pon sí o no en la columna de acuerdo. Si es no, escribe por qué en
la columna de al lado. Después tu nombre y la fecha.

Ya las apliqué al dataset con mi firma sola, y está marcado como tal: la columna
de procedencia dice "corregido contra el texto (una firma)" y esas cinco no
cuentan como doble lectura. Por eso la doble lectura bajó de 94,3 a 94,2 por
ciento. Cuando firmes pasan a consenso y vuelve a subir.

Con esto no se mueve ninguna cifra del manuscrito. Lo comprobé: el embudo da lo
mismo y las seis tablas se reconstruyen sin un solo cambio.
"""

CAB = ["Estudio", "Brazo", "Campo", "Lo que consta ahora", "De dónde venía",
       "Lo que propongo", "La frase del artículo", "Dónde está",
       "¿De acuerdo? (sí/no)", "Si es no, ¿por qué?", "Tu nombre", "Fecha"]
ANCHOS = [10, 7, 11, 20, 16, 20, 62, 22, 17, 34, 18, 12]

DONDE = {"EST-019": "Presentación del caso, y el propio resumen",
         "EST-085": "Resultados, y el pie de la figura 1D",
         "EST-118": "No hay texto completo que mirar",
         "EST-038": "Discusión",
         "EST-178": "Conclusiones"}


def main():
    OUT.mkdir(exist_ok=True)
    (OUT / "pdf_para_revisar").mkdir(exist_ok=True)
    filas = list(csv.DictReader(open(CORR, encoding="utf-8")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Correcciones"
    ws.append(CAB)
    cab = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="44546A")
    suyo = PatternFill("solid", fgColor="FFF2CC")
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill = cab, relleno
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = ANCHOS[i - 1]
    ws.row_dimensions[1].height = 32

    for r in filas:
        ws.append([r["study_id"], r["arm_id"], r["campo"],
                   ES.get(r["valor_anterior"], r["valor_anterior"]),
                   r["procedencia_anterior"],
                   ES.get(r["valor_corregido"], r["valor_corregido"]),
                   r["cita_literal"], DONDE.get(r["study_id"], ""),
                   "", "", "", ""])
    for fila in ws.iter_rows(min_row=2, max_row=len(filas) + 1):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in fila[8:]:
            c.fill = suyo
        ws.row_dimensions[fila[0].row].height = 74

    dv = DataValidation(type="list", formula1='"sí,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{len(filas) + 1}")
    ws.freeze_panes = "A2"

    n = wb.create_sheet("Qué es esto")
    for i, linea in enumerate(MENSAJE.strip().split("\n\n"), 1):
        n.cell(i, 1, linea.replace("\n", " ")).alignment = Alignment(
            wrap_text=True, vertical="top")
        n.row_dimensions[i].height = 58
    n.column_dimensions["A"].width = 118

    wb.save(OUT / "firma_correcciones_modalidad.xlsx")
    (OUT / "mensaje.md").write_text(MENSAJE, encoding="utf-8")

    copiados = []
    for r in filas:
        p = PDFS / f"{r['study_id']}.pdf"
        if p.exists():
            shutil.copy2(p, OUT / "pdf_para_revisar" / p.name)
            copiados.append(r["study_id"])
    sin = [r["study_id"] for r in filas if r["study_id"] not in copiados]

    print(f"paquete en {OUT}")
    print(f"  firma_correcciones_modalidad.xlsx  {len(filas)} correcciones")
    print(f"  mensaje.md")
    print(f"  pdf_para_revisar/  {len(copiados)} PDF -> {' '.join(copiados)}")
    if sin:
        print(f"  sin PDF (es su problema): {' '.join(sin)}")


if __name__ == "__main__":
    main()
