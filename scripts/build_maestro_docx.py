"""Convierte el informe extendido a .docx, con las citas resueltas.

QUE ES ESTE DOCUMENTO Y QUE NO. `paper/manuscrito_revision_sistematica.md` es
el manuscrito autoritativo: secciones numeradas, ~9 500 palabras, y el detalle
metodologico que el de la revista condensa. NO es lo que se envia a Journal of
Science and Research --alli va `manuscrito_JSR_final.docx`, y el sobre lleva un
solo manuscrito a proposito--. Este sirve para leerlo fuera del editor de texto:
para pasarselo a un director, a un comite o a un coautor.

LAS CITAS. El maestro escribe `[@Clave]` al estilo pandoc. Aqui se numeran por
orden de primera aparicion contra `Bibliography_base.bib` y se emite la lista de
referencias al final, en Vancouver. Se reutiliza la maquinaria de
`build_jsr_submission.py`, que ya lo hacia: no hay dos formas de numerar una
cita en este proyecto.

LAS TABLAS SON LAS DEL CANAL. Los pies del maestro nombran Tabla 1 a 7 y
Figura 1 y 2, y su correspondencia con los ficheros NO es la misma que en el
manuscrito de la revista --alli la Tabla 1 es la de criterios y aqui es la de
caracteristicas--. Por eso cada documento declara su propio mapa y no se
comparte uno.

Salida:
    paper/docx/informe_extendido.docx

Uso:
    python scripts/build_maestro_docx.py
"""
import pathlib
import re
import sys

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches, Mm
except ImportError:
    raise SystemExit("hace falta python-docx: python -m pip install python-docx")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_jsr_docx import escribe, tabla, tabla_csv, figura
from build_jsr_submission import lee_bib, numera_citas, referencia

ROOT = pathlib.Path(__file__).resolve().parent.parent
FUENTE = ROOT / "paper" / "manuscrito_revision_sistematica.md"
DEST = ROOT / "paper" / "docx"
TABLAS = ROOT / "paper" / "tablas"
FIGURAS = ROOT / "paper" / "figuras"
TNR = "Times New Roman"

# El mapa del MAESTRO. No coincide con el del manuscrito de la revista.
ADJUNTOS = {
    "Tabla 1.": TABLAS / "tabla_1_caracteristicas.csv",
    "Tabla 2.": TABLAS / "tabla_2_completitud.csv",
    "Tabla 3.": TABLAS / "tabla_3_sesgo_recuperacion.csv",
    "Tabla 4.": TABLAS / "tabla_4_exclusiones.csv",
    "Tabla 5.": TABLAS / "tabla_5_desenlaces.csv",
    "Tabla 6.": TABLAS / "tabla_6_embudo.csv",
    "Tabla 7.": TABLAS / "tabla_7_riesgo_sesgo.csv",
    "Figura 1.": FIGURAS / "figura_1_prisma.png",
    "Figura 2.": FIGURAS / "figura_2_composicion.png",
}

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def documento():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, lado, Inches(1))
    n = d.styles["Normal"]
    n.font.name, n.font.size = TNR, Pt(11)
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.4
    cab = s.header.paragraphs[0]
    cab.text = ("Informe extendido — NO es el manuscrito de envío. "
                "Fagoterapia en P. aeruginosa multirresistente")
    cab.runs[0].font.size, cab.runs[0].font.name = Pt(9), TNR
    cab.runs[0].italic = True
    return d


def main():
    if not FUENTE.exists():
        raise SystemExit("no encuentro %s" % FUENTE.name)
    bib = lee_bib(ROOT / "Bibliography_base.bib")
    orden = {}
    texto = numera_citas(FUENTE.read_text(encoding="utf-8"), orden)
    faltan = [c for c in orden if c not in bib]
    if faltan:
        print("AVISO: %d cita(s) sin entrada en el .bib: %s"
              % (len(faltan), ", ".join(sorted(faltan)[:6])))

    d = documento()
    titulos = 0
    lineas = texto.split("\n")
    i = 0
    while i < len(lineas):
        l = lineas[i].rstrip()
        if l.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].startswith("|"):
                bloque.append(lineas[i])
                i += 1
            if len(bloque) > 2:
                tabla(d, bloque)
            continue
        i += 1
        if not l or l == "---":
            continue
        if l.startswith("# "):
            titulos += 1
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6 if titulos == 1 else 12)
            escribe(p, l[2:], size=14, negrita=True)
            continue
        if l.startswith("### "):
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(5)
            escribe(p, l[4:], size=11, negrita=True)
            continue
        if l.startswith("## "):
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(7)
            escribe(p, l[3:], size=12, negrita=True)
            continue
        adj = next((k for k in ADJUNTOS if l.startswith("**" + k)), None)
        if adj:
            pie = d.add_paragraph()
            pie.paragraph_format.space_before = Pt(12)
            pie.paragraph_format.space_after = Pt(5)
            pie.paragraph_format.line_spacing = 1.0
            escribe(pie, l, size=9)
            ruta = ADJUNTOS[adj]
            if not ruta.exists():
                print("AVISO: falta %s, que la %s necesita" % (ruta.name, adj))
                continue
            (figura if ruta.suffix == ".png" else tabla_csv)(d, ruta)
            continue
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if re.match(r"^\d+\. ", l) or l.startswith(("- ", "• ")):
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            if l.startswith("- "):
                l = "• " + l[2:]
        escribe(p, l, size=11)

    # Referencias, en el orden en que el texto las cita
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    escribe(p, "Referencias", size=12, negrita=True)
    for clave, n in sorted(orden.items(), key=lambda kv: kv[1]):
        e = bib.get(clave)
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        escribe(p, "%d. %s" % (n, referencia(e) if e
                               else "[SIN ENTRADA EN EL .bib: %s]" % clave),
                size=10)

    DEST.mkdir(parents=True, exist_ok=True)
    salida = DEST / "informe_extendido.docx"
    try:
        d.save(salida)
    except PermissionError:
        salida = DEST / "informe_extendido_NUEVO.docx"
        d.save(salida)
        print("AVISO: estaba abierto en Word; escrito al lado como %s" % salida.name)
    print("escrito %s  (%d KB)" % (salida.relative_to(ROOT), salida.stat().st_size // 1024))
    print("  %d citas numeradas, %d tablas y figuras adjuntas"
          % (len(orden), sum(1 for k in ADJUNTOS if ADJUNTOS[k].exists())))
    print("  NO es el manuscrito de envío: ese es manuscrito_JSR_final.docx")


if __name__ == "__main__":
    main()
