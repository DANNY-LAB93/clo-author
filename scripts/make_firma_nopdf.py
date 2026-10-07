# -*- coding: utf-8 -*-
"""La tercera enmienda al protocolo: solo estudios con el PDF completo accesible.

El 2026-10-06 D. Valdiviezo pidio anadir un criterio de exclusion: «PDF que no
posean acceso completo». Preguntado por el alcance, respondio que se aplique a
los estudios con articulo cuyo texto no se obtuvo, a los resumenes de congreso
y a las fichas de registro sin articulo, y que el sesgo de recuperacion se
declare como efecto del criterio.

Excluir es un acto de autoria: este cuaderno lleva la enmienda, el codigo
nuevo (NOPDF), sus consecuencias medidas y la lista de los estudios que saca,
para que los dos lo firmen. `ingest_firma_nopdf.py` lo aplica con las dos
firmas.

SALIDA
    ~/Desktop/FIRMAR_texto_completo_2026-10-06.xlsx
"""
import collections
import csv
import json
import pathlib
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
RS = ROOT / "revision_sistematica"
QR = ROOT / "quality_reports"
CUADERNO = pathlib.Path.home() / "Desktop" / "FIRMAR_texto_completo_2026-10-06.xlsx"
OPCIONES = ["APROBAR la enmienda y excluir con NOPDF los estudios de la lista",
            "NO APROBAR"]
TIPO = {"extraible": "artículo publicado; el texto no se obtuvo",
        "solo-resumen": "solo resumen de congreso",
        "solo-registro": "solo ficha de registro, sin artículo"}


def leer(p, enc="utf-8-sig"):
    with open(p, encoding=enc, newline="") as fh:
        return list(csv.DictReader(fh))


def candidatos():
    """Los estudios del corpus sin PDF del articulo completo."""
    fuera = {r["study_id"] for r in leer(RS / "cribado" / "exclusiones_tras_texto_completo.csv")}
    pdfs = {p.stem for p in (RS / "textos_completos" / "pdf").iterdir()}
    grupos = collections.defaultdict(list)
    for g in leer(RS / "cribado" / "study_groups.csv"):
        grupos["EST-%03d" % int(g["estudio"])].append(g)
    pre = collections.defaultdict(set)
    for r in leer(RS / "extraccion" / "pre_extraccion_desde_resumen.csv"):
        if r.get("study_design"):
            pre[r["id_provisional"]].add(r["study_design"])
    brazos = collections.Counter(r["study_id"] for r in
                                 leer(RS / "extraccion" / "extraccion_adjudicada.csv", "utf-8"))
    out = []
    for e in sorted(grupos):
        if e in fuera or e in pdfs:
            continue
        g = next((x for x in grupos[e] if x["informe_para_extraer"] == "SI"), grupos[e][0])
        out.append({"estudio": e, "situacion": g["situacion"],
                    "tipo": TIPO.get(g["situacion"], g["situacion"]),
                    "titulo": " ".join((g["titulo"] or "").split()),
                    "revista_o_registro": g["revista"] or "",
                    "anio": g["anio"],
                    "diseno_segun_el_resumen": " / ".join(sorted(pre.get(e, []))) or "no consta",
                    "brazos_extraidos": brazos.get(e, 0)})
    return out


def main():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if CUADERNO.exists():
        wb = load_workbook(CUADERNO)
        if any(c.value not in (None, "") for ws in wb.worksheets for row in ws.iter_rows()
               for c in row if c.fill and str(c.fill.fgColor.rgb).endswith("FCE4D6")):
            raise SystemExit("NO SE ESCRIBE: el cuaderno ya tiene respuestas.")

    S = json.load(open(QR / "synthesis_scalars.json", encoding="utf-8"))
    O = json.load(open(QR / "outcome_scalars.json", encoding="utf-8"))
    lista = candidatos()
    c = collections.Counter(x["situacion"] for x in lista)
    disenos = collections.Counter()
    for x in lista:
        for d in x["diseno_segun_el_resumen"].split(" / "):
            disenos[d] += 1
    quedan = S["estudios"] - len(lista)
    brazos_fuera = sum(x["brazos_extraidos"] for x in lista)

    negrita, grande = Font(bold=True), Font(bold=True, size=13)
    blanca, azul = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    ojo = PatternFill("solid", fgColor="FCE4D6")
    arriba = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()
    h = wb.active
    h.title = "La enmienda"
    h.column_dimensions["A"].width = 34
    h.column_dimensions["B"].width = 100
    h.cell(row=1, column=1, value="Tercera enmienda: solo estudios con el PDF completo accesible").font = grande
    texto = [
        ("Criterio nuevo", "Se excluyen los estudios de los que no se dispone del PDF del artículo "
                           "completo con acceso: los que tienen artículo publicado pero cuyo texto no se "
                           "obtuvo, los que solo existen como resumen de congreso y los que solo existen "
                           "como ficha de registro de ensayo."),
        ("Código", "NOPDF — «%s». Duodécimo código del vocabulario de exclusiones." % (
            "Sin PDF del artículo completo con acceso: el texto no se obtuvo, o el estudio solo "
            "existe como resumen de congreso o ficha de registro")),
        ("Pidió", "D. Valdiviezo, 6 de octubre de 2026, con el alcance «a los 21 y a las 42 fichas» y "
                  "el sesgo de recuperación «declarado como efecto»."),
        ("Cuándo se adopta", "Con el cribado y la extracción terminados. En el manuscrito figura en la "
                             "Tabla 1 junto a los demás criterios, y en la sección 2.9 como la "
                             "ampliación de la segunda enmienda (NOREC, 2 de septiembre), con su fecha "
                             "y su efecto medido. No se presenta como parte del protocolo inicial: el "
                             "registro de decisiones que acompaña al manuscrito lo fecha."),
        ("Cuántos salen", "%d estudios: %d artículos sin texto, %d resúmenes de congreso y %d fichas de "
                          "registro. La lista está en la hoja siguiente." % (
                              len(lista), c["extraible"], c["solo-resumen"], c["solo-registro"])),
        ("Qué queda", "%d estudios, todos con el PDF leído. %d brazos de %d." % (
            quedan, O["brazos"] - brazos_fuera, O["brazos"])),
        ("Qué se pierde", "Diseños según el resumen de los que salen: %s. Los diseños comparativos "
                          "pasan de %d a %d y los ensayos aleatorizados de %d a %d." % (
                              ", ".join("%s %d" % (k, v) for k, v in disenos.most_common()),
                              S["estudios_comparativos"],
                              S["comparativos_si_se_excluye_lo_no_recuperado"],
                              S["ecas"], S["ecas_si_se_excluye_lo_no_recuperado"])),
        ("Qué deja de poder medirse", "La recuperación de texto completo pasa a ser del 100 % por "
                                      "construcción. El sesgo de recuperación se declara como efecto "
                                      "del criterio: qué diseños se pierden con él."),
        ("Lo que dice hoy el manuscrito", "Que NOREC no se generaliza precisamente por esto (sección "
                                          "2.9). Habrá que reescribir esa sección, la limitación "
                                          "correspondiente y la corriente de registros de PRISMA, cuyas "
                                          "fichas salen todas."),
    ]
    f = 3
    for k, v in texto:
        h.cell(row=f, column=1, value=k).font = negrita
        c2 = h.cell(row=f, column=2, value=v)
        c2.alignment = arriba
        h.row_dimensions[f].height = 48
        f += 1
    f += 1
    h.cell(row=f, column=1, value="Vuestra decisión").font = negrita
    cel = h.cell(row=f, column=2)
    cel.fill = ojo
    dv = DataValidation(type="list", allow_blank=True, formula1='"%s"' % ",".join(OPCIONES))
    h.add_data_validation(dv)
    dv.add(cel)
    h.cell(row=f + 1, column=1, value="Por qué").font = negrita
    h.cell(row=f + 1, column=2).fill = ojo

    s = wb.create_sheet("Los %d" % len(lista))
    cab = ["Estudio", "Qué es", "Título", "Revista o registro", "Año", "Diseño según el resumen",
           "Brazos extraídos"]
    for j, (cc, w) in enumerate(zip(cab, [10, 30, 70, 34, 7, 24, 10]), start=1):
        x = s.cell(row=1, column=j, value=cc)
        x.font, x.fill, x.alignment = blanca, azul, arriba
        s.column_dimensions[x.column_letter].width = w
    for i, x in enumerate(lista, start=2):
        for j, k in enumerate(["estudio", "tipo", "titulo", "revista_o_registro", "anio",
                               "diseno_segun_el_resumen", "brazos_extraidos"], start=1):
            s.cell(row=i, column=j, value=x[k]).alignment = arriba
    s.freeze_panes = "A2"
    s.auto_filter.ref = s.dimensions

    fi = wb.create_sheet("Firma")
    fi.column_dimensions["A"].width = 46
    fi.column_dimensions["B"].width = 54
    fi.cell(row=1, column=1, value="Firma de los dos autores").font = negrita
    for i, k in enumerate(["Revisor 1 (nombre completo)", "Revisor 2 (nombre completo)",
                           "Fecha (AAAA-MM-DD)",
                           "¿Habéis leído la lista y las consecuencias?"], start=4):
        fi.cell(row=i, column=1, value=k).font = negrita
        fi.cell(row=i, column=2, value="").fill = ojo
    wb.save(CUADERNO)
    print("cuaderno: %s" % CUADERNO)
    print("  salen %d: %s" % (len(lista), dict(c)))
    print("  quedan %d estudios; brazos %d -> %d" % (quedan, O["brazos"], O["brazos"] - brazos_fuera))
    print("  diseños que salen:", dict(disenos))


if __name__ == "__main__":
    main()
