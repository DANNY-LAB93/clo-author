"""Construye la tabla de criterios de inclusion y exclusion, a dos columnas.

D.V. pidio dos columnas enfrentadas --"Criterios de inclusion" y "Criterios de
exclusion"-- en lugar de la tabla por dominios (Poblacion, Intervencion,
Comparador...) que llevaba el manuscrito, y corrigio una primera version que
encabezaba cada fila con su dominio PICO: los criterios describen QUE ARTICULOS
entran y cuales no, no la pregunta PICO. Cada fila es ahora una frase sobre el
documento.

FICHAS DE REGISTRO Y RESUMENES DE CONGRESO. La tabla anterior no decia nada de
ellos, y sin embargo 60 de los 155 estudios del corpus existen unicamente como
ficha de registro de ensayo y 3 solo como resumen de congreso: 63 de 155,
admitidos de hecho pero no declarados en ningun criterio. Se declara ahora.

No confundir esa fila con el codigo PRO. PRO excluye ARTICULOS DE PROTOCOLO
publicados en revista (EST-035, EST-052, EST-083, EST-122; el de CYPHY tiene 223
ocurrencias de "will be"), que son un informe que anuncia lo que se hara. Una
ficha de registro es otra cosa: es el rastro del ensayo en la corriente de
registros que PRISMA 2020 obliga a separar, y se contabiliza como estudio
identificado aunque no aporte resultados. Las dos filas lo dicen expresamente,
porque juntas parecen contradecirse.

LOS CODIGOS NO SE IMPRIMEN. Cada motivo de exclusion lleva en el codigo fuente su
codigo del vocabulario cerrado (`scripts/exclusion_codes.py`) --ORG, VET, LAB...
el mismo con el que se cuentan las exclusiones en el diagrama PRISMA y en el
anexo S16--, pero la tabla sale sin ellos: D.V. los quito el 2026-09-03. Se
conserva el par motivo-codigo aqui porque es lo unico que ata cada fila de esta
tabla con su recuento; si alguien vuelve a quererlos visibles, basta con volver a
concatenarlos en `filas()`.

La ventana de publicacion no se teclea: sale de los escalares.

Salidas:
    paper/tablas/tabla_criterios_inclusion_exclusion.md
    paper/tablas/tabla_criterios_inclusion_exclusion.csv
    ~/Desktop/Envio_JSR_Fagoterapia_Pseudomonas/tabla_criterios_inclusion_exclusion.docx  -- el cuadro suelto
    paper/manuscrito_JSR_final.md   -- se le reemplaza la Tabla 1

Uso:
    python scripts/build_criteria_table.py
"""
import csv
import json
import pathlib
import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
MANUSCRITO = ROOT / "paper" / "manuscrito_JSR_final.md"
TABLAS = ROOT / "paper" / "tablas"
BASE = "tabla_criterios_inclusion_exclusion"


def ventana():
    d = json.loads((ROOT / "quality_reports" / "synthesis_scalars.json")
                   .read_text(encoding="utf-8"))
    d = d.get("escalares", d)
    return d["anio_min"], d["anio_max"]


def filas():
    a, b = ventana()
    inclusion = [
        "Artículos originales publicados entre %s y %s, donde la interfaz de "
        "búsqueda permitió aplicar el filtro." % (a, b),
        "Artículos redactados en español o en inglés.",
        "Artículos nacionales e internacionales, incluida la literatura regional "
        "indexada en BVS y SciELO.",
        "Artículos originales que reporten el uso terapéutico de bacteriófagos "
        "líticos en pacientes humanos con infección por *Pseudomonas aeruginosa* "
        "multirresistente (MDR), extremadamente resistente (XDR) o panresistente "
        "(PDR), según Magiorakos et al., o descrita en términos que impliquen esas "
        "categorías.",
        "Artículos en los que el bacteriófago se administre solo o combinado con "
        "antimicrobianos, por cualquier vía.",
        "Artículos sobre varios patógenos, cuando los datos de *P. aeruginosa* "
        "pueden separarse del resto.",
        "Ensayos clínicos aleatorizados y no aleatorizados, estudios de cohorte, "
        "series de casos y reportes de caso que describan pacientes tratados.",
        "Artículos que informen al menos uno de los desenlaces de interés: éxito "
        "clínico, erradicación microbiológica, mortalidad, eventos adversos o "
        "emergencia de resistencia al fago.",
        "Fichas de registro de ensayos clínicos y resúmenes de congreso que cumplan "
        "lo anterior, contabilizados como estudios identificados aunque no aporten "
        "resultados publicados.",
    ]
    exclusion = [
        ("Artículos publicados fuera del período de estudio.", None),
        ("Estudios publicados en un idioma diferente del español o el inglés.", "IDI"),
        ("Estudios sobre un organismo distinto de *P. aeruginosa* cuando sus datos "
         "no pueden separarse del resto.", "ORG"),
        ("Estudios en animales o sobre aislados veterinarios.", "VET"),
        ("Estudios de laboratorio, preclínicos o de modelización, sin pacientes "
         "tratados.", "LAB"),
        ("Artículos de revisión bibliográfica, editoriales y comentarios sin datos "
         "primarios propios.", "REV"),
        ("Revisiones sistemáticas y revisiones de alcance.", "SEC"),
        ("Artículos que no evalúan fagoterapia en pacientes: encuestas, estudios "
         "epidemiológicos, notas de prensa u otra terapia.", "OFF"),
        ("Estudios en los que la intervención no es un bacteriófago, como "
         "endolisinas u otros derivados administrados sin la partícula viral.", "INT"),
        ("Artículos de protocolo publicados en revista, que declaran lo que se hará "
         "sin presentar resultados.", "PRO"),
        ("Estudios cuyo texto completo no se pudo recuperar, de modo que los "
         "criterios no pudieron verificarse contra el artículo.", "NOREC"),
    ]
    # El codigo NO se imprime: D.V. lo quito de la tabla el 2026-09-03. Se
    # conserva junto a cada motivo porque es lo que ata esta fila con el recuento
    # del diagrama PRISMA y con el anexo de exclusiones, y sin el par aqui esa
    # correspondencia solo vive en la cabeza de quien escribio la tabla.
    exc = [t for t, _ in exclusion]
    # Se enfrentan fila a fila; la columna corta se rellena en blanco, como en
    # la tabla que sirvio de modelo.
    n = max(len(inclusion), len(exc))
    inc = inclusion + [""] * (n - len(inclusion))
    exc = exc + [""] * (n - len(exc))
    return list(zip(inc, exc))


CABECERA = ("Criterios de inclusión", "Criterios de exclusión")
# El manuscrito titula su seccion "Criterios de elegibilidad", que es como los
# llama el item 5 de PRISMA 2020, y el pie decia "Criterios de inclusion y
# exclusion". Es lo mismo con dos nombres, y la seccion remitia a una tabla que
# se llamaba de otra manera. El pie los ata; las columnas siguen diciendo
# inclusion y exclusion, que es el formato que pidio D.V.
PIE = ("**Tabla 1.** Criterios de elegibilidad: inclusión y exclusión de los "
       "artículos.")
TNR = "Times New Roman"
ESCRITORIO = pathlib.Path.home() / "Desktop" / "Envio_JSR_Fagoterapia_Pseudomonas"


def escribe(p, txt, size, negrita=False):
    """Resuelve *cursiva* y **negrita** en linea; el marcado no llega al .docx."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", txt):
        if not trozo:
            continue
        neg = trozo.startswith("**")
        r = p.add_run(trozo.strip("*"))
        r.font.name = TNR
        r.font.size = Pt(size)
        r.bold = negrita or neg
        r.italic = trozo.startswith("*") and not neg


def docx_suelto(f, destino):
    """El cuadro solo, en su propio .docx, para pegarlo donde haga falta."""
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, lado, Mm(25.4))
    n = d.styles["Normal"]
    n.font.name, n.font.size = TNR, Pt(11)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.line_spacing = 1.0

    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(CABECERA):
        cel = t.rows[0].cells[i]
        cel.text = ""
        cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        escribe(cel.paragraphs[0], c, 11, negrita=True)
    for inc, exc in f:
        cs = t.add_row().cells
        for cel, txt in zip(cs, (inc, exc)):
            cel.text = ""
            cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            escribe(cel.paragraphs[0], txt, 11)

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribe(p, PIE, 10)

    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        d.save(destino)
    except PermissionError:
        destino = destino.with_name(destino.stem + "_NUEVO.docx")
        d.save(destino)
        print("AVISO: estaba abierto en Word; se escribe al lado.")
    print("escrito %s" % destino)


def main():
    f = filas()
    TABLAS.mkdir(parents=True, exist_ok=True)

    md = ["| %s | %s |" % CABECERA, "|---|---|"]
    md += ["| %s | %s |" % (i, e) for i, e in f]
    (TABLAS / (BASE + ".md")).write_text("\n".join(md) + "\n", encoding="utf-8")

    with open(TABLAS / (BASE + ".csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(CABECERA)
        for i, e in f:
            w.writerow([i.replace("**", "").replace("*", ""),
                        e.replace("**", "").replace("*", "")])
    print("escritas %s.md y %s.csv  (%d filas)" % (BASE, BASE, len(f)))

    docx_suelto(f, ESCRITORIO / (BASE + ".docx"))

    # Sustituye la Tabla 1 dentro del manuscrito, pie incluido.
    #
    # ESTO BORRO LA SECCION RESULTADOS ENTERA. La version anterior buscaba
    # «**Tabla 1.», buscaba «**Tabla 2.» y reemplazaba TODO lo que hubiera en
    # medio, dando por hecho que las dos iban seguidas. Lo estaban cuando se
    # escribio esto. Hoy la Tabla 1 vive en METODOLOGIA y la Tabla 2 en
    # RESULTADOS, asi que «en medio» son setenta lineas: el final de Metodos,
    # el encabezado «## RESULTADOS» y el principio de los resultados. El guion
    # no fallaba: imprimia «Tabla 1 del manuscrito reemplazada» y se quedaba
    # tan ancho.
    #
    # Ahora el reemplazo se ata al bloque de tabla que sigue al pie --las
    # lineas que empiezan por «|»-- y nada mas.
    lineas = MANUSCRITO.read_text(encoding="utf-8").split("\n")
    ini = next((i for i, l in enumerate(lineas) if l.startswith("**Tabla 1.")), -1)
    if ini < 0:
        raise SystemExit("no encuentro el pie de la Tabla 1 en el manuscrito")
    fin = ini + 1
    while fin < len(lineas) and not lineas[fin].startswith("|"):
        if lineas[fin].strip():          # algo que no es la tabla ni un blanco
            raise SystemExit("tras el pie de la Tabla 1 no viene una tabla, sino: %r"
                             % lineas[fin][:60])
        fin += 1
    while fin < len(lineas) and lineas[fin].startswith("|"):
        fin += 1
    # RED DE SEGURIDAD. El tramo que se va a tirar tiene que ser una tabla
    # entera y nada mas: cabecera, separador y al menos una fila. Si no lo es,
    # los limites estan mal y lo que toca es morir, no escribir. Probado
    # metiendo un «## SECCION» dentro del bloque: sin esto el guion reemplazaba
    # media tabla y dejaba la otra mitad huerfana, con el mismo aire de exito.
    # OJO con el nombre: `filas` ya es la funcion que arma los criterios, unas
    # lineas mas arriba en este mismo `main()`. Llamar igual a la lista la
    # tapaba y reventaba con UnboundLocalError antes de llegar aqui.
    cuadricula = [l for l in lineas[ini:fin] if l.startswith("|")]
    separador = [l for l in cuadricula if set(l.replace("|", "").strip()) <= set("-: ")]
    if len(cuadricula) < 3 or len(separador) != 1:
        raise SystemExit(
            "el bloque de la Tabla 1 no parece una tabla completa: %d filas y "
            "%d separadores. No se toca el manuscrito."
            % (len(cuadricula), len(separador)))
    tragado = [l for l in lineas[ini:fin] if l.startswith("#")]
    if tragado:
        raise SystemExit("el reemplazo se comeria %d encabezado(s): %s"
                         % (len(tragado), tragado[:3]))
    # El pie sale de PIE, que es el que ya usa el cuadro suelto. Estaba
    # tecleado aparte aqui y decia otra cosa: el .docx suelto llevaba
    # «Criterios de elegibilidad: inclusion y exclusion» --la redaccion que ata
    # el pie con el nombre de la seccion, y que alguien corrigio a proposito--
    # y el manuscrito seguia con la vieja. La misma tabla con dos pies, y uno
    # de los dos viajando suelto en el sobre.
    nuevas = lineas[:ini] + [PIE, ""] + md + lineas[fin:]
    # NO SE REESCRIBE SI NO CAMBIA NADA. Tocar el fichero mueve su fecha, y
    # el guardian del .zip compara cada documento derivado contra la fecha
    # del manuscrito: reescribirlo por costumbre dejaba el .docx de la tabla
    # «caducado» un segundo despues de generarlo, y el sobre no se podia
    # comprimir sin volver a correrlo todo en circulo.
    if nuevas == lineas:
        print("Tabla 1 del manuscrito ya estaba al dia; no se reescribe")
        return
    MANUSCRITO.write_text("\n".join(nuevas), encoding="utf-8")
    print("Tabla 1 del manuscrito reemplazada (%d lineas dentro, %d fuera)"
          % (fin - ini, len(lineas)))


if __name__ == "__main__":
    main()
