"""Recoge la firma de la segunda revisora sobre las correcciones, o la rechaza.

QUÉ COMPRUEBA ANTES DE INGERIR

El cuaderno vuelve de otra persona y de otro ordenador, así que no se toma al
pie de la letra. Antes de tocar nada se verifica que cada fila devuelta sigue
siendo la fila que se mandó: mismo estudio, mismo brazo, mismo campo, mismo
valor propuesto y misma cita. Si algo de eso cambió, la fila NO se ingiere y se
informa: una firma sobre una propuesta distinta de la que se envió no es una
firma sobre esta corrección.

Después:

  · «sí» -> se añade su firma junto a la del primer revisor;
  · «no» -> NO se ingiere, y se informa con el motivo que escribió, porque eso
    reabre la casilla en vez de cerrarla;
  · en blanco -> no se ingiere. No contestar no es aceptar.

Uso:
    python scripts/ingest_firma_correcciones.py            # informe
    python scripts/ingest_firma_correcciones.py --escribir
"""
import csv
import pathlib
import sys

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
EXTR = ROOT / "revision_sistematica" / "extraccion"
CORR = EXTR / "correcciones_tras_texto_completo.csv"
FIRMADO = EXTR / "firma_nataly" / "firma_correcciones_modalidad.xlsx"

SI = {"sí", "si", "s", "yes", "de acuerdo"}
NO = {"no", "n"}
# Las mismas etiquetas legibles con que se construyó el cuaderno.
ES = {"phage monotherapy": "fago solo",
      "phage+antibiotic combination": "fago + antibiótico",
      "NA": "NA (no lo sabemos)"}


def main():
    escribir = "--escribir" in sys.argv
    filas = list(csv.DictReader(open(CORR, encoding="utf-8")))
    cols = list(filas[0].keys())
    idx = {(r["study_id"], r["arm_id"], r["campo"]): r for r in filas}

    ws = load_workbook(FIRMADO)["Correcciones"]
    firmadas, rechazadas, mudas, alteradas = [], [], [], []
    for f in ws.iter_rows(min_row=2, values_only=True):
        est, brazo, campo, _, _, propuesto, cita = f[0], f[1], f[2], f[3], f[4], f[5], f[6]
        acuerdo = (str(f[8] or "").strip().lower())
        porque = (str(f[9] or "").strip())
        quien = (str(f[10] or "").strip())
        cuando = f[11]
        r = idx.get((est, brazo, campo))
        if r is None:
            alteradas.append((est, brazo, campo, "esa fila no se envió"))
            continue
        # La propuesta devuelta tiene que ser la que se mandó.
        esperado = ES.get(r["valor_corregido"], r["valor_corregido"])
        if (propuesto or "").strip() != esperado:
            alteradas.append((est, brazo, campo,
                              f"el valor propuesto cambió: {propuesto!r}"))
            continue
        if (cita or "").strip() != r["cita_literal"].strip():
            alteradas.append((est, brazo, campo, "la cita del artículo cambió"))
            continue
        if acuerdo in NO:
            rechazadas.append((est, porque or "(sin motivo escrito)"))
            continue
        if acuerdo not in SI or not quien:
            mudas.append((est, f"acuerdo={acuerdo!r} firma={quien!r}"))
            continue
        fecha = (cuando.date().isoformat() if hasattr(cuando, "date")
                 else str(cuando or "").strip())
        firmadas.append((r, quien, fecha))

    print(f"cuaderno devuelto: {ws.max_row - 1} filas\n")
    for est, quien, fecha in ((x[0]["study_id"], x[1], x[2]) for x in firmadas):
        print(f"  FIRMADA   {est}  por {quien}  {fecha}")
    for est, motivo in rechazadas:
        print(f"  RECHAZADA {est}  -> {motivo}")
    for est, det in mudas:
        print(f"  SIN FIRMA {est}  {det}")
    for est, brazo, campo, det in alteradas:
        print(f"  ALTERADA  {est} {brazo} {campo}: {det}  NO SE INGIERE")

    if rechazadas:
        print("\n  Una corrección rechazada NO se revierte sola: reabre la "
              "casilla y hay que decidirla entre los dos.")
    if not escribir:
        print("\n(informe. Añade --escribir para guardar las firmas)")
        return

    for r, quien, fecha in firmadas:
        if quien.lower() not in r["firmado_por"].lower():
            r["firmado_por"] = f"{r['firmado_por']}; {quien}"
        r["fecha"] = f"{r['fecha']}; {fecha}"
    with open(CORR, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, cols)
        w.writeheader()
        w.writerows(filas)
    print(f"\n{len(firmadas)} firmas guardadas en {CORR.name}.")
    print("Ahora: python scripts/build_adjudicated_dataset.py")


if __name__ == "__main__":
    main()
