"""Construye el cuaderno UNICO de consenso del riesgo de sesgo.

POR QUE UNO Y NO DOS. Los dos cuadernos de los comparativos no son dos lecturas
independientes: la columna de texto libre trae el mismo texto, caracter por
caracter, en los once estudios anotados, y eso solo pasa si uno salio del otro.
Con esa base no hay kappa que reportar --medirla seria medir cuantas celdas se
tocaron despues de copiar-- ni se puede declarar evaluacion independiente.

Lo que SI hay es una evaluacion hecha con el articulo delante, con frases
citadas del texto. Eso es una evaluacion por consenso, y como tal se declara.
Este guion la convierte en lo que es: UN cuaderno, con UN juicio por dominio,
firmado por los dos autores.

QUE SE PRECARGA Y QUE NO. Donde los dos ficheros coinciden, el valor entra ya
puesto: es el acuerdo que ya existia y no hay nada que decidir. Donde difieren
--13 celdas, todas de ROBINS-I-- la celda queda VACIA y en amarillo fuerte, con
un comentario que dice que habia en cada fichero. Esas trece son las unicas
decisiones que quedan, y tienen que tomarlas los dos autores mirando el
articulo, no heredarse de un fichero que ya no significa lo que parecia.

LA FIRMA NO ES DECORACION. `ingest_rob.py --consenso` no ingiere el cuaderno sin
los dos nombres y la fecha en la hoja «Firma». Sin eso, "consenso" es una
palabra en el manuscrito sin nada detras.

Salida:
    revision_sistematica/riesgo_sesgo/riesgo_sesgo_comparativos_consenso.xlsx

Uso:
    python scripts/make_rob_consensus.py
"""
import pathlib
import sys

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    raise SystemExit("hace falta openpyxl: python -m pip install openpyxl")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from compare_rob import lee_libro, notas
from make_rob_forms import CONTEXTO, corpus, evidencia
from rob_instruments import COMPARATIVOS, NO_EVALUABLE, PENDIENTE, SIMPLIFICADOS

ROOT = pathlib.Path(__file__).resolve().parent.parent
DEST = ROOT / "revision_sistematica" / "riesgo_sesgo"
A = DEST / "riesgo_sesgo_comparativos_danny_valdiviezo.xlsx"
B = DEST / "riesgo_sesgo_comparativos_nataly_trelles.xlsx"
SALIDA = DEST / "riesgo_sesgo_comparativos_consenso.xlsx"

CAB = PatternFill("solid", fgColor="D9D9D9")
CABF = Font(bold=True, size=9)
FIJO = PatternFill("solid", fgColor="EDEDED")     # contexto, no se toca
ACORDADO = PatternFill("solid", fgColor="E2EFDA")  # ya coincidian, revisable
DECIDIR = PatternFill("solid", fgColor="FFD966")   # discrepan: hay que decidir
FIRMA = PatternFill("solid", fgColor="FFF2CC")

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def hoja_firma(wb, n_decidir):
    ws = wb.create_sheet("Firma")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46
    ws["A1"] = "Evaluación del riesgo de sesgo — por consenso"
    ws["A1"].font = Font(bold=True, size=13)
    texto = (
        "Este cuaderno NO es una evaluación por duplicado independiente, y el "
        "manuscrito no la declara como tal. Es una evaluación única, acordada "
        "entre los dos autores con el artículo delante.\n\n"
        "Las celdas verdes traen el juicio en que los dos ficheros previos ya "
        "coincidían: revísenlas, pero no hay nada que decidir en ellas.\n\n"
        "Las %d celdas AMARILLAS están vacías porque los dos ficheros daban "
        "valores distintos. Son las únicas decisiones que quedan. El comentario "
        "de cada una dice qué había en cada fichero; decidan mirando el "
        "artículo, no el comentario.\n\n"
        "Sin los dos nombres y la fecha de abajo, `ingest_rob.py --consenso` no "
        "ingiere nada." % n_decidir)
    ws["A3"] = texto
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:B3")
    ws.row_dimensions[3].height = 190
    for i, (etq, ayuda) in enumerate((
            ("Evaluado y acordado por (1):", "nombre completo"),
            ("Evaluado y acordado por (2):", "nombre completo"),
            ("Fecha (AAAA-MM-DD):", "")), start=5):
        ws.cell(row=i, column=1, value=etq).font = Font(bold=True)
        c = ws.cell(row=i, column=2)
        c.fill = FIRMA
        if ayuda:
            c.comment = Comment(ayuda, "firma")
    return ws


def hoja(wb, inst, filas, ev, VA, VB, NT):
    ws = wb.create_sheet(inst["hoja"][:31])
    ws.freeze_panes = "C3"
    ws["A1"] = inst["nombre"] + " — CONSENSO"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Un solo juicio por dominio, acordado entre los dos autores.   ·   "
                "Fuente: " + inst["fuente"])
    ws["A2"].font = Font(italic=True, size=9)

    cab = [c[1] for c in CONTEXTO]
    cab += ["%s. %s" % (it["codigo"], it["texto_es"]) for it in inst["items"]]
    cab += ["JUICIO GLOBAL del estudio", "¿En qué frase te apoyaste?",
            "Dudas o desacuerdo con el diseño"]
    for j, c in enumerate(cab, 1):
        cel = ws.cell(row=3, column=j, value=c)
        cel.fill, cel.font = CAB, CABF
        cel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 78
    n_ctx = len(CONTEXTO)
    for j, (_, _, ancho, _) in enumerate(CONTEXTO, 1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    for j in range(n_ctx + 1, len(cab) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 30

    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"%s"' % ",".join(inst["juicios"]))
    ws.add_data_validation(dv)

    quedan = 0
    for i, f in enumerate(filas, start=4):
        s = f["study_id"]
        for j, (campo, _, _, _) in enumerate(CONTEXTO, 1):
            cel = ws.cell(row=i, column=j, value=f.get(campo, ""))
            cel.fill = FIJO
            cel.alignment = Alignment(wrap_text=(campo == "titulo"), vertical="top")
        codigos = [it["codigo"] for it in inst["items"]] + ["GLOBAL"]
        for k, cod in enumerate(codigos):
            cel = ws.cell(row=i, column=n_ctx + 1 + k)
            dv.add(cel)
            va = VA.get((s, inst["clave"], cod), "")
            vb = VB.get((s, inst["clave"], cod), "")
            trozos = []
            if k < len(inst["items"]):
                dom = inst["items"][k].get("dominio_evidencia", "")
                frases = ev.get(s, {}).get(dom, [])
                if frases:
                    trozos.append("Del artículo:\n\n" + "\n\n".join(frases)[:1500])
            if va == vb and va:
                cel.value, cel.fill = va, ACORDADO
            else:
                # Vacia a proposito: heredar uno de los dos valores seria elegir
                # por ellos, y es justo lo que hay que evitar.
                cel.fill = DECIDIR
                quedan += 1
                trozos.insert(0, "DECIDIR ENTRE LOS DOS.\n"
                                 "Fichero de D. Valdiviezo: %s\n"
                                 "Fichero de N. Trelles: %s" % (va or "(vacío)",
                                                                vb or "(vacío)"))
            if trozos:
                cel.comment = Comment("\n\n".join(trozos)[:2000], "consenso")
        col = n_ctx + 1 + len(codigos)
        # El texto libre se arrastra tal cual: es el trabajo de lectura ya
        # hecho, y es identico en los dos ficheros.
        frase, duda = NT.get((s, inst["clave"]), ("", ""))
        for valor in (frase, duda):
            cel = ws.cell(row=i, column=col, value=valor or None)
            cel.fill = FIRMA
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            col += 1
    return quedan


def main():
    for p in (A, B):
        if not p.exists():
            raise SystemExit("no encuentro %s" % p.name)
    VA, VB = lee_libro(A, SIMPLIFICADOS), lee_libro(B, SIMPLIFICADOS)
    NT = notas(A, SIMPLIFICADOS)
    ev = evidencia()
    filas = [c for c in corpus() if c["diseno"] in COMPARATIVOS
             and c["instrumento"] not in (NO_EVALUABLE, PENDIENTE)]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pendientes = 0
    hojas = []
    for inst in SIMPLIFICADOS:
        suyas = [f for f in filas if f["instrumento"] == inst["clave"]]
        if suyas:
            hojas.append((inst, suyas))
    for inst, suyas in hojas:
        pendientes += hoja(wb, inst, suyas, ev, VA, VB, NT)
    firma = hoja_firma(wb, pendientes)
    wb.move_sheet(firma, offset=-len(wb.sheetnames) + 1)

    if SALIDA.exists():
        raise SystemExit(
            "ya existe %s.\nNo se sobrescribe: podría llevar decisiones "
            "tomadas. Bórralo a mano si de verdad quieres regenerarlo."
            % SALIDA.name)
    wb.save(SALIDA)
    total = sum(len(s) * (len(i["items"]) + 1) for i, s in hojas)
    print("escrito %s" % SALIDA.relative_to(ROOT))
    print("   %d juicios en total" % total)
    print("   %d ya coincidían y entran precargados" % (total - pendientes))
    print("   %d en amarillo, VACÍOS: hay que decidirlos entre los dos"
          % pendientes)
    print("   falta además la firma de los dos autores en la hoja «Firma»")


if __name__ == "__main__":
    main()
