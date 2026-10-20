"""Hoja de datos para rellenar la carta de cesion de derechos de JSR.

QUE ES Y QUE NO ES. NO es la carta. La carta la provee la revista y hay que
descargarla de su seccion «Archivos y formatos descargables»; su lista de
comprobacion dice que el envio se devuelve si no llega en el formato de ellos.
Esto es la hoja con lo que va en cada hueco de ESE documento, para no tener que
buscar el titulo exacto ni volver a teclearlo.

DE DONDE SALE. El titulo se lee del manuscrito de envio, no se teclea. Lo
demas son datos de autor que el proyecto no tiene y que solo pueden poner
ellos: documento de identidad, ORCID y el correo de N. Trelles. Se dejan en
blanco a proposito y marcados.

LO QUE SE MIDE. Desde el 2026-10-12 tampoco se teclea lo que se dice del
manuscrito: el numero de referencias, las palabras del titulo y del resumen,
las palabras clave y la extension salen de las mismas funciones con las que
`build_jsr_submission.py` escribe el LEEME (`normas_formato`), y la
declaracion de uso de IA se comprueba con su misma guarda. La hoja decia
«20 referencias», «19 palabras de 20», «resumen de 250 palabras» y «las
cumple todas» tecleados, y la extension, que el PDF del 8 de octubre ya
ponia en 27 paginas contra un limite de 25, no la miraba nadie. La
extension se mide sobre el PDF maquetado: esta hoja va bien despues de
`build_jsr_pdf.py`; si va antes, lo dice.

LO QUE AVISA. Desde el 2026-10-13, ademas: si el protocolo de la fase
definitiva sigue diciendo que el informe piloto «no se enviará como tal» (y
firmar la cesion es enviarlo), y si el manuscrito hace responder por el trabajo
a los dos autores cuando N.T. no lo ha aceptado. Las dos cosas se leen de los
ficheros con las funciones del sobre, no se teclean. Desde el 2026-10-14 el
aviso del protocolo va al principio, y la hoja nombra la frase que deja
pendiente la responsabilidad de N.T. y la ultima extension medida. Desde el
2026-10-15, quien firmo las exclusiones a texto completo y lo que N.T. tiene que
conocer o ratificar antes de aprobar, con el mismo texto que el LEEME
(`pendiente_de_nt`) y las cifras de `synthesis_scalars.json`.

Salida:
    paper/docx/datos_carta_cesion.docx

Uso:
    python scripts/build_cesion_datos.py
"""
import json
import pathlib
import re
import sys

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches, Mm, RGBColor
except ImportError:
    raise SystemExit("hace falta python-docx: python -m pip install python-docx")

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_jsr_docx import escribe
from build_jsr_submission import (PROTOCOLO, TOPE_PAGINAS, TOPE_TITULO, con_punto,
                                   exige_declaracion_ia, firmas_exclusiones,
                                   frase_piloto, marcadores,
                                   medidas_formato, normas_formato, pendiente_de_nt,
                                   responsabilidad_ajena, responsabilidad_pendiente)

ROOT = pathlib.Path(__file__).resolve().parent.parent
JSR = ROOT / "paper" / "manuscrito_JSR_final.md"
DEST = ROOT / "paper" / "docx"
ESCALARES = ROOT / "quality_reports" / "synthesis_scalars.json"
TNR = "Times New Roman"

FORMULARIO = "https://revistas.utb.edu.ec/index.php/sr/about/submissions"

# Los ocho puntos que el firmante acepta, copiados de la carta de la revista el
# 14 de septiembre de 2026. Se reproducen aqui PARA COMPROBARLOS, no para
# sustituir al documento: junto a cada uno va si este trabajo lo cumple y donde
# consta. El que firme tiene que poder responder por los ocho.
#
# El texto de cada punto es el de la revista y no se toca. El estado que va al
# lado se rellena en `main` con lo que dice el manuscrito HOY: los huecos
# {entre llaves} no se escriben aqui a mano.
PUNTOS = [
    # DECIA solo «Lo firman los dos autores», y la aprobacion escrita de N.T.
    # no consta: el estado depende del marcador de su CRediT.
    ("Se esta de acuerdo con el contenido publicado en el mismo.",
     "{aprobacion}"),
    ("Cesion a titulo gratuito de la totalidad de los derechos patrimoniales de "
     "autor derivados del articulo, a favor de la citada publicacion.",
     "Decisión de los autores. No afecta a nada de lo construido."),
    # DECIA «N. Trelles segunda», con genero; y el orden no se miraba.
    ("Se esta de acuerdo con el orden en el que aparecen los autores.",
     "{orden}"),
    # DECIA «Cierto: la revision la hicieron dos personas». Es verdad de las
    # personas, pero callaba que un modelo de lenguaje intervino en todas las
    # fases, que es justo lo que este punto obliga a pensar. No se resume la
    # declaracion de IA: se cita su frase sobre la autoria.
    ("No existen personas adicionales que satisfagan los criterios de autoria y "
     "no hayan sido incluidas.",
     "{personas}"),
    ("La persona designada como autor de correspondencia es el unico contacto y "
     "el responsable de comunicarse con el resto de los autores y de autorizar "
     "la version final de publicacion.",
     "D. Valdiviezo figura como autor de correspondencia. ATENCIÓN: esto obliga "
     "a D. Valdiviezo a haber recabado la aprobación escrita de N. Trelles ANTES "
     "de firmar."),
    # DECIA «20 referencias en Vancouver, resueltas desde el .bib. Ninguna cita
    # sin entrada», con la cifra tecleada; la lista del manuscrito de la revista
    # no la resuelve ningun guion desde el .bib, asi que lo segundo no se
    # comprobaba. Y remitia a una declaracion de IA que ese manuscrito no llevo
    # hasta el 2026-10-08: ahora se comprueba con la guarda del sobre.
    ("El articulo contiene material original e inedito, el trabajo de terceros "
     "esta citado, y las citas textuales, parafrasis y referencias estan "
     "identificadas en el texto.",
     "{originalidad}"),
    ("El articulo no ha sido publicado anteriormente ni esta sometido a "
     "publicacion en otra revista.",
     "Cierto. El manuscrito preparado para Clinical Microbiology and Infection "
     "nunca se envió; sigue en el repositorio sin someter."),
    # DECIA que las normas «estan corregidas» y, abajo, «las cumple todas», sin
    # haber medido la extension.
    ("Se aceptan las condiciones de la Revista JSR en cuanto a normas, "
     "procedimientos, formato, edicion grafica, correccion y otros "
     "requerimientos que se solicitan en la convocatoria.",
     "{formato}"),
]

# Las cinco se corrigieron el 14 de septiembre. La lista se queda vacia a
# proposito y no se borra: si la revista cambia una norma que no se puede medir
# o alguien toca el constructor, aqui es donde vuelve a aparecer. Lo que si se
# mide (titulo, resumen, palabras clave, extension) no va aqui: sale de
# `normas_formato` en cada ejecucion.
INCUMPLE = []

# Lo que aplica el constructor del .docx y aqui no se mide. Las tres entradas
# con cifras --«resumen de 250 palabras», «cinco palabras clave» y «19
# palabras de 20»-- estaban tecleadas y ahora las mide `normas_formato`.
APLICADAS = [
    "A4", "márgenes de 3 cm", "Times New Roman 12",
    "una columna a interlineado 1,5", "párrafos justificados y sin espacio entre "
    "consecutivos",
    "las keywords en el mismo orden que las palabras clave, como pide la norma",
    "decimales con coma", "millares con espacio fino",
    "título en mayúscula sostenida, negrita y 18 pt",
    "resalte en cursiva y no en negrita",
    "palabras clave en negrita y cursiva, como pide la directriz",
    "figuras en PNG, además del PDF vectorial",
]

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def documento():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, lado, Inches(0.9))
    n = d.styles["Normal"]
    n.font.name, n.font.size = TNR, Pt(11)
    n.paragraph_format.space_after = Pt(7)
    n.paragraph_format.line_spacing = 1.25
    cab = s.header.paragraphs[0]
    cab.text = ("Hoja de datos — NO es la carta. La carta se descarga de la revista "
                "y se firma a mano.")
    cab.runs[0].font.size, cab.runs[0].font.name = Pt(9), TNR
    cab.runs[0].italic = True
    return d


def h(d, txt, size=13, antes=14):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(5)
    escribe(p, txt, size=size, negrita=True)


def fila(d, etiqueta, valor, rojo=False):
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(etiqueta + "  ")
    r.font.name, r.font.size, r.bold = TNR, Pt(11), True
    r2 = p.add_run(valor)
    r2.font.name, r2.font.size = TNR, Pt(11)
    if rojo:
        r2.font.color.rgb = RGBColor(0xA0, 0x30, 0x20)


def main():
    if not JSR.exists():
        raise SystemExit("no encuentro %s" % JSR.name)
    cuerpo = JSR.read_text(encoding="utf-8")
    lineas = cuerpo.split("\n")
    titulo_es = next(l[2:].strip() for l in lineas if l.startswith("# "))
    titulo_es = titulo_es.replace("*", "")
    fm = medidas_formato()
    cumple, falla, sin_medir = normas_formato(
        APLICADAS, despues="`scripts/build_cesion_datos.py`")
    # La frase la comparte el LEEME, que es markdown; en Word las comillas
    # invertidas se verian tal cual.
    sin_medir = sin_medir.replace("`", "") if sin_medir else None
    falla += ["%s: la norma %s; %s" % x for x in INCUMPLE]

    # El estado de cada punto, con lo que dice el manuscrito hoy.
    trelles = next((x for _, ms in marcadores(cuerpo) for x in ms
                    if x.startswith("‹‹") and "TRELLES" in x.upper()), None)
    try:
        ia, ia_mal = exige_declaracion_ia(), ""
    except SystemExit as e:
        ia, ia_mal = "", str(e).split("\n")[0]
    autoria_ia = ("Ninguna herramienta de IA figura como autora, conforme a las "
                  "recomendaciones del ICMJE.")
    # DECIA la frase entera tecleada --«redactó los borradores del manuscrito,
    # de los anexos y de la carta de presentación»-- y bastaba con que la
    # declaración añadiera «y de los registros de decisión» (2026-10-13) para
    # que la hoja dijera que no pasaba la guarda. Se cita lo que la declaración
    # dice, hasta el punto o el punto y coma, en vez de exigir unas palabras.
    m_red = re.search(r"redactó los borradores[^.;]*", ia.replace("*", ""))
    redaccion_ia = m_red.group(0).strip() if m_red else ""
    # La responsabilidad por el trabajo: N.T. no ha aceptado responder por él.
    resp = responsabilidad_ajena(cuerpo)
    # Y la frase que la deja pendiente, que cambia con el mismo escrito.
    resp_pendiente = responsabilidad_pendiente(cuerpo)
    piloto = frase_piloto()
    # Lo que N.T. tiene que saber o ratificar antes de aprobar (2026-10-15): 27
    # exclusiones registradas por orden de D.V. sin su firma, el mensaje del
    # 1-sep que redacto el modelo en primera persona de D.V. y 2 descartes de
    # solapamiento del modelo sin firma. La hoja no lo decia, y firmar el punto 1
    # es estar de acuerdo con ese contenido. Lo mismo que el LEEME.
    S = json.loads(ESCALARES.read_text(encoding="utf-8"))
    sabe_nt = pendiente_de_nt(S)
    m_orden = re.search(r"Danny Valdiviezo\*\*.{0,20}\*\*Nataly Trelles", cuerpo)
    aprobacion = ("ATENCIÓN: lo firman los dos autores, y la aprobación escrita de "
                  "N. Trelles sobre esta versión no consta todavía: su CRediT lleva "
                  "aún el marcador %s. Ver el punto pendiente al final." % trelles
                  if trelles else
                  "Lo firman los dos autores. Ver el punto pendiente al final.")
    if sabe_nt:
        aprobacion += (" Antes de aprobar, N. Trelles tiene que conocer o ratificar "
                       "lo que se lista al final de esta hoja.")
    if resp:
        # El manuscrito no puede hacer responder por el trabajo a quien no lo
        # ha aceptado; si vuelve a decirlo, la hoja lo avisa aqui, en rojo.
        aprobacion = ("%s Y el manuscrito dice «%s», cuando N. Trelles no ha "
                      "aceptado responder por el trabajo: mientras no conste por "
                      "escrito, la frase solo puede hablar por D. Valdiviezo."
                      % (aprobacion if aprobacion.startswith("ATENCIÓN") else
                         "ATENCIÓN: " + aprobacion, resp.rstrip(".")))
    datos = {
        "aprobacion": aprobacion,
        "orden": ("D. Valdiviezo en primer lugar y N. Trelles en segundo, como en "
                  "el manuscrito." if m_orden else
                  "OJO: no se encuentra en el manuscrito la línea de autores con "
                  "D. Valdiviezo y N. Trelles por ese orden; comprobadlo antes de "
                  "firmar."),
        "personas": (
            "Cierto en cuanto a personas: las únicas que intervinieron son los dos "
            "autores. Un modelo de lenguaje intervino en todas las fases y no "
            "puede ser autor; la declaración «Uso de inteligencia artificial» del "
            "manuscrito lo dice así: «%s»" % autoria_ia
            if autoria_ia in ia else
            "OJO: la declaración «Uso de inteligencia artificial» del manuscrito "
            "no dice ya que ninguna herramienta de IA figura como autora. Un modelo "
            "de lenguaje intervino en el trabajo: comprobadlo antes de firmar."),
        "originalidad": (
            "%d referencias en Vancouver en la lista del manuscrito. Ver también la "
            "declaración «Uso de inteligencia artificial» del manuscrito, en "
            "Declaraciones, que no se retira: dice, entre otras cosas, que el "
            "modelo de lenguaje «%s»." % (fm["referencias"], redaccion_ia)
            if ia and redaccion_ia else
            "OJO: la declaración «Uso de inteligencia artificial» del manuscrito "
            "no pasa la guarda del sobre (%s). No se firma este punto sin ella."
            % (ia_mal or "ya no dice que el modelo redactó los borradores")),
        # Decia «hoy no se cumplen 1»: el plural no se ajustaba.
        "formato": (
            "OJO: hoy no se %s %d: %s. El detalle, al final de esta hoja."
            % ("cumple" if len(falla) == 1 else "cumplen", len(falla),
               "; ".join(falla)) if falla else
            # Sin medir decia solo eso, y callaba que la ultima medida ya pasaba
            # del limite (2026-10-14). La cifra sale del PDF anterior.
            "OJO: la extensión no está medida sobre esta versión del manuscrito%s. "
            "El detalle, al final de esta hoja."
            % ("; la última medida, sobre el PDF maquetado anterior, daba %d páginas "
               "de cuerpo contra un límite de %d" % (fm["paginas_previas"], TOPE_PAGINAS)
               if fm["paginas_previas"] and fm["paginas_previas"] > TOPE_PAGINAS else "")
            if sin_medir else
            "Las cinco normas de formato que el manuscrito incumplía el 14 de "
            "septiembre están corregidas, y las que se miden se cumplen. El "
            "detalle, al final de esta hoja."),
    }

    d = documento()
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribe(p, "Datos para la carta de cesión de derechos de JSR", size=15, negrita=True)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    escribe(p, "Esta hoja **no es la carta**. La carta la provee la revista, se llama "
               "«Carta de originalidad y cesión de derechos» y se descarga de la sección "
               "«Archivos y formatos descargables» de %s. Su lista de comprobación avisa "
               "de que **no se inicia la revisión** si no llegan el artículo, esta carta y "
               "el formato de información de autores, en los formatos de ellos. Aquí está "
               "lo que va en cada hueco." % FORMULARIO, size=10.5)
    # ANTES QUE NADA, SI SE ENVIA. La hoja preparaba la firma de una cesion sin
    # decir que el protocolo de la fase definitiva dice que este informe no se
    # enviara como tal (consejo del 2026-10-13). Se cita del protocolo; si la
    # frase cambia, este parrafo cambia o desaparece. Iba al final de la hoja,
    # despues de todos los datos que rellenar: va aqui, antes de ellos
    # (2026-10-14).
    if piloto:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribe(p, "**Antes de firmar nada: si este informe se envía.** El protocolo "
                   "de la fase definitiva (borrador para OSF, %s) dice: «%s» Firmar "
                   "esta cesión es enviarlo. Los autores deciden antes: o no se envía, "
                   "o se cambia esa frase del protocolo antes de depositarlo y el "
                   "manuscrito se presenta como piloto." % (PROTOCOLO.name, piloto),
                size=10.5)

    h(d, "Lo que hay que escribir en la carta")
    fila(d, "Lugar y fecha:", "Cuenca, Ecuador — dd/mm/aa el día que se firme")
    fila(d, "Nombre del artículo:", titulo_es)
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    # DECIA «dentro del máximo de 20» sin compararlo: lo habría dicho igual con 21.
    escribe(p, "*(%d palabras, %s del máximo de %d que pide la revista. El binomio "
               "va en cursiva: Pseudomonas aeruginosa.)*"
            % (fm["titulo"], "dentro" if fm["titulo"] <= TOPE_TITULO else "POR ENCIMA",
               TOPE_TITULO), size=9.5)

    h(d, "Firma del Autor I", size=12, antes=12)
    fila(d, "Nombres y apellidos:", "Danny Valdiviezo")
    # Decia «solo lo tiene él»: sin pronombres de genero para los autores.
    fila(d, "Documento de identidad:", "PENDIENTE — solo lo tiene D. Valdiviezo",
         rojo=True)
    fila(d, "Correo electrónico:", "dvchiqui@gmail.com")
    fila(d, "Número ORCID:", "PENDIENTE (opcional para la revista)", rojo=True)
    fila(d, "Filiación:", "Facultad de Medicina, Universidad Católica de Cuenca, "
                          "Cuenca, Ecuador")

    h(d, "Firma del Autor II", size=12, antes=12)
    fila(d, "Nombres y apellidos:", "Nataly Trelles")
    fila(d, "Documento de identidad:", "PENDIENTE", rojo=True)
    fila(d, "Correo electrónico:", "PENDIENTE", rojo=True)
    fila(d, "Número ORCID:", "PENDIENTE (opcional para la revista)", rojo=True)
    fila(d, "Filiación:", "Facultad de Medicina, Universidad Católica de Cuenca, "
                          "Cuenca, Ecuador")

    h(d, "Los ocho puntos que se firman, y cómo está cada uno")
    p = d.add_paragraph()
    escribe(p, "Copiados de la carta de la revista el 14 de septiembre de 2026, para "
               "comprobarlos antes de firmar. Quien firma responde por los ocho.", size=10)
    for i, (punto, estado) in enumerate(PUNTOS, 1):
        estado = estado.format(**datos)
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(2)
        escribe(p, "%d. %s" % (i, punto), size=10.5)
        q = d.add_paragraph()
        q.paragraph_format.left_indent = Inches(0.6)
        q.paragraph_format.space_after = Pt(7)
        r = q.add_run(estado)
        r.font.name, r.font.size, r.italic = TNR, Pt(10), True
        r.font.color.rgb = (RGBColor(0xA0, 0x30, 0x20)
                            if estado.startswith(("OJO", "ATENCION", "ATENCIÓN")) else
                            RGBColor(0x44, 0x44, 0x44))

    h(d, "El punto 8: las normas de formato")
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    corregidas = ("Las cinco que fallaban ese día se corrigieron entonces: márgenes de "
                  "2,54 a 3 cm, el espacio entre párrafos consecutivos a cero, las "
                  "palabras clave en orden alfabético y con las keywords en ese mismo "
                  "orden, el título en mayúscula sostenida a 18 pt, y el resalte del "
                  "cuerpo pasado de negrita a cursiva. La negrita se conserva donde la "
                  "norma la pide o la da por supuesta: los títulos, los encabezados de "
                  "sección y de tabla, y la propia lista de palabras clave, que la "
                  "directriz exige «en negrita y cursiva».")
    if not falla and sin_medir is None:
        escribe(p, "Firmar el punto 8 es aceptar las normas de formato de la revista. "
                   "Leídas en su web el 14 de septiembre de 2026, el manuscrito **las "
                   "cumple todas**, medidas sobre esta versión. %s" % corregidas,
                size=10.5)
    else:
        # DECIA «las cumple todas» aunque la extension no se hubiera medido.
        escribe(p, "Firmar el punto 8 es aceptar las normas de formato de la revista, "
                   "leídas en su web el 14 de septiembre de 2026. %s %s"
                % (corregidas,
                   "Hoy **no cumple %d**:" % len(falla) if falla else
                   "Hoy no se puede decir que las cumpla todas:"), size=10.5)
    for que in falla:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(3)
        escribe(p, "• %s." % que, size=10.5)
    if sin_medir:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(3)
        escribe(p, "• %s" % sin_medir, size=10.5)
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    escribe(p, "Cumple en cambio: %s." % ", ".join(cumple), size=10.5)

    h(d, "Y lo que ninguna hoja resuelve")
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    escribe(p, "El punto 5 dice que el autor de correspondencia es **responsable de "
               "autorizar la versión final**. D. Valdiviezo no puede firmarlo de buena fe "
               "mientras N. Trelles no haya aprobado por escrito esta versión. Es el mismo "
               "criterio ICMJE que el manuscrito lleva pendiente, y aquí pasa de ser una "
               "buena práctica a ser una declaración firmada.%s"
            % (" Del mismo escrito depende el marcador %s del CRediT de N. Trelles en "
               "el manuscrito: lo que nombra se añade solo si consta por escrito, y si "
               "no consta el marcador se borra sin añadirlo." % trelles
               if trelles else ""), size=10.5)
    # La frase de responsabilidad de la declaracion de uso de IA depende del
    # mismo escrito de N.T. (2026-10-13): se cita del manuscrito.
    if resp_pendiente:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribe(p, "Del mismo escrito depende la frase de la declaración «Uso de "
                   "inteligencia artificial»: «%s» Cuando conste, se cambia en el "
                   "maestro y se copia igual al manuscrito de la revista; mientras no "
                   "conste, se queda como está." % resp_pendiente, size=10.5)
    # Lo que N.T. tiene que saber o ratificar antes de aprobar, con el texto del
    # LEEME (`pendiente_de_nt`); las comillas invertidas son de markdown.
    if sabe_nt:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribe(p, "**Antes de que N. Trelles apruebe esta versión y firme.** %s "
                   "Firmar el punto 1 es estar de acuerdo con ese contenido, y N. "
                   "Trelles tiene que conocer o ratificar antes esto:"
                % con_punto(firmas_exclusiones(S)), size=10.5)
        for titulo, texto in sabe_nt:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.space_after = Pt(3)
            escribe(p, "• **%s** %s" % (con_punto(titulo), texto.replace("`", "")),
                    size=10.5)

    DEST.mkdir(parents=True, exist_ok=True)
    salida = DEST / "datos_carta_cesion.docx"
    try:
        d.save(salida)
    except PermissionError:
        salida = DEST / "datos_carta_cesion_NUEVO.docx"
        d.save(salida)
        print("AVISO: estaba abierto en Word; escrito al lado como %s" % salida.name)
    print("escrito %s  (%d KB)" % (salida.relative_to(ROOT), salida.stat().st_size // 1024))
    print("  titulo leido del manuscrito: %d palabras" % fm["titulo"])
    print("  %d normas de formato incumplidas%s"
          % (len(falla), "; extension SIN MEDIR sobre esta version" if sin_medir else ""))
    print("  declaracion de uso de IA: %s" % ("literal desde el maestro" if ia else
                                              "NO PASA LA GUARDA: " + ia_mal))
    if sabe_nt:
        print("  pendiente de N.T. antes de aprobar: %d punto(s)" % len(sabe_nt))
    if resp:
        print("  AVISO: el manuscrito hace responder por el trabajo a los dos autores")
    if piloto:
        print("  AVISO: el protocolo dice que el informe piloto no se enviara como tal")


if __name__ == "__main__":
    main()
