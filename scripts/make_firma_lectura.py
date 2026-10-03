# -*- coding: utf-8 -*-
"""Lo que la lectura firmada del 2026-09-30 deja abierto, para firmar.

POR QUE UN CUADERNO. `ingest_lectura_firmada.py` aplico lo que la lectura deja
inequivoco. Lo que queda son decisiones de autoria que un guion no toma:

  1. Nueve estudios sin ningun paciente con P. aeruginosa tratado con fagos.
     No cumplen el criterio de poblacion. Excluir es un acto de autoria.
  2. Tres estudios de un solo brazo que la lectura llevo por debajo del umbral
     de multirresistencia (EST-053, EST-057, EST-088). EST-001 y EST-094 estan
     en la misma situacion y se firmo mantenerlos el 2026-09-22; estos tres no
     tienen esa firma.
  3. Nueve casos de clase o de DTR que la lectura deja con mas de un valor
     posible para el brazo.
  4. Como se hizo la lectura. El manuscrito tiene que decir quien leyo los
     articulos y si intervino un modelo de lenguaje, y eso solo lo sabeis
     vosotros. Mientras no se conteste, los Metodos dicen solo lo que consta:
     que cada conclusion lleva la firma de los dos.

Cada fila trae lo que firmasteis el 30 de septiembre, tal cual, y lo que cambia
segun lo que decidais. Las casillas naranjas son las vuestras.

EL SEGURO. Si el cuaderno ya existe con alguna casilla naranja rellena, no se
escribe.

SALIDA
    ~/Desktop/FIRMAR_lectura_2026-10-01.xlsx
"""
import csv
import json
import pathlib
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
RS = ROOT / "revision_sistematica"
QR = ROOT / "quality_reports"
LP = RS / "lectura_pendiente"
CUADERNO = pathlib.Path.home() / "Desktop" / "FIRMAR_lectura_2026-10-01.xlsx"

# Estudios con P. aeruginosa en el articulo, pero que no era la diana del fago
# cuando se aplico. La lectura los dejo en «no verificable», no en «no aplica»;
# van aparte, para que se decidan sabiendo que son un caso distinto.
DUDOSOS = {
    "EST-051": "El cóctel administrado (Intesti) solo fue activo contra E. coli, y la "
               "P. aeruginosa no volvió a aislarse antes de aplicar el fago.",
    "EST-096": "P. aeruginosa se detectó 85 días antes de la fagoterapia y la orina ya "
               "era negativa al iniciarla (paciente 10).",
}

OPC_EXCL = ["EXCLUIR con el código ORG", "MANTENER en el corpus"]
OPC_UMBRAL = ["MANTENER y declararlo, como EST-001 y EST-094", "EXCLUIR"]
OPC_LECTURA = [
    "La leímos los dos, cada uno los artículos completos",
    "La leyó un autor y la revisó el otro",
    "Con un modelo de lenguaje y la revisamos los dos",
    "Otra: la explicamos en «por qué»",
]

# Lo abierto de clase y DTR: la pregunta sale de `queda_abierto`; las opciones,
# de lo que la propia lectura dice que es posible.
OPC_ABIERTO = {
    "EST-004": ["Dejar A en «yes» y B en «no»", "Poner los dos en «not-derivable»"],
    "EST-019": ["Dejar el brazo en «no»", "Poner «not-derivable»"],
    "EST-042": ["Corregir a «not-classifiable»", "Dejar «PDR» y declarar que no se sostiene",
                "Partir el brazo por paciente (requiere reextraer desenlaces)"],
    "EST-047": ["Dejar «no»", "Poner «yes» (manda el aislado 2)", "Poner «not-derivable»"],
    "EST-048": ["Corregir a «not-classifiable»", "Dejar «MDR»"],
    "EST-058": ["Dejar «no» (la «I» de EUCAST es sensible)",
                "Poner «yes» (la «I» cuenta como no sensible, Kadri)"],
    "EST-106": ["Corregir a «below-MDR-threshold»", "Corregir a «not-classifiable»",
                "Dejar «MDR»"],
    "EST-124": ["Manda el aislado del día de la fagoterapia: «below-MDR-threshold»",
                "Manda el aislado PA02: «XDR» y DTR «yes»", "Corregir a «not-classifiable»"],
    "EST-169": ["Dejar «not-classifiable» y declararlo mixto",
                "Partir el brazo por paciente (requiere reextraer desenlaces)"],
}


def leer(p, enc="utf-8-sig"):
    with open(p, encoding=enc, newline="") as fh:
        return list(csv.DictReader(fh))


def ya_contestado():
    if not CUADERNO.exists():
        return []
    from openpyxl import load_workbook
    wb = load_workbook(CUADERNO)
    puestas = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if (c.fill and c.fill.fgColor and str(c.fill.fgColor.rgb).endswith("FCE4D6")
                        and c.value not in (None, "")):
                    puestas.append("%s!%s" % (ws.title, c.coordinate))
    return puestas


def que_cambia(est, adj, rob, grupo, sol, S):
    """Lo que el estudio pesa hoy en el manuscrito, para decidir sabiendolo."""
    br = adj.get(est, [])
    partes = ["%d brazo(s), %s paciente(s)" % (
        len(br), "+".join(b["n_arm"] for b in br) or "?")]
    disenos = sorted({b["study_design"] for b in br})
    partes.append("diseño " + " / ".join(disenos))
    if est in rob:
        partes.append("evaluado con %s en riesgo de sesgo (sale de la Tabla 7)"
                      % {"rob2": "RoB 2", "robins": "ROBINS-I"}.get(rob[est], rob[est]))
    if grupo.get(est) == "si":
        partes.append("es uno de los %d con grupo de comparación"
                      % S["comparativos_con_grupo_real"])
    pares = [p for p in sol if est in (p["estudio_1"], p["estudio_2"])
             and p["veredicto"].startswith("SOLAPAMIENTO")]
    if pares:
        partes.append("comparte pacientes con %s" % ", ".join(
            p["estudio_2"] if p["estudio_1"] == est else p["estudio_1"] for p in pares))
    return "; ".join(partes)


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    puestas = ya_contestado()
    if puestas:
        print("NO SE ESCRIBE: el cuaderno ya tiene respuestas en %s" % ", ".join(puestas[:6]),
              file=sys.stderr)
        raise SystemExit(1)

    S = json.load(open(QR / "synthesis_scalars.json", encoding="utf-8"))
    lect = {r["study_id"]: r for r in leer(LP / "lectura_resistencia_firmada.csv")}
    adj = {}
    for r in leer(RS / "extraccion" / "extraccion_adjudicada.csv", "utf-8"):
        adj.setdefault(r["study_id"], []).append(r)
    rob = {}
    for r in leer(RS / "riesgo_sesgo" / "riesgo_sesgo_comparativos_adjudicado.csv"):
        rob[r["study_id"]] = r["instrumento"]
    grupo = {r["study_id"]: r["tiene_grupo_de_comparacion"]
             for r in leer(RS / "riesgo_sesgo" / "grupo_de_comparacion_real.csv", "utf-8")}
    sol = leer(QR / "pendiente2_solapamiento.csv")
    titulos = {}
    for r in leer(QR / "pendiente_lectura_resistencia.csv"):
        titulos[r["study_id"]] = r["titulo"]

    negrita, grande = Font(bold=True), Font(bold=True, size=13)
    blanca = Font(bold=True, color="FFFFFF")
    azul = PatternFill("solid", fgColor="44546A")
    ojo = PatternFill("solid", fgColor="FCE4D6")
    arriba = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()
    h = wb.active
    h.title = "Empieza aqui"
    h.column_dimensions["A"].width = 110
    texto = [
        ("Lo que vuestra lectura del 30 de septiembre deja para decidir", True),
        ("", False),
        ("La lectura ya está ingerida en todo lo que dejaba inequívoco: 35 correcciones de", False),
        ("clase, procedencia y DTR, los 32 pares de solapamiento y los 14 comparativos.", False),
        ("Lo que queda aquí no lo decide un script:", False),
        ("", False),
        ("  1. %d estudios sin ningún paciente con P. aeruginosa tratado con fagos, y %d dudosos."
         % (S["lectura_clase_no_aplica"], len(DUDOSOS)), False),
        ("  2. %d estudios de un solo brazo que pasaron por debajo del umbral de multirresistencia."
         % S["bajo_umbral_sin_decision"], False),
        ("  3. %d casos de clase o de DTR que dejan el brazo con más de un valor posible."
         % S["lectura_queda_abierto"], False),
        ("  4. Cómo se hizo la lectura, que el manuscrito tiene que declarar.", False),
        ("", False),
        ("Cada fila lleva lo que firmasteis, tal cual, y lo que pesa hoy ese estudio en el", False),
        ("manuscrito. Escoged en el desplegable de la casilla naranja y, si hace falta,", False),
        ("explicad en «por qué». Sin las dos firmas de la última hoja no se ingiere nada.", False),
        ("", False),
        ("Mientras tanto, el manuscrito declara estos puntos como pendientes de vuestra", True),
        ("decisión, con los estudios nombrados. No se puede enviar así.", True),
    ]
    for i, (t, b) in enumerate(texto, start=1):
        h.cell(row=i, column=1, value=t).font = negrita if b else Font()

    def hoja(nombre, titulo, cab, filas, anchos, col_dec, opciones_por_fila):
        s = wb.create_sheet(nombre)
        s.cell(row=1, column=1, value=titulo).font = grande
        for j, c in enumerate(cab, start=1):
            x = s.cell(row=3, column=j, value=c)
            x.font, x.fill, x.alignment = blanca, azul, arriba
            s.column_dimensions[x.column_letter].width = anchos[j - 1]
        for i, f in enumerate(filas, start=4):
            for j, v in enumerate(f, start=1):
                s.cell(row=i, column=j, value=v).alignment = arriba
            for j in range(col_dec, len(cab) + 1):
                s.cell(row=i, column=j).fill = ojo
            ops = opciones_por_fila(i - 4)
            dv = DataValidation(type="list", allow_blank=True,
                                formula1='"%s"' % ",".join(o.replace(",", ";")[:80] for o in ops))
            dv.error = "Escoja una de las opciones."
            s.add_data_validation(dv)
            dv.add(s.cell(row=i, column=col_dec))
            s.row_dimensions[i].height = 150
        s.freeze_panes = "B4"
        return s

    # 1. sin P. aeruginosa
    cab = ["Estudio", "Título", "Lo que firmasteis el 30-sep", "Frase en que os apoyasteis",
           "Qué pesa hoy en el manuscrito", "Vuestra decisión", "Por qué", "Frase del artículo"]
    filas = []
    for e in S["lectura_clase_no_aplica_ids"] + sorted(DUDOSOS):
        r = lect[e]
        firm = r["clase_verificada_firmada"]
        if e in DUDOSOS:
            firm = "[DUDOSO] " + DUDOSOS[e] + "\n\nLo firmado: " + firm
        filas.append([e, titulos.get(e, ""), firm, r["criterio_o_antibiograma_firmado"][:900],
                      que_cambia(e, adj, rob, grupo, sol, S), "", "", ""])
    hoja("1 Sin P. aeruginosa",
         "¿Salen del corpus los estudios sin ningún paciente con P. aeruginosa tratado con fagos?",
         cab, filas, [10, 40, 60, 60, 42, 26, 36, 40], 6, lambda i: OPC_EXCL)

    # 2. bajo el umbral, sin decision
    filas = []
    for e in S["bajo_umbral_sin_decision_ids"]:
        r = lect[e]
        filas.append([e, titulos.get(e, ""), r["clase_verificada_firmada"],
                      r["criterio_o_antibiograma_firmado"][:900],
                      que_cambia(e, adj, rob, grupo, sol, S), "", "", ""])
    hoja("2 Bajo el umbral",
         "¿Se mantienen, como EST-001 y EST-094, los estudios de un solo brazo por debajo del umbral?",
         cab, filas, [10, 40, 60, 60, 42, 30, 36, 40], 6, lambda i: OPC_UMBRAL)

    # 3. clase y DTR abiertos
    ids = S["lectura_queda_abierto_ids"]
    cab3 = ["Estudio", "Lo que queda abierto", "Lo que firmasteis el 30-sep",
            "Opciones", "Vuestra decisión", "Por qué", "Frase del artículo"]
    filas = []
    for e in ids:
        r = lect[e]
        filas.append([e, r["queda_abierto"], r["clase_verificada_firmada"],
                      "\n".join("· " + o for o in OPC_ABIERTO.get(e, [])), "", "", ""])
    hoja("3 Clase y DTR abiertos", "Brazos con más de un valor posible de clase o de DTR",
         cab3, filas, [10, 50, 70, 42, 30, 36, 40], 5,
         lambda i: OPC_ABIERTO.get(ids[i], ["Otra: explicadla en «por qué»"]))

    # 4. como se hizo
    s = wb.create_sheet("4 Cómo se hizo")
    s.column_dimensions["A"].width = 58
    s.column_dimensions["B"].width = 70
    s.cell(row=1, column=1, value="Cómo se hizo la lectura del 30 de septiembre").font = grande
    preguntas = [
        ("¿Quién leyó los artículos y rellenó las tres hojas?", OPC_LECTURA),
        ("Si intervino un modelo de lenguaje: ¿cuál, y qué hizo exactamente? "
         "(buscar en el PDF, transcribir figuras, redactar la conclusión...)", None),
        ("Las 23 transcripciones marcadas «[imagen]» (antibiogramas leídos de figuras): "
         "¿las comprobó una persona contra la figura?",
         ["Sí, todas", "Algunas: decid cuáles en la casilla de al lado", "No"]),
    ]
    f = 3
    for q, ops in preguntas:
        s.cell(row=f, column=1, value=q).alignment = arriba
        c = s.cell(row=f, column=2)
        c.fill, c.alignment = ojo, arriba
        if ops:
            dv = DataValidation(type="list", allow_blank=True,
                                formula1='"%s"' % ",".join(o.replace(",", ";") for o in ops))
            s.add_data_validation(dv)
            dv.add(c)
        s.row_dimensions[f].height = 48
        f += 2
    s.cell(row=f, column=1, value=(
        "Por qué se pregunta: el manuscrito no puede presentar como humano un "
        "procedimiento que no lo fue, y declara el uso de modelos de lenguaje en el "
        "cribado. Las 457 citas de texto del cuaderno se comprobaron de forma mecánica "
        "y todas están en su PDF y en la página que dice; las 23 transcripciones de "
        "figuras no se pueden comprobar así.")).alignment = arriba
    s.row_dimensions[f].height = 80

    fi = wb.create_sheet("Firma")
    fi.column_dimensions["A"].width = 46
    fi.column_dimensions["B"].width = 54
    fi.cell(row=1, column=1, value="Firma de los dos autores").font = negrita
    fi.cell(row=2, column=1, value="Las decisiones son de los dos: se firman juntas o no se "
                                   "firma ninguna.")
    for i, k in enumerate(["Revisor 1 (nombre completo)", "Revisor 2 (nombre completo)",
                           "Fecha (AAAA-MM-DD)",
                           "¿Habéis leído los artículos citados en cada hoja?"], start=4):
        fi.cell(row=i, column=1, value=k).font = negrita
        fi.cell(row=i, column=2, value="").fill = ojo

    wb.save(CUADERNO)
    print("cuaderno: %s" % CUADERNO)
    print("  1. sin P. aeruginosa      %d (+%d dudosos)" % (S["lectura_clase_no_aplica"], len(DUDOSOS)))
    print("  2. bajo el umbral         %d" % S["bajo_umbral_sin_decision"])
    print("  3. clase y DTR abiertos   %d" % S["lectura_queda_abierto"])
    print("  4. cómo se hizo la lectura")


if __name__ == "__main__":
    main()
