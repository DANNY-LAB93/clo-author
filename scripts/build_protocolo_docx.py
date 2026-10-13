"""El protocolo de la fase definitiva, en .docx, para adjuntarlo al registro de OSF.

La fuente unica es `revision_sistematica/protocolo/protocolo_revision_sistematica.md`.
El .docx se rehace desde ahi y no se edita a mano: si se edita, la siguiente
pasada lo pisa, y el texto depositado dejaria de coincidir con el del
repositorio.

Se niega a escribir si el protocolo conserva algun marcador ‹‹...›› sin
rellenar, salvo con --borrador: un protocolo con huecos no se deposita.

Uso:
    python scripts/build_protocolo_docx.py --borrador
    python scripts/build_protocolo_docx.py
"""
import argparse
import pathlib
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parent.parent
FUENTE = ROOT / "revision_sistematica" / "protocolo" / "protocolo_revision_sistematica.md"
SALIDA = FUENTE.with_suffix(".docx")
TNR = "Times New Roman"


def runs(p, texto, size=11, mono=False):
    """Resuelve **negrita**, *cursiva* y `codigo` en linea."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto):
        if not trozo:
            continue
        r = p.add_run(trozo.strip("*`") if trozo[0] in "*`" else trozo)
        r.font.name = "Consolas" if (mono or trozo.startswith("`")) else TNR
        r.font.size = Pt(size)
        r.bold = trozo.startswith("**")
        r.italic = trozo.startswith("*") and not trozo.startswith("**")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--borrador", action="store_true",
                    help="permite marcadores sin rellenar y marca el documento como borrador")
    args = ap.parse_args()

    texto = FUENTE.read_text(encoding="utf-8")
    huecos = re.findall(r"‹‹[^›]*››", texto)
    # Se cuenta tambien la apertura a secas: un marcador con chevrones simples
    # dentro («‹k›») no casa con la expresion y el del umbral de agrupacion se
    # escapaba asi, el unico que de verdad tiene que bloquear el deposito.
    if texto.count("‹‹") != len(huecos):
        huecos += ["‹‹(marcador mal cerrado)››"] * (texto.count("‹‹") - len(huecos))
    if huecos and not args.borrador:
        raise SystemExit("el protocolo tiene %d marcadores sin rellenar, por ejemplo %s. "
                         "Un protocolo con huecos no se deposita; usa --borrador para "
                         "verlo." % (len(huecos), huecos[0][:80]))

    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, lado, Mm(25.4))
    n = d.styles["Normal"]
    n.font.name, n.font.size = TNR, Pt(11)

    if args.borrador:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(p, "**BORRADOR — NO DEPOSITAR: %d marcadores sin rellenar**" % len(huecos))

    en_codigo = False
    for linea in texto.split("\n"):
        if linea.startswith("```"):
            en_codigo = not en_codigo
            continue
        if en_codigo:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            runs(p, linea, size=9, mono=True)
            continue
        if not linea.strip() or linea.strip() == "---":
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", linea)
        if m:
            h = d.add_heading(level=min(len(m.group(1)), 3))
            runs(h, m.group(2), size={1: 15, 2: 13}.get(len(m.group(1)), 11))
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", linea)
        if m:
            p = d.add_paragraph(style="List Bullet" if not m.group(1) else "List Bullet 2")
            runs(p, m.group(2))
            continue
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", linea)
        if m:
            p = d.add_paragraph(style="List Number")
            runs(p, m.group(2))
            continue
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        runs(p, linea)

    try:
        d.save(SALIDA)
    except PermissionError:
        raise SystemExit("%s esta abierto en Word; cierralo y vuelve a correr" % SALIDA.name)
    print("escrito %s%s" % (SALIDA.relative_to(ROOT),
                            "  (BORRADOR, %d marcadores)" % len(huecos) if args.borrador else ""))


if __name__ == "__main__":
    main()
